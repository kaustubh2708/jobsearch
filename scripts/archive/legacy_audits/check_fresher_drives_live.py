import urllib.request, ssl, json, time, re
from bs4 import BeautifulSoup
import openpyxl

print("=== CHECKING ACTIVE FRESHER HIRING DRIVES (OCTOBER 2026) ===")

wb = openpyxl.load_workbook('data/Job Search.xlsx')
ws = wb['Fresher Hiring Drives']

ctx = ssl.create_default_context()
ctx.check_hostname = False
ctx.verify_mode = ssl.CERT_NONE
headers = {
    'User-Agent': 'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36'
}

# 1. Inspect existing drives in sheet
sheet_drives = []
for r in range(2, 78):
    comp = ws.cell(r, 1).value
    prog = ws.cell(r, 2).value
    loc = ws.cell(r, 3).value
    win = ws.cell(r, 4).value
    cell_u = ws.cell(r, 5)
    u = cell_u.hyperlink.target if cell_u.hyperlink else cell_u.value
    if comp:
        sheet_drives.append({
            'row': r,
            'company': comp.strip(),
            'program': prog.strip() if prog else '',
            'location': loc.strip() if loc else '',
            'window': win.strip() if win else '',
            'url': u.strip() if u else ''
        })

print(f"Total documented drive programs: {len(sheet_drives)}")

# Filter for programs currently active in October 2026
active_window_drives = []
for d in sheet_drives:
    win = d['window'].lower()
    comp = d['company'].lower()
    if 'amazon' in comp or 'mongodb' in comp or 'sarvam' in comp:
        continue
    # Active signals for October
    is_active = any(k in win for k in ['2026-10', '2026-11', '2026-12', 'october', 'annual intake', 'annual campus', 'annual college'])
    if is_active:
        active_window_drives.append(d)

print(f"Programs with active October 2026 windows: {len(active_window_drives)}")
for d in active_window_drives[:15]:
    print(f"  * {d['company']} | {d['program']} | Window: {d['window']} | Link: {d['url'] or 'TBD'}")

