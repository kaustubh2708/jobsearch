import openpyxl, json
from copy import copy
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

with open('data/curated_wave13_extended_final.json') as f:
    new_roles = json.load(f)

# Styles
verified_fill = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid') # Soft green
white_fill = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')
font_regular = Font(name='Calibri', size=11, bold=False, color='000000')
font_link = Font(name='Calibri', size=11, bold=False, color='0563C1', underline='single')
center_align = Alignment(horizontal='center', vertical='center')
left_align = Alignment(horizontal='left', vertical='center')
thin_side = Side(border_style='thin', color='D9D9D9')
thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

files = ['data/Companies(1).xlsx', 'data/companies.xlsx']

for target_file in files:
    wb = openpyxl.load_workbook(target_file)
    ws = wb['Wave 13']
    
    start_row = ws.max_row + 1
    current_index = ws.max_row # last serial number was max_row - 1
    
    print(f"Target: {target_file}. Current max_row: {ws.max_row}. Appending {len(new_roles)} rows starting at {start_row}...")

    for i, role in enumerate(new_roles):
        r = start_row + i
        s_no = current_index + i
        
        # Cols:
        # 1: #
        # 2: Company
        # 3: Role Title
        # 4: Location (live)
        # 5: Experience (per JD)
        # 6: Key Tech Stack (estimate)
        # 7: Estimated Compensation (estimate)
        # 8: Match Score
        # 9: Priority / Match Tier
        # 10: Link Verification Status
        # 11: Direct Application URL
        # 12: Notes

        c1 = ws.cell(r, 1, s_no)
        c1.alignment = center_align
        c1.font = font_regular
        c1.fill = white_fill
        c1.border = thin_border

        c2 = ws.cell(r, 2, role['company'])
        c2.alignment = left_align
        c2.font = font_regular
        c2.fill = white_fill
        c2.border = thin_border

        c3 = ws.cell(r, 3, role['title'])
        c3.alignment = left_align
        c3.font = font_regular
        c3.fill = white_fill
        c3.border = thin_border

        c4 = ws.cell(r, 4, role['location'])
        c4.alignment = left_align
        c4.font = font_regular
        c4.fill = white_fill
        c4.border = thin_border

        c5 = ws.cell(r, 5, role['experience'])
        c5.alignment = left_align
        c5.font = font_regular
        c5.fill = white_fill
        c5.border = thin_border

        c6 = ws.cell(r, 6, role['tech_stack'])
        c6.alignment = left_align
        c6.font = font_regular
        c6.fill = white_fill
        c6.border = thin_border

        c7 = ws.cell(r, 7, role['comp'])
        c7.alignment = left_align
        c7.font = font_regular
        c7.fill = white_fill
        c7.border = thin_border

        c8 = ws.cell(r, 8, role['score'])
        c8.alignment = center_align
        c8.font = font_regular
        c8.fill = white_fill
        c8.border = thin_border

        c9 = ws.cell(r, 9, role['tier'])
        c9.alignment = left_align
        c9.font = font_regular
        c9.fill = white_fill
        c9.border = thin_border

        c10 = ws.cell(r, 10, role['link_status'])
        c10.alignment = left_align
        c10.font = font_regular
        c10.fill = verified_fill
        c10.border = thin_border

        c11 = ws.cell(r, 11, 'Exact Requisition Link ↗')
        c11.hyperlink = role['url']
        c11.alignment = center_align
        c11.font = font_link
        c11.fill = white_fill
        c11.border = thin_border

        c12 = ws.cell(r, 12, role['notes'])
        c12.alignment = left_align
        c12.font = font_regular
        c12.fill = white_fill
        c12.border = thin_border

    wb.save(target_file)
    print(f"Successfully saved {target_file}. New max_row: {ws.max_row}.")

