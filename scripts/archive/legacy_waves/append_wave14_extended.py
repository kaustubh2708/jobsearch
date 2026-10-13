import json, openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.load_workbook('data/Job Search.xlsx')
ws = wb['Wave 14']

with open('data/wave14_extended_pristine_final.json') as f:
    new_roles = json.load(f)

start_row = ws.max_row + 1
start_num = ws.max_row # row 29 has # 28, so row 30 gets # 29

print(f"Current max_row: {ws.max_row}. Appending {len(new_roles)} new roles starting at row {start_row} (# {start_num}).")

font_regular = Font(name='Calibri', size=10, bold=False, color='000000')
font_bold = Font(name='Calibri', size=10, bold=True, color='000000')
font_link = Font(name='Calibri', size=10, color='000563C1', underline='single')

align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)
align_center = Alignment(horizontal='center', vertical='center')

border_thin = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

fill_even = PatternFill(start_color='F9FAFB', end_color='F9FAFB', fill_type='solid')
fill_odd = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')

for idx, role in enumerate(new_roles):
    current_row = start_row + idx
    current_num = start_num + idx
    fill = fill_even if current_row % 2 == 0 else fill_odd

    ws.row_dimensions[current_row].height = 24

    # 1. #
    c1 = ws.cell(row=current_row, column=1, value=current_num)
    c1.font = font_bold
    c1.alignment = align_center
    c1.border = border_thin
    c1.fill = fill

    # 2. Company
    c2 = ws.cell(row=current_row, column=2, value=role['company'])
    c2.font = font_bold
    c2.alignment = align_left
    c2.border = border_thin
    c2.fill = fill

    # 3. Role Title
    c3 = ws.cell(row=current_row, column=3, value=role['title'])
    c3.font = font_regular
    c3.alignment = align_left
    c3.border = border_thin
    c3.fill = fill

    # 4. Location
    c4 = ws.cell(row=current_row, column=4, value=role['location'])
    c4.font = font_regular
    c4.alignment = align_left
    c4.border = border_thin
    c4.fill = fill

    # 5. Domain / Track
    c5 = ws.cell(row=current_row, column=5, value=role['domain'])
    c5.font = font_regular
    c5.alignment = align_left
    c5.border = border_thin
    c5.fill = fill

    # 6. Exact Stated Experience in JD (0–2 YOE)
    c6 = ws.cell(row=current_row, column=6, value=role['experience_stated'])
    c6.font = font_regular
    c6.alignment = align_left
    c6.border = border_thin
    c6.fill = fill

    # 7. Key Tech Stack
    c7 = ws.cell(row=current_row, column=7, value=role['tech_stack'])
    c7.font = font_regular
    c7.alignment = align_left
    c7.border = border_thin
    c7.fill = fill

    # 8. Estimated / Stated Compensation
    c8 = ws.cell(row=current_row, column=8, value=role['compensation'])
    c8.font = font_bold
    c8.alignment = align_left
    c8.border = border_thin
    c8.fill = fill

    # 9. Match Score
    c9 = ws.cell(row=current_row, column=9, value=role['match_score'])
    c9.font = font_bold
    c9.alignment = align_center
    c9.border = border_thin
    c9.fill = fill

    # 10. Priority / Match Tier
    c10 = ws.cell(row=current_row, column=10, value=role['priority_tier'])
    c10.font = font_regular
    c10.alignment = align_center
    c10.border = border_thin
    c10.fill = fill

    # 11. Link Verification Status
    c11 = ws.cell(row=current_row, column=11, value=role['verification_status'])
    c11.font = font_regular
    c11.alignment = align_center
    c11.border = border_thin
    c11.fill = fill

    # 12. Direct Application URL
    c12 = ws.cell(row=current_row, column=12, value='Direct Link')
    c12.hyperlink = role['url']
    c12.font = font_link
    c12.alignment = align_left
    c12.border = border_thin
    c12.fill = fill

    # 13. Role Fit & Strategy Notes
    c13 = ws.cell(row=current_row, column=13, value=role['notes'])
    c13.font = font_regular
    c13.alignment = align_left
    c13.border = border_thin
    c13.fill = fill

wb.save('data/Job Search.xlsx')
print(f"Successfully appended {len(new_roles)} rows into Wave 14. New max_row: {ws.max_row} (Total roles: {ws.max_row - 1}).")

