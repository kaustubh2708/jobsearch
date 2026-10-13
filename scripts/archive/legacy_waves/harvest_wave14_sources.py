import urllib.request, json, ssl, time, re
import openpyxl

# 1. Compile all history from Prev Waves, Wave 13
wb = openpyxl.load_workbook('data/Job Search.xlsx')
history_companies = set()
history_urls = set()

ws_prev = wb['Prev Waves']
for r in range(2, ws_prev.max_row + 1):
    c = ws_prev.cell(r, 3).value
    u = ws_prev.cell(r, 10).hyperlink.target if ws_prev.cell(r, 10).hyperlink else ws_prev.cell(r, 10).value
    if c: history_companies.add(c.strip().lower())
    if u: history_urls.add(u.strip())

ws_w13 = wb['Wave 13']
for r in range(2, ws_w13.max_row + 1):
    c = ws_w13.cell(r, 2).value
    u = ws_w13.cell(r, 12).hyperlink.target if ws_w13.cell(r, 12).hyperlink else ws_w13.cell(r, 12).value
    if c: history_companies.add(c.strip().lower())
    if u: history_urls.add(u.strip())

print(f"Master baseline for Wave 14: {len(history_companies)} companies, {len(history_urls)} URLs")

with open('data/master_history_dedup_wave14.json', 'w') as f:
    json.dump({'companies': sorted(list(history_companies)), 'urls': sorted(list(history_urls))}, f, indent=2)

ctx = ssl.create_default_context()
ctx.check_hostname = False
ctx.verify_mode = ssl.CERT_NONE

# 2. Harvest fresh Greenhouse boards
gh_boards = [
    'datadog', 'mongodb', 'yugabyte', 'confluent', 'tibco', 'canonical',
    'qualys', 'rapid7', 'cyberark', 'sailpoint', 'pingidentity', 'docusign',
    'box', 'dropbox', 'samsara', 'toast', 'verkada', 'cloudflare', 'palantir',
    'snowflake', 'nutanix', 'informatica', 'synopsys', 'cadence', 'marvell',
    'arm', 'siemens', 'nxp', 'mathworks', 'akamai', 'splunk', 'netapp',
    'citrix', 'teradata', 'veritas', 'commvault', 'cohesity', 'purestorage',
    'tanium', 'crowdstrike', 'sentinelone', 'fortinet', 'paloaltonetworks',
    'zscaler', 'fastly', 'checkpoint', 'proofpoint', 'elastic', 'gitlab',
    'pagerduty', 'twilio', 'zendesk', 'hubspot', 'rubrik', 'benchling',
    'cohere', 'dbtlabs', 'ramp', 'retool', 'snyk', 'vanta', 'warp', 'deepgram',
    'mistral', 'together', 'couchbase', 'fivetran', 'starburst', 'astronomer',
    'prefect', 'airbyte', 'postman', 'browserstack', 'inmobi', 'clevertap'
]

gh_results = []
for b in gh_boards:
    if b in history_companies or b in ['amazon', 'sarvam ai', 'mongodb']:
        continue
    url = f'https://boards-api.greenhouse.io/v1/boards/{b}/jobs?content=true'
    try:
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, context=ctx, timeout=4) as resp:
            data = json.loads(resp.read().decode('utf-8'))
            jobs = data.get('jobs', [])
            for j in jobs:
                loc = (j.get('location', {}).get('name') or '').lower()
                title = j.get('title', '')
                tl = title.lower()

                if not any(k in loc for k in ['india', 'bengaluru', 'bangalore', 'noida', 'gurgaon', 'gurugram', 'hyderabad', 'pune', 'mumbai', 'remote']):
                    continue

                if any(w in tl for w in ['senior', 'sr.', 'principal', 'staff', 'lead', 'manager', 'director', 'intern', 'apprentice', 'trainee', 'head']):
                    continue

                if not any(w in tl for w in ['engineer', 'developer', 'sde', 'qa', 'analyst', 'systems', 'cloud', 'software', 'technology']):
                    continue

                gh_results.append({
                    'company': b.capitalize(),
                    'title': title,
                    'location': j.get('location', {}).get('name'),
                    'url': j.get('absolute_url'),
                    'desc_html': j.get('content', ''),
                    'source': 'Greenhouse'
                })
    except:
        pass

print(f"Total Greenhouse candidates harvested for Wave 14: {len(gh_results)}")
with open('data/wave14_gh_raw.json', 'w') as f:
    json.dump(gh_results, f, indent=2)
