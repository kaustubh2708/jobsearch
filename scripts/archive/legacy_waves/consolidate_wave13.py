import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

wb = openpyxl.load_workbook('data/Job Search.xlsx')

if 'Wave 13' not in wb.sheetnames:
    print("Wave 13 not found in workbook, already consolidated.")
    exit(0)

ws_prev = wb['Prev Waves']
ws_w13 = wb['Wave 13']

start_row = ws_prev.max_row + 1
start_num = ws_prev.max_row # row 253 has # 252, so row 254 gets # 253

font_regular = Font(name='Calibri', size=10, bold=False, color='000000')
font_bold = Font(name='Calibri', size=10, bold=True, color='000000')
font_link = Font(name='Calibri', size=10, color='000563C1', underline='single')

align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)
align_center = Alignment(horizontal='center', vertical='center')

border_thin = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

fill_even = PatternFill(start_color='F9FAFB', end_color='F9FAFB', fill_type='solid')
fill_odd = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')

appended_count = 0
for r in range(2, ws_w13.max_row + 1):
    comp = ws_w13.cell(r, 2).value
    title = ws_w13.cell(r, 3).value
    loc = ws_w13.cell(r, 4).value
    domain = ws_w13.cell(r, 5).value
    exp = ws_w13.cell(r, 6).value
    tech = ws_w13.cell(r, 7).value
    score = ws_w13.cell(r, 9).value
    status_str = ws_w13.cell(r, 11).value
    cell_u = ws_w13.cell(r, 12)
    u = cell_u.hyperlink.target if cell_u.hyperlink else cell_u.value

    curr_row = start_row + appended_count
    curr_num = start_num + appended_count
    fill = fill_even if curr_row % 2 == 0 else fill_odd

    ws_prev.row_dimensions[curr_row].height = 24

    c1 = ws_prev.cell(curr_row, 1, curr_num)
    c1.font = font_bold; c1.alignment = align_center; c1.border = border_thin; c1.fill = fill

    c2 = ws_prev.cell(curr_row, 2, 'Wave 13')
    c2.font = font_bold; c2.alignment = align_center; c2.border = border_thin; c2.fill = fill

    c3 = ws_prev.cell(curr_row, 3, comp)
    c3.font = font_bold; c3.alignment = align_left; c3.border = border_thin; c3.fill = fill

    c4 = ws_prev.cell(curr_row, 4, title)
    c4.font = font_regular; c4.alignment = align_left; c4.border = border_thin; c4.fill = fill

    c5 = ws_prev.cell(curr_row, 5, loc)
    c5.font = font_regular; c5.alignment = align_left; c5.border = border_thin; c5.fill = fill

    c6 = ws_prev.cell(curr_row, 6, domain)
    c6.font = font_regular; c6.alignment = align_left; c6.border = border_thin; c6.fill = fill

    c7 = ws_prev.cell(curr_row, 7, exp)
    c7.font = font_regular; c7.alignment = align_left; c7.border = border_thin; c7.fill = fill

    c8 = ws_prev.cell(curr_row, 8, tech)
    c8.font = font_regular; c8.alignment = align_left; c8.border = border_thin; c8.fill = fill

    c9_val = f"{status_str} ({score}/100)" if status_str and 'CLOSED' in status_str else f"Strong Match ({score}/100)"
    c9 = ws_prev.cell(curr_row, 9, c9_val)
    c9.font = font_regular; c9.alignment = align_center; c9.border = border_thin; c9.fill = fill

    c10 = ws_prev.cell(curr_row, 10, 'Direct Link')
    c10.hyperlink = u
    c10.font = font_link; c10.alignment = align_left; c10.border = border_thin; c10.fill = fill

    appended_count += 1

del wb['Wave 13']
wb.save('data/Job Search.xlsx')
print(f"Successfully appended {appended_count} roles from Wave 13 into Prev Waves. New Prev Waves max_row: {ws_prev.max_row}. Deleted Wave 13 sheet.")
print("Sheets remaining:", wb.sheetnames)

