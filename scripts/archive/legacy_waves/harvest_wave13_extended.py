import urllib.request, json, ssl, re, time
from bs4 import BeautifulSoup
from selenium import webdriver
from selenium.webdriver.chrome.options import Options
import openpyxl

ctx = ssl.create_default_context()
ctx.check_hostname = False
ctx.verify_mode = ssl.CERT_NONE
headers = {
    'User-Agent': 'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36'
}

# 1. Load seen URLs and job IDs
all_seen_urls = set()
all_seen_req_ids = set()

with open('data/vrinda_all_seen_urls.json') as f:
    dedup = json.load(f)
    for u in dedup.get('seen_urls', []):
        all_seen_urls.add(u.strip().lower().rstrip('/'))
    for jid in dedup.get('seen_ids', []):
        all_seen_req_ids.add(str(jid).strip().lower())

wb_check = openpyxl.load_workbook('data/Companies(1).xlsx', data_only=False)
for sname in ['past wave', 'Wave 12', 'Wave 13']:
    ws = wb_check[sname]
    url_col = 13 if sname == 'past wave' else 11
    for r in range(2, ws.max_row + 1):
        c = ws.cell(row=r, column=url_col)
        val = c.hyperlink.target if c.hyperlink else c.value
        if val and str(val).strip().startswith('http'):
            all_seen_urls.add(str(val).strip().lower().rstrip('/'))

paused = {'amazon', 'sarvam', 'sarvam ai', 'mongodb', 'microsoft'}

print(f"Tracking {len(all_seen_urls)} seen URLs, {len(all_seen_req_ids)} seen IDs.")

# Targeted queries for Vrinda Gupta (~3 YOE, C#/.NET, Azure, Node.js/TS, Java/Distributed Systems)
queries = [
    # .NET / C# in NCR
    ('.NET%20developer%20OR%20C%23', 'Gurugram', 0),
    ('.NET%20developer%20OR%20C%23', 'Gurugram', 25),
    ('.NET%20developer%20OR%20C%23', 'Noida', 0),
    ('.NET%20developer%20OR%20C%23', 'Noida', 25),
    ('.NET%20Core%20backend', 'Gurugram', 0),
    ('.NET%20Core%20backend', 'Noida', 0),
    ('C%23%20Software%20Engineer', 'Gurugram', 0),
    ('C%23%20Software%20Engineer', 'Noida', 0),
    ('C%23%20Software%20Engineer', 'Bengaluru', 0),
    
    # SDE 2 / Software Engineer 2 / Backend in NCR & Pan-India
    ('Software%20Engineer%20II%20backend', 'Gurugram', 0),
    ('Software%20Engineer%20II%20backend', 'Noida', 0),
    ('Software%20Engineer%20II%20backend', 'Bengaluru', 0),
    ('SDE%202%20backend', 'Gurugram', 0),
    ('SDE%202%20backend', 'Noida', 0),
    ('SDE%202%20backend', 'Bengaluru', 0),
    ('SDE%20II', 'Gurugram', 0),
    ('SDE%20II', 'Noida', 0),

    # Node.js / TypeScript Backend
    ('Node.js%20backend%20engineer', 'Gurugram', 0),
    ('Node.js%20backend%20engineer', 'Noida', 0),
    ('TypeScript%20backend%20engineer', 'Gurugram', 0),
    ('TypeScript%20backend%20engineer', 'Bengaluru', 0),

    # Cloud & Distributed Systems
    ('Azure%20backend%20engineer', 'Gurugram', 0),
    ('Azure%20backend%20engineer', 'Noida', 0),
    ('Azure%20backend%20engineer', 'Bengaluru', 0),
    ('Distributed%20Systems%20engineer', 'Gurugram', 0),
    ('Distributed%20Systems%20engineer', 'Bengaluru', 0),
    ('Microservices%20backend%20engineer', 'Gurugram', 0),
    ('Microservices%20backend%20engineer', 'Noida', 0),
    
    # AI Agents / LLM Systems
    ('AI%20Agent%20engineer', 'India', 0),
    ('Agentic%20engineer', 'India', 0),
]

disqualifiers = [
    'field engineer', 'solutions architect', 'product support', 'gtm engineer',
    'customer reliability', 'support engineer', 'sales engineer', 'network engineer',
    'business analyst', 'marketing', 'trainee', 'türkiye', 'intern', 'manager', 'lead',
    'principal', 'director', 'vp', 'head of', 'cursor', 'staff software engineer',
    'lead engineer', 'associate solutions engineer'
]

raw_leads = []
seen_current = set()

print("Harvesting candidate listings...")
for q, loc, start in queries:
    url = f"https://www.linkedin.com/jobs-guest/jobs/api/seeMoreJobPostings/search?keywords={q}&location={loc}&start={start}"
    try:
        req = urllib.request.Request(url, headers=headers)
        with urllib.request.urlopen(req, timeout=6, context=ctx) as r:
            html = r.read().decode('utf-8')
            soup = BeautifulSoup(html, 'html.parser')
            for card in soup.find_all('li'):
                title_el = card.find('h3', class_='base-search-card__title')
                comp_el = card.find('h4', class_='base-search-card__subtitle')
                loc_el = card.find('span', class_='job-search-card__location')
                link_el = card.find('a', class_='base-card__full-link')
                if title_el and comp_el and link_el:
                    title = title_el.get_text(strip=True)
                    comp = comp_el.get_text(strip=True)
                    location = loc_el.get_text(strip=True) if loc_el else loc
                    href = link_el.get('href', '').split('?')[0]

                    # Filter paused
                    if comp.lower() in paused: continue

                    t_lower = title.lower()
                    if any(dk in t_lower for dk in disqualifiers): continue
                    if 'python developer' in t_lower or 'sr python' in t_lower or 'python engineer' in t_lower: continue
                    if not any(ek in t_lower for ek in ['software engineer', 'backend', 'developer', 'sde', 'swe', 'full stack', 'cloud engineer', 'devops', 'systems engineer']):
                        continue

                    clean_href = href.lower().rstrip('/')
                    if clean_href in all_seen_urls or clean_href in seen_current: continue

                    m_id = re.search(r'(\d{8,12})', href)
                    jid = m_id.group(1) if m_id else None
                    if jid and jid in all_seen_req_ids: continue

                    seen_current.add(clean_href)
                    raw_leads.append({
                        'company': comp,
                        'title': title,
                        'location': location,
                        'url': href,
                        'job_id': jid
                    })
        time.sleep(0.3)
    except Exception as e:
        pass

print(f"Collected {len(raw_leads)} unique candidate listings.")
with open('data/raw_wave13_extended_leads.json', 'w') as f:
    json.dump(raw_leads, f, indent=2)

