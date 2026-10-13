import openpyxl, time, json, urllib.request, ssl
from bs4 import BeautifulSoup
from selenium import webdriver
from selenium.webdriver.chrome.options import Options

print("Starting robust past waves link verification audit...")

wb = openpyxl.load_workbook('data/Companies(1).xlsx', data_only=True)
ws_past = wb['past wave']

# Collect all links from 'past wave'
rows_to_check = []
for r in range(2, ws_past.max_row + 1):
    comp = ws_past.cell(r, 3).value
    title = ws_past.cell(r, 4).value
    url_cell = ws_past.cell(r, 12)
    url = url_cell.hyperlink.target if url_cell.hyperlink else url_cell.value
    wave = ws_past.cell(r, 2).value
    
    if url and str(url).startswith('http'):
        rows_to_check.append({
            'row': r,
            'wave': str(wave),
            'company': str(comp),
            'title': str(title),
            'url': str(url).strip()
        })

print(f"Total links to verify in past wave: {len(rows_to_check)}")

options = Options()
options.add_argument('--headless=new')
options.add_argument('--no-sandbox')
options.add_argument('--disable-dev-shm-usage')
options.add_argument('--user-agent=Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36')

driver = webdriver.Chrome(options=options)

ctx = ssl.create_default_context()
ctx.check_hostname = False
ctx.verify_mode = ssl.CERT_NONE
headers = {'User-Agent': 'Mozilla/5.0'}

closed_indicators = [
    'page not found', 'no longer available', 'job closed',
    'no longer accepting applications', 'this job is no longer active',
    'this posting has expired', 'the job you are looking for is no longer open',
    '404 not found', 'job not found', 'no longer open'
]

results = {
    'total_checked': 0,
    'still_open': 0,
    'closed_or_expired': 0,
    'error_or_blocked': 0,
    'details': []
}

def restart_driver():
    global driver
    try:
        driver.quit()
    except Exception:
        pass
    driver = webdriver.Chrome(options=options)

try:
    for idx, item in enumerate(rows_to_check, 1):
        u = item['url']
        results['total_checked'] += 1
        status = 'unknown'
        reason = ''
        
        # 1. Fast HTTP probe
        try:
            req = urllib.request.Request(u, headers=headers)
            with urllib.request.urlopen(req, timeout=4, context=ctx) as resp:
                code = resp.getcode()
                if code == 404 or code == 410:
                    status = 'closed_or_expired'
                    reason = f'HTTP {code}'
        except urllib.error.HTTPError as e:
            if e.code in [404, 410]:
                status = 'closed_or_expired'
                reason = f'HTTP {e.code}'
        except Exception:
            pass

        # 2. Browser check if still unknown
        if status == 'unknown':
            try:
                driver.get(u)
                time.sleep(1.2)
                b_text = driver.find_element('tag name', 'body').text.lower()
                
                is_closed = any(ci in b_text for ci in closed_indicators)
                if is_closed or len(b_text) < 200:
                    status = 'closed_or_expired'
                    reason = 'Rendered text contains closed/expired indicator or empty shell'
                else:
                    status = 'still_open'
                    reason = 'Active requisition page rendered'
            except Exception as e:
                err_msg = str(e)
                if 'invalid session' in err_msg or 'disconnected' in err_msg:
                    print("Restarting webdriver session...")
                    restart_driver()
                    try:
                        driver.get(u)
                        time.sleep(1.2)
                        b_text = driver.find_element('tag name', 'body').text.lower()
                        is_closed = any(ci in b_text for ci in closed_indicators)
                        if is_closed or len(b_text) < 200:
                            status = 'closed_or_expired'
                            reason = 'Rendered text contains closed indicator'
                        else:
                            status = 'still_open'
                            reason = 'Active requisition page rendered'
                    except Exception as e2:
                        status = 'error_or_blocked'
                        reason = str(e2)[:80]
                else:
                    status = 'error_or_blocked'
                    reason = err_msg[:80]

        if status == 'still_open':
            results['still_open'] += 1
        elif status == 'closed_or_expired':
            results['closed_or_expired'] += 1
        else:
            results['error_or_blocked'] += 1

        results['details'].append({
            'row': item['row'],
            'wave': item['wave'],
            'company': item['company'],
            'title': item['title'],
            'url': item['url'],
            'status': status,
            'reason': reason
        })

        if idx % 25 == 0 or idx == len(rows_to_check):
            print(f"Audit Progress: [{idx}/{len(rows_to_check)}] Checked. Open: {results['still_open']}, Closed: {results['closed_or_expired']}, Error: {results['error_or_blocked']}")
            with open('data/past_waves_link_audit_report.json', 'w') as f:
                json.dump(results, f, indent=2)

finally:
    driver.quit()

print("\nFinal Link Audit Complete!")
print(f"Total Evaluated: {results['total_checked']}")
print(f"Still Open: {results['still_open']} ({results['still_open']/results['total_checked']*100:.1f}%)")
print(f"Closed / Expired: {results['closed_or_expired']} ({results['closed_or_expired']/results['total_checked']*100:.1f}%)")
print(f"Error / Blocked: {results['error_or_blocked']}")

with open('data/past_waves_link_audit_report.json', 'w') as f:
    json.dump(results, f, indent=2)
print("Saved final report to data/past_waves_link_audit_report.json")
