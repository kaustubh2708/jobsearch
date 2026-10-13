import urllib.request, urllib.parse, ssl, json, time, re
from bs4 import BeautifulSoup
import openpyxl

ctx = ssl.create_default_context()
ctx.check_hostname = False
ctx.verify_mode = ssl.CERT_NONE

headers = {
    'User-Agent': 'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36',
    'Accept-Language': 'en-US,en;q=0.9',
}

# 1. Compile existing URLs to avoid duplicate links
wb = openpyxl.load_workbook('data/Job Search.xlsx')
history_urls = set()

for sheetname in ['Prev Waves', 'Wave 13', 'Wave 14']:
    ws = wb[sheetname]
    for r in range(2, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(row=r, column=c)
            u = cell.hyperlink.target if cell.hyperlink else (cell.value if isinstance(cell.value, str) and cell.value.startswith('http') else None)
            if u:
                history_urls.add(u.strip().split('?')[0])

print(f"Loaded {len(history_urls)} existing URLs for deduplication.")

paused_companies = {'amazon', 'microsoft', 'mongodb', 'sarvam ai'}

target_giants = [
    'Qualcomm', 'Cisco', 'Oracle', 'Adobe', 'Salesforce', 'Nvidia',
    'Intel', 'AMD', 'Atlassian', 'Uber', 'PayPal', 'Intuit',
    'SAP', 'IBM', 'Siemens', 'Walmart Global Tech', 'Morgan Stanley',
    'Goldman Sachs', 'JPMorgan Chase', 'Google', 'Apple', 'Mastercard',
    'Visa', 'ServiceNow', 'Snowflake', 'Databricks', 'Stripe', 'Autodesk',
    'Synopsys', 'Cadence', 'MathWorks', 'Akamai', 'Broadcom', 'Micron Technology',
    'Sony', 'Philips', 'Honeywell', 'Barclays', 'HSBC', 'UBS',
    'Deutsche Bank', 'Citi', 'American Express', 'Fidelity Investments',
    'Societe Generale', 'Wells Fargo', 'BNY Mellon', 'BlackRock', 'NatWest Group'
]

discovered = []
seen_urls = set(history_urls)

# Search targeted queries for each giant
search_queries = []
for g in target_giants:
    search_queries.append((g, f"{g} Software Engineer India"))
    search_queries.append((g, f"{g} Associate Engineer India"))
    search_queries.append((g, f"{g} Graduate Engineer India"))

# General fresher queries for top MNCs
general_queries = [
    'Associate Software Engineer India MNC',
    'Junior Software Engineer India Product',
    'Graduate Software Engineer India Bangalore',
    'SDE 1 India Fintech',
    'Associate Data Engineer India',
    'Associate Cloud Engineer India'
]

print(f"Harvesting LinkedIn for {len(target_giants)} tech giants...")
for g, q in search_queries:
    kw = urllib.parse.quote(q)
    url = f"https://in.linkedin.com/jobs-guest/jobs/api/seeMoreJobPostings/search?keywords={kw}&location=India&geoId=102713980&start=0"
    try:
        req = urllib.request.Request(url, headers=headers)
        with urllib.request.urlopen(req, context=ctx, timeout=6) as resp:
            soup = BeautifulSoup(resp.read().decode('utf-8', errors='ignore'), 'html.parser')
            cards = soup.find_all('div', class_='base-card') or soup.find_all('li')
            for c in cards:
                t_el = c.find('h3', class_='base-search-card__title') or c.find('h3')
                comp_el = c.find('h4', class_='base-search-card__subtitle') or c.find('h4')
                loc_el = c.find('span', class_='job-search-card__location') or c.find('span')
                link_el = c.find('a', class_='base-card__full-link') or c.find('a')

                if t_el and comp_el and link_el:
                    comp = comp_el.get_text(strip=True)
                    title = t_el.get_text(strip=True)
                    u = link_el.get('href', '').split('?')[0].strip()
                    loc = loc_el.get_text(strip=True) if loc_el else 'India'

                    if u in seen_urls:
                        continue
                    if any(p in comp.lower() for p in paused_companies):
                        continue

                    seen_urls.add(u)
                    discovered.append({
                        'company': comp,
                        'title': title,
                        'location': loc,
                        'url': u,
                        'source': 'LinkedIn',
                        'target_giant': g
                    })
    except Exception as e:
        pass
    time.sleep(0.5)

print(f"Harvested {len(discovered)} initial LinkedIn candidate cards.")

with open('data/wave14_extended_raw_linkedin.json', 'w') as f:
    json.dump(discovered, f, indent=2)

