import urllib.request, urllib.parse, ssl, json, time, re
from bs4 import BeautifulSoup
import openpyxl

ctx = ssl.create_default_context()
ctx.check_hostname = False
ctx.verify_mode = ssl.CERT_NONE

headers = {
    'User-Agent': 'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36',
    'Accept-Language': 'en-US,en;q=0.9',
}

wb = openpyxl.load_workbook('data/Job Search.xlsx')
history_urls = set()
for sname in ['Prev Waves', 'Wave 14']:
    ws = wb[sname]
    for r in range(2, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(row=r, column=c)
            u = cell.hyperlink.target if cell.hyperlink else (cell.value if isinstance(cell.value, str) and cell.value.startswith('http') else None)
            if u:
                history_urls.add(u.strip().split('?')[0])

with open('data/wave15_zero_tolerance_approved.json') as f:
    existing = json.load(f)

for r in existing:
    history_urls.add(r['url'].strip().split('?')[0])

paused = {'amazon', 'microsoft', 'mongodb', 'sarvam ai'}

queries = [
    'Associate Software Engineer India Fresher',
    'Junior Software Engineer India 0-1',
    'Graduate Software Engineer Bangalore 2026',
    'Junior Python Developer India Fresher',
    'Junior Java Developer India 0-1',
    'Junior Backend Developer India 0-1',
    'Associate Data Engineer India Fresher',
    'Associate Cloud Engineer India Fresher',
    'Junior Full Stack Developer India 0-1',
    'Software Engineer Campus Hire India',
    'Associate Technical Consultant India Fresher',
    'Software Engineer Early Career India Bangalore',
    'Graduate Engineer India Pune',
    'Associate QA Engineer India Fresher',
    'Junior Systems Engineer India 0-1',
    'Cisco Graduate Software Engineer India',
    'Oracle Associate Software Engineer India',
    'Salesforce Futureforce Associate India',
    'Adobe College Graduate India',
    'Google University Graduate India',
    'Atlassian Graduate Engineer India'
]

discovered = []
seen_urls = set(history_urls)

for q in queries:
    for start in [0, 25]:
        kw = urllib.parse.quote(q)
        url = f"https://in.linkedin.com/jobs-guest/jobs/api/seeMoreJobPostings/search?keywords={kw}&location=India&geoId=102713980&start={start}"
        try:
            req = urllib.request.Request(url, headers=headers)
            with urllib.request.urlopen(req, context=ctx, timeout=6) as resp:
                soup = BeautifulSoup(resp.read().decode('utf-8', errors='ignore'), 'html.parser')
                cards = soup.find_all('div', class_='base-card') or soup.find_all('li')
                for c in cards:
                    t_el = c.find('h3', class_='base-search-card__title') or c.find('h3')
                    comp_el = c.find('h4', class_='base-search-card__subtitle') or c.find('h4')
                    loc_el = c.find('span', class_='job-search-card__location') or c.find('span')
                    link_el = c.find('a', class_='base-card__full-link') or c.find('a')

                    if t_el and comp_el and link_el:
                        comp = comp_el.get_text(strip=True)
                        title = t_el.get_text(strip=True)
                        u = link_el.get('href', '').split('?')[0].strip()
                        loc = loc_el.get_text(strip=True) if loc_el else 'India'

                        if u in seen_urls: continue
                        if any(p in comp.lower() for p in paused): continue

                        seen_urls.add(u)
                        discovered.append({
                            'company': comp,
                            'title': title,
                            'location': loc,
                            'url': u
                        })
        except:
            pass
        time.sleep(0.3)

print(f"Discovered {len(discovered)} supplemental candidate cards.")
with open('data/wave15_fresher_supplement_raw.json', 'w') as f:
    json.dump(discovered, f, indent=2)

