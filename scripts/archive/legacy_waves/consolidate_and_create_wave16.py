import openpyxl, json
from copy import copy
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

with open('data/curated_wave16_final.json') as f:
    wave16_roles = json.load(f)

# Styles
white_fill = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')
verified_fill = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid') # Soft green
header_fill = PatternFill(start_color='1F497D', end_color='1F497D', fill_type='solid') # Navy
header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')

font_regular = Font(name='Calibri', size=11, bold=False, color='000000')
font_bold = Font(name='Calibri', size=11, bold=True, color='000000')
font_link = Font(name='Calibri', size=11, bold=False, color='0563C1', underline='single')

center_align = Alignment(horizontal='center', vertical='center')
left_align = Alignment(horizontal='left', vertical='center')
thin_side = Side(border_style='thin', color='D9D9D9')
thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

files = ['data/Companies(1).xlsx', 'data/companies.xlsx']

for target_file in files:
    print(f"\nProcessing {target_file}...")
    wb = openpyxl.load_workbook(target_file)
    
    p_ws = wb['past wave']
    w14_ws = wb['Wave 14']

    # 1. Consolidate Wave 14 into past wave
    current_max_p = p_ws.max_row
    last_s_no = p_ws.cell(current_max_p, 1).value or (current_max_p - 1)
    
    print(f"  Consolidating Wave 14 ({w14_ws.max_row - 1} rows) into past wave (starting at row {current_max_p + 1})...")
    
    for r in range(2, w14_ws.max_row + 1):
        target_row = current_max_p + (r - 1)
        s_no = last_s_no + (r - 1)

        # Wave 14 columns:
        # 1: #
        # 2: Company
        # 3: Role Title
        # 4: Location (live)
        # 5: Experience (per JD)
        # 6: Key Tech Stack (estimate)
        # 7: Estimated Compensation (estimate)
        # 8: Match Score
        # 9: Priority / Match Tier
        # 10: Link Verification Status
        # 11: Direct Application URL
        # 12: Notes

        comp = w14_ws.cell(r, 2).value
        title = w14_ws.cell(r, 3).value
        loc = w14_ws.cell(r, 4).value
        exp = w14_ws.cell(r, 5).value
        tech = w14_ws.cell(r, 6).value
        comp_est = w14_ws.cell(r, 7).value
        score = w14_ws.cell(r, 8).value
        tier = w14_ws.cell(r, 9).value
        link_status = w14_ws.cell(r, 10).value
        url_cell = w14_ws.cell(r, 11)
        url_target = url_cell.hyperlink.target if url_cell.hyperlink else url_cell.value
        notes = w14_ws.cell(r, 12).value

        # Populate past wave (14 cols):
        c1 = p_ws.cell(target_row, 1, s_no)
        c1.alignment = center_align; c1.font = font_regular; c1.border = thin_border
        
        c2 = p_ws.cell(target_row, 2, 'Wave 14')
        c2.alignment = center_align; c2.font = font_regular; c2.border = thin_border

        c3 = p_ws.cell(target_row, 3, comp)
        c3.alignment = left_align; c3.font = font_regular; c3.border = thin_border

        c4 = p_ws.cell(target_row, 4, title)
        c4.alignment = left_align; c4.font = font_regular; c4.border = thin_border

        c5 = p_ws.cell(target_row, 5, loc)
        c5.alignment = left_align; c5.font = font_regular; c5.border = thin_border

        c6 = p_ws.cell(target_row, 6, exp)
        c6.alignment = left_align; c6.font = font_regular; c6.border = thin_border

        c7 = p_ws.cell(target_row, 7, tech)
        c7.alignment = left_align; c7.font = font_regular; c7.border = thin_border

        c8 = p_ws.cell(target_row, 8, comp_est)
        c8.alignment = left_align; c8.font = font_regular; c8.border = thin_border

        c9 = p_ws.cell(target_row, 9, score)
        c9.alignment = center_align; c9.font = font_regular; c9.border = thin_border

        c10 = p_ws.cell(target_row, 10, tier)
        c10.alignment = left_align; c10.font = font_regular; c10.border = thin_border

        c11 = p_ws.cell(target_row, 11, link_status)
        c11.alignment = left_align; c11.font = font_regular; c11.border = thin_border

        c12 = p_ws.cell(target_row, 12, 'Open')
        c12.alignment = center_align; c12.font = Font(name='Calibri', size=11, bold=True, color='276A3C'); c12.fill = verified_fill; c12.border = thin_border

        c13 = p_ws.cell(target_row, 13, 'Exact Requisition Link ↗')
        if url_target: c13.hyperlink = url_target
        c13.alignment = center_align; c13.font = font_link; c13.border = thin_border

        c14 = p_ws.cell(target_row, 14, notes)
        c14.alignment = left_align; c14.font = font_regular; c14.border = thin_border

    # Remove Wave 14 sheet
    wb.remove(w14_ws)
    print("  Removed Wave 14 standalone sheet.")

    # Create Wave 16 sheet
    w16_ws = wb.create_sheet(title='Wave 16')
    headers_w16 = [
        '#', 'Company', 'Role Title', 'Location (live)', 'Experience (per JD)',
        'Key Tech Stack (estimate)', 'Estimated Compensation (estimate)', 'Match Score',
        'Priority / Match Tier', 'Link Verification Status', 'Direct Application URL', 'Notes'
    ]
    
    col_widths = {
        'A': 6.0, 'B': 22.0, 'C': 34.0, 'D': 24.0, 'E': 20.0,
        'F': 32.0, 'G': 28.0, 'H': 14.0, 'I': 28.0, 'J': 26.0,
        'K': 24.0, 'L': 45.0
    }

    for col_letter, width in col_widths.items():
        w16_ws.column_dimensions[col_letter].width = width

    for c_idx, h in enumerate(headers_w16, 1):
        cell = w16_ws.cell(1, c_idx, h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border

    print(f"  Populating Wave 16 sheet with {len(wave16_roles)} verified roles...")
    for i, role in enumerate(wave16_roles, 1):
        r = 1 + i
        
        c1 = w16_ws.cell(r, 1, i)
        c1.alignment = center_align; c1.font = font_regular; c1.fill = white_fill; c1.border = thin_border

        c2 = w16_ws.cell(r, 2, role['company'])
        c2.alignment = left_align; c2.font = font_regular; c2.fill = white_fill; c2.border = thin_border

        c3 = w16_ws.cell(r, 3, role['title'])
        c3.alignment = left_align; c3.font = font_regular; c3.fill = white_fill; c3.border = thin_border

        c4 = w16_ws.cell(r, 4, role['location'])
        c4.alignment = left_align; c4.font = font_regular; c4.fill = white_fill; c4.border = thin_border

        c5 = w16_ws.cell(r, 5, role['experience'])
        c5.alignment = left_align; c5.font = font_regular; c5.fill = white_fill; c5.border = thin_border

        c6 = w16_ws.cell(r, 6, role['tech_stack'])
        c6.alignment = left_align; c6.font = font_regular; c6.fill = white_fill; c6.border = thin_border

        c7 = w16_ws.cell(r, 7, role['comp'])
        c7.alignment = left_align; c7.font = font_regular; c7.fill = white_fill; c7.border = thin_border

        c8 = w16_ws.cell(r, 8, role['score'])
        c8.alignment = center_align; c8.font = font_regular; c8.fill = white_fill; c8.border = thin_border

        c9 = w16_ws.cell(r, 9, role['tier'])
        c9.alignment = left_align; c9.font = font_regular; c9.fill = white_fill; c9.border = thin_border

        c10 = w16_ws.cell(r, 10, role['link_status'])
        c10.alignment = left_align; c10.font = font_regular; c10.fill = verified_fill; c10.border = thin_border

        c11 = w16_ws.cell(r, 11, 'Exact Requisition Link ↗')
        c11.hyperlink = role['url']
        c11.alignment = center_align; c11.font = font_link; c11.fill = white_fill; c11.border = thin_border

        c12 = w16_ws.cell(r, 12, role['notes'])
        c12.alignment = left_align; c12.font = font_regular; c12.fill = white_fill; c12.border = thin_border

    wb.save(target_file)
    print(f"  Successfully saved {target_file}. Sheets: {wb.sheetnames}.")

