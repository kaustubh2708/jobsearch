import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = openpyxl.load_workbook('data/Job Search.xlsx')
prev_ws = wb['Prev Waves']
w11_ws = wb['Wave 11']

print(f"Prev Waves initial rows: {prev_ws.max_row}")
print(f"Wave 11 rows: {w11_ws.max_row}")

next_row = prev_ws.max_row + 1

thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

for r in range(2, w11_ws.max_row + 1):
    c_comp = w11_ws.cell(r, 2).value
    if not c_comp:
        continue
    c_title = w11_ws.cell(r, 3).value
    c_loc = w11_ws.cell(r, 4).value
    c_domain = w11_ws.cell(r, 5).value
    c_exp = w11_ws.cell(r, 6).value
    c_tech = w11_ws.cell(r, 7).value
    c_score = w11_ws.cell(r, 9).value
    c_tier = w11_ws.cell(r, 10).value

    url_cell = w11_ws.cell(r, 12)
    c_url = url_cell.hyperlink.target if url_cell.hyperlink else url_cell.value

    status_score = f"{c_tier} ({c_score}/100)" if c_score else (c_tier or "Potential Match")

    # Prev Waves columns:
    # 1: #
    # 2: Wave
    # 3: Company
    # 4: Role Title
    # 5: Location
    # 6: Domain / Verdict
    # 7: Stated JD Experience Requirement
    # 8: Technical Stack & Fit / Notes
    # 9: Status / Match Score
    # 10: Application URL

    prev_ws.cell(next_row, 1, next_row - 1)
    prev_ws.cell(next_row, 2, 'Wave 11')
    prev_ws.cell(next_row, 3, c_comp)
    prev_ws.cell(next_row, 4, c_title)
    prev_ws.cell(next_row, 5, c_loc)
    prev_ws.cell(next_row, 6, c_domain)
    prev_ws.cell(next_row, 7, c_exp)
    prev_ws.cell(next_row, 8, c_tech)
    prev_ws.cell(next_row, 9, status_score)

    u_target = prev_ws.cell(next_row, 10, 'Direct Link' if c_url else 'N/A')
    if c_url:
        u_target.hyperlink = c_url
        u_target.font = Font(name='Calibri', size=10, color='0563C1', underline='single')
    else:
        u_target.font = Font(name='Calibri', size=10)

    for col_idx in range(1, 10):
        cell = prev_ws.cell(next_row, col_idx)
        cell.font = Font(name='Calibri', size=10)
        cell.border = thin_border
    prev_ws.cell(next_row, 10).border = thin_border

    next_row += 1

del wb['Wave 11']
wb.save('data/Job Search.xlsx')
print(f"Successfully consolidated Wave 11 into Prev Waves. Prev Waves max row now: {prev_ws.max_row}")
print("Active sheets:", wb.sheetnames)
