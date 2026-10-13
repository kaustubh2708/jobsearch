import json, openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

wb = openpyxl.load_workbook('data/Job Search.xlsx')

if 'Wave 15' in wb.sheetnames:
    del wb['Wave 15']

ws = wb.create_sheet('Wave 15')

with open('data/wave15_final_approved.json') as f:
    roles = json.load(f)

headers = [
    '#', 'Company', 'Role Title', 'Location', 'Domain / Track',
    'Exact Stated Experience in JD (0–2 YOE)', 'Key Tech Stack',
    'Estimated / Stated Compensation', 'Match Score', 'Priority / Match Tier',
    'Link Verification Status', 'Direct Application URL', 'Role Fit & Strategy Notes'
]

# Header styling
header_fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

border_thin = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

ws.row_dimensions[1].height = 28
for col_idx, h in enumerate(headers, 1):
    cell = ws.cell(1, col_idx, h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_align
    cell.border = border_thin

font_regular = Font(name='Calibri', size=10, bold=False, color='000000')
font_bold = Font(name='Calibri', size=10, bold=True, color='000000')
font_link = Font(name='Calibri', size=10, color='000563C1', underline='single')

align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)
align_center = Alignment(horizontal='center', vertical='center')

fill_even = PatternFill(start_color='F9FAFB', end_color='F9FAFB', fill_type='solid')
fill_odd = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')

for idx, role in enumerate(roles, 1):
    row_idx = idx + 1
    fill = fill_even if row_idx % 2 == 0 else fill_odd
    ws.row_dimensions[row_idx].height = 24

    # 1. #
    c1 = ws.cell(row_idx, 1, idx)
    c1.font = font_bold; c1.alignment = align_center; c1.border = border_thin; c1.fill = fill

    # 2. Company
    c2 = ws.cell(row_idx, 2, role['company'])
    c2.font = font_bold; c2.alignment = align_left; c2.border = border_thin; c2.fill = fill

    # 3. Role Title
    c3 = ws.cell(row_idx, 3, role['title'])
    c3.font = font_regular; c3.alignment = align_left; c3.border = border_thin; c3.fill = fill

    # 4. Location
    c4 = ws.cell(row_idx, 4, role['location'])
    c4.font = font_regular; c4.alignment = align_left; c4.border = border_thin; c4.fill = fill

    # 5. Domain / Track
    c5 = ws.cell(row_idx, 5, role['domain'])
    c5.font = font_regular; c5.alignment = align_left; c5.border = border_thin; c5.fill = fill

    # 6. Exact Stated Experience in JD (0–2 YOE)
    c6 = ws.cell(row_idx, 6, role['experience_stated'])
    c6.font = font_regular; c6.alignment = align_left; c6.border = border_thin; c6.fill = fill

    # 7. Key Tech Stack
    c7 = ws.cell(row_idx, 7, role['tech_stack'])
    c7.font = font_regular; c7.alignment = align_left; c7.border = border_thin; c7.fill = fill

    # 8. Estimated / Stated Compensation
    c8 = ws.cell(row_idx, 8, role['compensation'])
    c8.font = font_bold; c8.alignment = align_left; c8.border = border_thin; c8.fill = fill

    # 9. Match Score
    c9 = ws.cell(row_idx, 9, role['match_score'])
    c9.font = font_bold; c9.alignment = align_center; c9.border = border_thin; c9.fill = fill

    # 10. Priority / Match Tier
    c10 = ws.cell(row_idx, 10, role['priority_tier'])
    c10.font = font_regular; c10.alignment = align_center; c10.border = border_thin; c10.fill = fill

    # 11. Link Verification Status
    c11 = ws.cell(row_idx, 11, role['verification_status'])
    c11.font = font_regular; c11.alignment = align_center; c11.border = border_thin; c11.fill = fill

    # 12. Direct Application URL
    c12 = ws.cell(row_idx, 12, 'Direct Link')
    c12.hyperlink = role['url']
    c12.font = font_link; c12.alignment = align_left; c12.border = border_thin; c12.fill = fill

    # 13. Role Fit & Strategy Notes
    c13 = ws.cell(row_idx, 13, role['notes'])
    c13.font = font_regular; c13.alignment = align_left; c13.border = border_thin; c13.fill = fill

# Column widths
col_widths = {1: 6, 2: 24, 3: 36, 4: 28, 5: 30, 6: 45, 7: 35, 8: 30, 9: 14, 10: 24, 11: 24, 12: 18, 13: 45}
for col_idx, width in col_widths.items():
    ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

wb.save('data/Job Search.xlsx')
print(f"Successfully populated Wave 15 with {len(roles)} roles in data/Job Search.xlsx!")
print("Active sheets:", wb.sheetnames)

