import urllib.request, json, ssl, re, time
from bs4 import BeautifulSoup
import openpyxl

ctx = ssl.create_default_context()
ctx.check_hostname = False
ctx.verify_mode = ssl.CERT_NONE
headers = {
    'User-Agent': 'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36'
}

# 1. Load seen URLs and job IDs
all_seen_urls = set()
all_seen_req_ids = set()

with open('data/vrinda_all_seen_urls.json') as f:
    dedup = json.load(f)
    for u in dedup.get('seen_urls', []):
        all_seen_urls.add(u.strip().lower().rstrip('/'))
    for jid in dedup.get('seen_ids', []):
        all_seen_req_ids.add(str(jid).strip().lower())

wb_check = openpyxl.load_workbook('data/Companies(1).xlsx', data_only=False)
for sname in ['past wave', 'Wave 14', 'Wave 15']:
    if sname in wb_check.sheetnames:
        ws = wb_check[sname]
        url_col = 13 if sname == 'past wave' else 11
        for r in range(2, ws.max_row + 1):
            c = ws.cell(row=r, column=url_col)
            val = c.hyperlink.target if c.hyperlink else c.value
            if val and str(val).strip().startswith('http'):
                all_seen_urls.add(str(val).strip().lower().rstrip('/'))

paused = {'amazon', 'sarvam', 'sarvam ai', 'mongodb', 'microsoft'}

print(f"Tracking {len(all_seen_urls)} seen URLs, {len(all_seen_req_ids)} seen IDs.")

# Target boards for top enterprise MNCs, tier-1 giants & GCCs
gh_boards = [
    'uber', 'atlassian', 'stripe', 'databricks', 'postman', 'browserstack',
    'inmobi', 'swiggy', 'clevertap', 'hasura', 'chargebee', 'druva',
    'innovaccer', 'couchbase', 'yugabyte', 'confluent', 'docusign',
    'cloudflare', 'snowflake', 'nutanix', 'qualys', 'sailpoint', 'samsara',
    'verkada', 'affirm', 'okta', 'paloaltonetworks', 'zscaler', 'crowdstrike',
    'sentinelone', 'purestorage', 'cohesity', 'rubrik', 'synopsys', 'cadence',
    'marvell', 'arm', 'siemens', 'nxp', 'mathworks', 'akamai', 'splunk',
    'netapp', 'citrix', 'teradata', 'veritas', 'commvault', 'tanium', 'fastly'
]

ats_leads = []
print("Scanning Greenhouse boards...")
for b in gh_boards:
    if b in paused: continue
    url = f'https://boards-api.greenhouse.io/v1/boards/{b}/jobs?content=true'
    try:
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, context=ctx, timeout=4) as resp:
            data = json.loads(resp.read().decode('utf-8'))
            for j in data.get('jobs', []):
                title = j.get('title', '')
                loc = (j.get('location', {}).get('name') or '').lower()
                tl = title.lower()

                # India check
                if not any(k in loc for k in ['india', 'bengaluru', 'bangalore', 'noida', 'gurgaon', 'gurugram', 'hyderabad', 'pune', 'delhi', 'remote']):
                    continue

                # Exclude non-tech / over-senior / support
                if any(w in tl for w in ['senior', 'sr.', 'principal', 'staff', 'lead', 'manager', 'director', 'intern', 'apprentice', 'head', 'architect', 'solutions engineer', 'support']):
                    continue

                # Must be core engineering
                if not any(w in tl for w in ['engineer', 'developer', 'sde', 'backend', 'cloud', 'systems', 'platform']):
                    continue

                abs_url = j.get('absolute_url', '')
                if abs_url.lower().rstrip('/') in all_seen_urls:
                    continue

                ats_leads.append({
                    'company': b.capitalize(),
                    'title': title,
                    'location': j.get('location', {}).get('name') or 'India',
                    'url': abs_url,
                    'job_id': str(j.get('id')),
                    'desc_html': j.get('content', ''),
                    'source': 'Greenhouse'
                })
    except Exception as e:
        pass

print(f"Greenhouse leads: {len(ats_leads)}")

# Ashby boards
ashby_boards = [
    'linear', 'warp', 'vanta', 'retool', 'synthesia', 'descript', 'harvey',
    'snyk', 'glean', 'cohere', 'mistral', 'together', 'groq', 'perplexity',
    'supabase', 'hex', 'deepgram', 'dust', 'atlan', 'anyscale'
]

for org in ashby_boards:
    if org in paused: continue
    url = f'https://api.ashbyhq.com/posting-api/job-board/{org}'
    try:
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, context=ctx, timeout=4) as resp:
            data = json.loads(resp.read().decode('utf-8'))
            for j in data.get('jobs', []):
                title = j.get('title', '')
                loc = (j.get('locationName') or '').lower()
                sec_locs = [str(x).lower() for x in (j.get('secondaryLocations') or [])]
                all_locs = loc + ' ' + ' '.join(sec_locs)
                if not any(k in all_locs for k in ['india', 'bengaluru', 'bangalore', 'noida', 'gurgaon', 'hyderabad', 'pune', 'delhi', 'remote']):
                    continue
                tl = title.lower()
                if any(w in tl for w in ['senior', 'sr.', 'principal', 'staff', 'lead', 'manager', 'director', 'intern', 'apprentice', 'head', 'architect', 'solutions', 'support']):
                    continue
                if not any(w in tl for w in ['engineer', 'developer', 'sde', 'backend', 'platform']):
                    continue
                job_url = j.get('jobUrl', '')
                if job_url.lower().rstrip('/') in all_seen_urls:
                    continue
                ats_leads.append({
                    'company': org.capitalize(),
                    'title': title,
                    'location': j.get('locationName') or 'India / Remote',
                    'url': job_url,
                    'job_id': str(j.get('id')),
                    'desc_html': j.get('descriptionHtml', ''),
                    'source': 'Ashby'
                })
    except Exception as e:
        pass

print(f"Total ATS leads (Greenhouse + Ashby): {len(ats_leads)}")

# LinkedIn live API searches across core tech clusters (Gurugram, Noida, Delhi NCR, Bengaluru)
queries = [
    # .NET / C#
    ('.NET%20Core%20developer', 'Gurugram', 0),
    ('.NET%20Core%20developer', 'Noida', 0),
    ('C%23%20engineer', 'Gurugram', 0),
    ('C%23%20engineer', 'Noida', 0),
    ('Azure%20Cloud%20developer', 'Gurugram', 0),
    ('Azure%20Cloud%20developer', 'Noida', 0),
    ('.NET%20backend%20engineer', 'Bengaluru', 0),
    
    # SDE II / SWE 2 / Backend
    ('Software%20Engineer%20II%20backend', 'Gurugram', 0),
    ('Software%20Engineer%20II%20backend', 'Noida', 0),
    ('Software%20Engineer%20II%20backend', 'Bengaluru', 0),
    ('SDE%202%20backend', 'Gurugram', 0),
    ('SDE%202%20backend', 'Noida', 0),
    ('SDE%202%20backend', 'Bengaluru', 0),
    ('SDE%20II', 'Gurugram', 0),
    ('SDE%20II', 'Noida', 0),

    # Node.js / TypeScript & Microservices / Distributed Systems
    ('Node.js%20backend%20engineer', 'Gurugram', 0),
    ('Node.js%20backend%20engineer', 'Noida', 0),
    ('TypeScript%20backend%20engineer', 'Gurugram', 0),
    ('TypeScript%20backend%20engineer', 'Bengaluru', 0),
    ('Distributed%20Systems%20engineer', 'Gurugram', 0),
    ('Distributed%20Systems%20engineer', 'Noida', 0),
    ('Distributed%20Systems%20engineer', 'Bengaluru', 0),
    ('AI%20Agent%20engineer', 'India', 0),
    ('Agentic%20engineer', 'India', 0),
]

disqualifiers = [
    'field engineer', 'solutions architect', 'solutions engineer', 'product support',
    'gtm engineer', 'customer reliability', 'support engineer', 'sales engineer',
    'network engineer', 'business analyst', 'marketing', 'trainee', 'türkiye',
    'intern', 'manager', 'lead', 'principal', 'director', 'vp', 'head of',
    'cursor', 'staff software engineer', 'lead engineer', 'associate solutions engineer',
    'senior staff', 'qa engineer', 'test engineer', 'automation engineer',
    'data analyst', 'sdet', 'data engineer', 'salesforce', 'sap abap', 'servicenow',
    'boomi', 'ruby on rails'
]

li_leads = []
seen_current = set()

print("Harvesting candidate listings from LinkedIn...")
for q, loc, start in queries:
    url = f"https://www.linkedin.com/jobs-guest/jobs/api/seeMoreJobPostings/search?keywords={q}&location={loc}&start={start}"
    try:
        req = urllib.request.Request(url, headers=headers)
        with urllib.request.urlopen(req, timeout=6, context=ctx) as r:
            html = r.read().decode('utf-8')
            soup = BeautifulSoup(html, 'html.parser')
            for card in soup.find_all('li'):
                title_el = card.find('h3', class_='base-search-card__title')
                comp_el = card.find('h4', class_='base-search-card__subtitle')
                loc_el = card.find('span', class_='job-search-card__location')
                link_el = card.find('a', class_='base-card__full-link')
                if title_el and comp_el and link_el:
                    title = title_el.get_text(strip=True)
                    comp = comp_el.get_text(strip=True)
                    location = loc_el.get_text(strip=True) if loc_el else loc
                    href = link_el.get('href', '').split('?')[0]

                    if comp.lower() in paused: continue

                    t_lower = title.lower()
                    if any(dk in t_lower for dk in disqualifiers): continue
                    if 'python developer' in t_lower or 'sr python' in t_lower or 'python engineer' in t_lower: continue
                    if not any(ek in t_lower for ek in ['software engineer', 'backend', 'developer', 'sde', 'swe', 'full stack', 'cloud engineer', 'systems engineer']):
                        continue

                    clean_href = href.lower().rstrip('/')
                    if clean_href in all_seen_urls or clean_href in seen_current: continue

                    m_id = re.search(r'(\d{8,12})', href)
                    jid = m_id.group(1) if m_id else None
                    if jid and jid in all_seen_req_ids: continue

                    seen_current.add(clean_href)
                    li_leads.append({
                        'company': comp,
                        'title': title,
                        'location': location,
                        'url': href,
                        'job_id': jid,
                        'source': 'LinkedIn'
                    })
        time.sleep(0.3)
    except Exception as e:
        pass

print(f"Collected {len(li_leads)} unique candidate LinkedIn listings.")

combined_candidates = ats_leads + li_leads
print(f"Total raw candidates for Wave 16: {len(combined_candidates)}")

with open('data/raw_wave16_candidates.json', 'w') as f:
    json.dump(combined_candidates, f, indent=2)

