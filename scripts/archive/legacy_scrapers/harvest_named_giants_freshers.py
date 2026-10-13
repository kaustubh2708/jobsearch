import urllib.request, urllib.parse, ssl, json, time, re
from bs4 import BeautifulSoup
import openpyxl

ctx = ssl.create_default_context()
ctx.check_hostname = False
ctx.verify_mode = ssl.CERT_NONE

headers = {
    'User-Agent': 'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36',
    'Accept-Language': 'en-US,en;q=0.9',
}

wb = openpyxl.load_workbook('data/Job Search.xlsx')
history_urls = set()
for sname in ['Prev Waves', 'Wave 14']:
    ws = wb[sname]
    for r in range(2, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(row=r, column=c)
            u = cell.hyperlink.target if cell.hyperlink else (cell.value if isinstance(cell.value, str) and cell.value.startswith('http') else None)
            if u:
                history_urls.add(u.strip().split('?')[0])

paused = {'amazon', 'microsoft', 'mongodb', 'sarvam ai'}

queries = [
    'Cisco Graduate Software Engineer India',
    'Oracle Associate Software Engineer India',
    'Salesforce Futureforce Associate MTS India',
    'Goldman Sachs Analyst Software Bengaluru',
    'Morgan Stanley Technology Analyst Bengaluru',
    'JPMorgan Chase Software Engineer Program India',
    'Walmart Global Tech Associate Software Engineer India',
    'Adobe Early Career Software Engineer India',
    'Atlassian Graduate Software Engineer India',
    'Intuit Software Engineer 1 India',
    'PayPal Software Engineer 1 India',
    'SAP Associate Developer India',
    'ServiceNow Associate Software Engineer India',
    'Snowflake Software Engineer University India',
    'Target Engineer Target India Campus',
    'Arm Graduate Engineer India',
    'Barclays Graduate Developer India',
    'Deutsche Bank Graduate Analyst Software Pune',
    'Mastercard Associate Software Engineer India',
    'Visa Associate Software Engineer India',
    'Uber Software Engineer University India',
    'Apple Software Engineer University India'
]

discovered = []
seen_urls = set(history_urls)

for q in queries:
    kw = urllib.parse.quote(q)
    url = f"https://in.linkedin.com/jobs-guest/jobs/api/seeMoreJobPostings/search?keywords={kw}&location=India&geoId=102713980&start=0"
    try:
        req = urllib.request.Request(url, headers=headers)
        with urllib.request.urlopen(req, context=ctx, timeout=6) as resp:
            soup = BeautifulSoup(resp.read().decode('utf-8', errors='ignore'), 'html.parser')
            cards = soup.find_all('div', class_='base-card') or soup.find_all('li')
            for c in cards:
                t_el = c.find('h3', class_='base-search-card__title') or c.find('h3')
                comp_el = c.find('h4', class_='base-search-card__subtitle') or c.find('h4')
                loc_el = c.find('span', class_='job-search-card__location') or c.find('span')
                link_el = c.find('a', class_='base-card__full-link') or c.find('a')

                if t_el and comp_el and link_el:
                    comp = comp_el.get_text(strip=True)
                    title = t_el.get_text(strip=True)
                    u = link_el.get('href', '').split('?')[0].strip()
                    loc = loc_el.get_text(strip=True) if loc_el else 'India'

                    if u in seen_urls: continue
                    if any(p in comp.lower() for p in paused): continue

                    seen_urls.add(u)
                    discovered.append({
                        'company': comp,
                        'title': title,
                        'location': loc,
                        'url': u,
                        'query': q
                    })
    except:
        pass
    time.sleep(0.4)

print(f"Harvested {len(discovered)} named giant campus cards.")
with open('data/wave15_named_giants_campus.json', 'w') as f:
    json.dump(discovered, f, indent=2)

