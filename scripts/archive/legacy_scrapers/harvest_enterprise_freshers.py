import urllib.request, urllib.parse, ssl, json, time, re
from bs4 import BeautifulSoup
import openpyxl

ctx = ssl.create_default_context()
ctx.check_hostname = False
ctx.verify_mode = ssl.CERT_NONE

headers = {
    'User-Agent': 'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36',
    'Accept-Language': 'en-US,en;q=0.9',
}

wb = openpyxl.load_workbook('data/Job Search.xlsx')
history_urls = set()
for sname in ['Prev Waves', 'Wave 14']:
    ws = wb[sname]
    for r in range(2, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(row=r, column=c)
            u = cell.hyperlink.target if cell.hyperlink else (cell.value if isinstance(cell.value, str) and cell.value.startswith('http') else None)
            if u:
                history_urls.add(u.strip().split('?')[0])

with open('data/wave15_final_pure_freshers_30.json') as f:
    existing = json.load(f)

for r in existing:
    history_urls.add(r['url'].strip().split('?')[0])

paused = {'amazon', 'microsoft', 'mongodb', 'sarvam ai'}

queries = [
    'Dell Associate Software Engineer India',
    'VMware Associate Software Engineer India',
    'Broadridge Associate Software Engineer India',
    'EPAM Junior Software Engineer India',
    'Persistent Systems Associate Software Engineer India',
    'Amdocs Associate Software Engineer India',
    'LTIMindtree Associate Software Engineer India',
    'Virtusa Associate Software Engineer India',
    'Hexaware Associate Software Engineer India',
    'KPIT Associate Software Engineer India',
    'Zensar Associate Software Engineer India',
    'Cisco Associate Engineer India Fresher',
    'Oracle Associate Developer India Fresher',
    'SAP Associate Developer India Fresher',
    'IBM Associate Developer India Fresher',
    'Accenture Associate Software Engineer India',
    'Cognizant Programmer Analyst India',
    'Capgemini Associate Software Engineer India',
    'Infosys Specialist Programmer India Fresher',
    'TCS Digital Software Engineer India Fresher',
    'Honeywell Associate Software Engineer India',
    'Philips Associate Software Engineer India',
    'Sony Associate Software Engineer India',
    'Qualcomm Associate Software Engineer India Fresher',
    'Nvidia Associate Software Engineer India Fresher'
]

discovered = []
seen_urls = set(history_urls)

for q in queries:
    for start in [0, 25]:
        kw = urllib.parse.quote(q)
        url = f"https://in.linkedin.com/jobs-guest/jobs/api/seeMoreJobPostings/search?keywords={kw}&location=India&geoId=102713980&start={start}"
        try:
            req = urllib.request.Request(url, headers=headers)
            with urllib.request.urlopen(req, context=ctx, timeout=6) as resp:
                soup = BeautifulSoup(resp.read().decode('utf-8', errors='ignore'), 'html.parser')
                cards = soup.find_all('div', class_='base-card') or soup.find_all('li')
                for c in cards:
                    t_el = c.find('h3', class_='base-search-card__title') or c.find('h3')
                    comp_el = c.find('h4', class_='base-search-card__subtitle') or c.find('h4')
                    loc_el = c.find('span', class_='job-search-card__location') or c.find('span')
                    link_el = c.find('a', class_='base-card__full-link') or c.find('a')

                    if t_el and comp_el and link_el:
                        comp = comp_el.get_text(strip=True)
                        title = t_el.get_text(strip=True)
                        u = link_el.get('href', '').split('?')[0].strip()
                        loc = loc_el.get_text(strip=True) if loc_el else 'India'

                        if u in seen_urls: continue
                        if any(p in comp.lower() for p in paused): continue

                        seen_urls.add(u)
                        discovered.append({
                            'company': comp,
                            'title': title,
                            'location': loc,
                            'url': u
                        })
        except:
            pass
        time.sleep(0.3)

print(f"Discovered {len(discovered)} enterprise candidate cards.")
with open('data/wave15_enterprise_candidates_raw.json', 'w') as f:
    json.dump(discovered, f, indent=2)

