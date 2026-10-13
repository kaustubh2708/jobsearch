import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

wb = openpyxl.load_workbook('data/Job Search.xlsx')
ws = wb['Wave 14']

r = 27
ws.cell(row=r, column=2, value='Qualcomm')
ws.cell(row=r, column=3, value='2027 Campus Hire_Associate Engineer_SW')
ws.cell(row=r, column=4, value='Chennai, Tamil Nadu, India')
ws.cell(row=r, column=5, value='Full Stack & Systems Software')
ws.cell(row=r, column=6, value="Bachelor's / Master's degree in Computer Science, Information Systems, ECE; 0 YOE campus fresher intake.")
ws.cell(row=r, column=7, value='C, C++, Python, Data Structures, Algorithms, Linux, Git')
ws.cell(row=r, column=8, value='INR 14 - 30 LPA (Base est. / Tier-1 MNC)')
ws.cell(row=r, column=9, value=95)
ws.cell(row=r, column=10, value='Strong Match (Tech Giant)')
ws.cell(row=r, column=11, value='Verified Live (HTTP 200)')

c12 = ws.cell(row=r, column=12, value='Direct Link')
c12.hyperlink = 'https://in.linkedin.com/jobs/view/2027-campus-hire-associate-engineer-sw-at-qualcomm-4440758187'
c12.font = Font(name='Calibri', size=10, color='000563C1', underline='single')

ws.cell(row=r, column=13, value='Full-time permanent FTE campus entry-level role at Qualcomm. 100% compliant with 0-1 YOE fresher constraints & >=14 LPA tier.')

wb.save('data/Job Search.xlsx')
print("Successfully replaced Row 27 with Qualcomm Campus Hire Software Engineer!")

