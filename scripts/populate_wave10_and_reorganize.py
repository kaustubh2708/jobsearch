import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import json
import hashlib

def run():
    print("Loading Job Search.xlsx...")
    wb = openpyxl.load_workbook('data/Job Search.xlsx')

    # 1. Fold Wave 8 into Prev Waves
    ws_prev = wb['Prev Waves']
    ws_w8 = wb['Wave 8']

    current_max_prev = ws_prev.max_row
    print(f"Prev Waves has {current_max_prev} rows currently.")
    print(f"Wave 8 has {ws_w8.max_row} rows (including header).")

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    regular_font = Font(name='Arial', size=10)
    link_font = Font(name='Arial', size=10, color='0563C1', underline='single')

    next_idx = current_max_prev
    for r in range(2, ws_w8.max_row + 1):
        comp = ws_w8.cell(r, 2).value
        title = ws_w8.cell(r, 3).value
        loc = ws_w8.cell(r, 4).value
        domain = ws_w8.cell(r, 5).value
        exp = ws_w8.cell(r, 6).value
        tech = ws_w8.cell(r, 7).value
        notes = ws_w8.cell(r, 13).value
        score = ws_w8.cell(r, 9).value
        tier = ws_w8.cell(r, 10).value
        url_cell = ws_w8.cell(r, 12)

        target_url = url_cell.hyperlink.target if url_cell.hyperlink else (url_cell.value if isinstance(url_cell.value, str) and url_cell.value.startswith('http') else None)

        target_row = ws_prev.max_row + 1
        # Prev Waves cols: ['#', 'Wave', 'Company', 'Role Title', 'Location', 'Domain / Verdict', 'Stated JD Experience Requirement', 'Technical Stack & Fit / Notes', 'Status / Match Score', 'Application URL']
        ws_prev.cell(target_row, 1, next_idx).font = regular_font
        ws_prev.cell(target_row, 2, 'Wave 8').font = regular_font
        ws_prev.cell(target_row, 3, comp).font = regular_font
        ws_prev.cell(target_row, 4, title).font = regular_font
        ws_prev.cell(target_row, 5, loc).font = regular_font
        ws_prev.cell(target_row, 6, domain).font = regular_font
        ws_prev.cell(target_row, 7, exp).font = regular_font
        ws_prev.cell(target_row, 8, f"{tech} | {notes}").font = regular_font
        ws_prev.cell(target_row, 9, f"{score} / 100 ({tier})").font = regular_font

        c_url = ws_prev.cell(target_row, 10, 'Application Link ↗' if target_url else 'Direct Link')
        if target_url:
            c_url.hyperlink = target_url
            c_url.font = link_font
        else:
            c_url.font = regular_font

        for col in range(1, 11):
            cell = ws_prev.cell(target_row, col)
            cell.border = thin_border
            cell.alignment = Alignment(vertical='center')

        next_idx += 1

    print(f"Prev Waves now has {ws_prev.max_row} rows after folding Wave 8.")

    # Remove Wave 8 sheet
    del wb['Wave 8']
    print("Deleted standalone 'Wave 8' sheet.")

    # 2. Create Wave 10 sheet
    if 'Wave 10' in wb.sheetnames:
        del wb['Wave 10']
    ws_w10 = wb.create_sheet('Wave 10')

    # Headers
    headers = [
        '#', 'Company', 'Role Title', 'Location', 'Domain / Track',
        'Exact Stated Experience in JD (0–2 YOE)', 'Key Tech Stack',
        'Estimated / Stated Compensation', 'Match Score', 'Priority / Match Tier',
        'Link Verification Status', 'Direct Application URL', 'Role Fit & Strategy Notes'
    ]

    header_font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='001F4E79', end_color='001F4E79', fill_type='solid')

    for c_idx, h in enumerate(headers, start=1):
        cell = ws_w10.cell(1, c_idx, h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = thin_border

    # Load final verified roles
    with open('data/wave10_final_verified_roles.json') as f:
        roles = json.load(f)

    print(f"Writing {len(roles)} verified roles to Wave 10...")

    zebra_fill = PatternFill(start_color='00F9FAFB', end_color='00F9FAFB', fill_type='solid')

    for r_idx, r in enumerate(roles, start=2):
        row_num = r_idx - 1
        ws_w10.cell(r_idx, 1, row_num).font = regular_font
        ws_w10.cell(r_idx, 2, r['company']).font = Font(name='Arial', size=10, bold=True)
        ws_w10.cell(r_idx, 3, r['title']).font = regular_font
        ws_w10.cell(r_idx, 4, r['location']).font = regular_font
        ws_w10.cell(r_idx, 5, r.get('domain', 'Software Engineering / Backend / Data')).font = regular_font
        ws_w10.cell(r_idx, 6, r['experience']).font = regular_font
        ws_w10.cell(r_idx, 7, r['tech_stack']).font = regular_font
        ws_w10.cell(r_idx, 8, r['compensation']).font = regular_font
        ws_w10.cell(r_idx, 9, r['score']).font = Font(name='Arial', size=10, bold=True)
        ws_w10.cell(r_idx, 10, r['tier']).font = Font(name='Arial', size=10, bold=True, color='00276A3C' if 'strong' in r['tier'] else '001F4E79')
        ws_w10.cell(r_idx, 11, 'Verified Live Exact Requisition (HTTP 200)').font = regular_font

        c_link = ws_w10.cell(r_idx, 12, 'Exact Requisition Link ↗')
        c_link.hyperlink = r['url']
        c_link.font = link_font

        ws_w10.cell(r_idx, 13, r['notes']).font = regular_font

        is_even = (r_idx % 2 == 0)
        for c in range(1, 14):
            cell = ws_w10.cell(r_idx, c)
            cell.border = thin_border
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            if is_even:
                cell.fill = zebra_fill

    # Set column widths
    col_widths = {
        'A': 6, 'B': 24, 'C': 34, 'D': 30, 'E': 28,
        'F': 36, 'G': 32, 'H': 26, 'I': 14, 'J': 20,
        'K': 28, 'L': 24, 'M': 44
    }
    for col_letter, width in col_widths.items():
        ws_w10.column_dimensions[col_letter].width = width

    # Save workbook
    wb.save('data/Job Search.xlsx')
    print("Saved data/Job Search.xlsx successfully!")
    print(f"Final sheet names in Job Search.xlsx: {wb.sheetnames}")

    # 3. Update all_existing_urls.json
    all_urls = set()
    with open('data/all_existing_urls.json') as f:
        for u in json.load(f):
            all_urls.add(u.strip().rstrip('/'))

    for r in roles:
        all_urls.add(r['url'].strip().rstrip('/'))

    with open('data/all_existing_urls.json', 'w') as f:
        json.dump(sorted(all_urls), f, indent=2)
    print(f"Updated all_existing_urls.json: {len(all_urls)} tracked URLs.")

    # 4. Check untouched master MD5
    with open('data/jobs_vrinda_discovery_wave2_verified.xlsx', 'rb') as f:
        h = hashlib.md5(f.read()).hexdigest()
    print(f"MD5 of jobs_vrinda_discovery_wave2_verified.xlsx: {h}")
    assert h == "bbf425b95af9798d23500f3799065386", f"Master hash changed! Got {h}"
    print("Master workbook verified UNTOUCHED!")

if __name__ == '__main__':
    run()
