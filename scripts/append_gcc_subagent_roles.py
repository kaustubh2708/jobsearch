import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

wb = openpyxl.load_workbook('data/Job Search.xlsx')
ws = wb['New Additions (Wave 5)']

# GCC roles from Subagent 2
gcc_roles = [
    {
        'company': 'Electronic Arts (EA)',
        'title': 'AI Full Stack Intern',
        'location': 'Hyderabad, Telangana (Hybrid)',
        'notes': 'Central Technology (CT) team at global entertainment leader EA. Requisition ID: 215939. Direct match for Python, FastAPI, LangChain, and LLM integrations. Stipend: ~₹50k-70k/mo, PPO CTC: ₹14-18 LPA.',
        'status': '88 / 100 (Strong Match)',
        'url': 'https://jobs.ea.com/en_US/careers/JobDetail/215939',
        'fill_green': True
    },
    {
        'company': 'IQVIA',
        'title': 'Software Engineer Intern (AI & Automation)',
        'location': 'Gurugram, Haryana / Bengaluru (Hybrid)',
        'notes': 'Fortune 500 Healthcare Tech Enterprise GCC in Gurugram (NCR). Requisition ID: R1567860. Exact match for Vaanya: Python (Mandatory), FastAPI, REST APIs, LLM/Agentic AI, RAG. Stipend: ~₹35k-50k/mo, PPO CTC: ₹10-13 LPA.',
        'status': '93 / 100 (Strong Match)',
        'url': 'https://jobs.iqvia.com/en/jobs/R1567860-0',
        'fill_green': True
    },
    {
        'company': 'insightsoftware',
        'title': 'Associate Software Engineer (Internship)',
        'location': 'Hyderabad / Bengaluru (Remote/Hybrid)',
        'notes': 'Global Financial Analytics leader. Workday Requisition ID: REQ001103-1. Explicit internship for fresh B.Tech graduates. Python, SQL, Cloud APIs, AI-first workflows. Stipend: ~₹35k-45k/mo, PPO CTC: ₹9-12 LPA.',
        'status': '82 / 100 (Strong Match)',
        'url': 'https://magnitudesoftware.wd1.myworkdayjobs.com/External/job/India---Hyderabad---Remote/Associate-Software-Engineer_REQ001103-1',
        'fill_green': True
    }
]

green_fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
bold_green = Font(color='006100', bold=True)

# Check if URL already in sheet
existing_in_sheet = set()
for r in range(2, ws.max_row + 1):
    val = ws.cell(r, 6).hyperlink.target if ws.cell(r, 6).hyperlink else ws.cell(r, 6).value
    if val:
        existing_in_sheet.add(str(val).lower().rstrip('/'))

added_count = 0
for gr in gcc_roles:
    u = gr['url'].lower().rstrip('/')
    if u in existing_in_sheet:
        continue
    
    row_idx = ws.max_row + 1
    ws.append([gr['company'], gr['title'], gr['location'], gr['notes'], gr['status'], gr['url']])
    
    for c in range(1, 7):
        cell = ws.cell(row=row_idx, column=c)
        cell.fill = green_fill
        cell.alignment = Alignment(wrap_text=True, vertical='top')
    
    ws.cell(row=row_idx, column=5).font = bold_green
    
    url_c = ws.cell(row=row_idx, column=6)
    url_c.hyperlink = gr['url']
    url_c.font = Font(color='0563C1', underline='single')
    added_count += 1

wb.save('data/Job Search.xlsx')
print(f"Added {added_count} high-priority GCC roles to 'New Additions (Wave 5)'. Total rows now: {ws.max_row - 1}")
