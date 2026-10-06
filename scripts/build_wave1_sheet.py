import csv, re, json, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# 1. Load probed URLs
with open('data/wave1_probed_urls.json') as f:
    probed_urls = json.load(f)

# 2. Load existing workbook to build duplicate indexes
wb = openpyxl.load_workbook('data/jobs_vrinda_discovery_wave2_verified.xlsx')

rq_lookup = {}
ws_rq = wb['Review Queue']
for r in range(2, ws_rq.max_row + 1):
    c = str(ws_rq.cell(row=r, column=2).value or '').strip().lower()
    t = str(ws_rq.cell(row=r, column=3).value or '').strip().lower()
    u = str(ws_rq.cell(row=r, column=11).value or '').strip()
    rq_lookup[f'{c}|{t}'] = r
    if u:
        m = re.search(r'jobs/view/(\d+)', u)
        if m: rq_lookup[f'linkedin_{m.group(1)}'] = r
        rq_lookup[u.lower().rstrip('/')] = r

nb_lookup = {}
ws_nb = wb['New Batch']
for r in range(2, ws_nb.max_row + 1):
    c = str(ws_nb.cell(row=r, column=2).value or '').strip().lower()
    t = str(ws_nb.cell(row=r, column=3).value or '').strip().lower()
    u = str(ws_nb.cell(row=r, column=11).value or '').strip()
    nb_lookup[f'{c}|{t}'] = r
    if u:
        m = re.search(r'jobs/view/(\d+)', u)
        if m: nb_lookup[f'linkedin_{m.group(1)}'] = r
        nb_lookup[u.lower().rstrip('/')] = r

# 3. Read raw CSV
with open('data/raw_user_wave1_input.csv', 'r', encoding='utf-8') as f:
    rows = list(csv.reader(f))

entries = []
seen_in_wave1 = set()

def determine_link_status(url):
    if not url:
        return 'No Direct Link (Referral / HR Lead)'
    clean_u = url.rstrip(').;\"\'')
    code = probed_urls.get(clean_u)
    if not code:
        # try matching prefix or substring
        for pu, pcode in probed_urls.items():
            if pu in clean_u or clean_u in pu:
                code = pcode
                break
    if code == '200':
        return '200 Active Requisition'
    elif code == '202':
        return '202 Accepted (Active Portal)'
    elif code == '410':
        return '410 Expired / Closed'
    elif code == '404':
        return '404 Dead Link'
    elif code in ('401', '403'):
        return f'{code} Portal / WAF Gated'
    elif code == '429':
        return '429 LinkedIn Rate Limited (View in Browser)'
    elif url.startswith('http'):
        return 'Active URL (Unverified)'
    else:
        return 'No Direct Link'

def evaluate_eligibility(role, exp_text, notes, tech):
    combined = f'{role} {exp_text} {notes} {tech}'.lower()
    if any(k in combined for k in ['7+', '7 years', '8+', '8 years', 'lead', 'manager', 'director', 'vp', 'executive']):
        if 'sde ii' not in combined and 'sde 2' not in combined:
            return 'Not Eligible (Experience 7+ Yrs / Over-Senior)', 'Requires 7+ years or lead/management seniority exceeding 3y candidate experience.'
    if any(k in combined for k in ['ui heavy', 'frontend-developer-sde-i', '6 months to 3 years, not applying', 'fresher', 'intern']):
        return 'Not Eligible (UI Heavy / Fresher / Excluded)', 'Explicitly flagged as UI heavy, entry level, or out of scope for backend focus.'
    if any(k in combined for k in ['sde 2', 'sde ii', 'sde-2', 'sde-ii', 'software engineer 2', 'software engineer ii', 'backend engineer', 'cloud engineer']):
        if any(k in combined for k in ['c#', '.net', 'azure', 'node.js', 'typescript', 'kafka', 'microservices', 'mcp']):
            return 'Eligible (High Match - Bullseye)', 'Perfect alignment with candidate stack (C#/.NET, Azure, Node.js/TS, Kafka, Distributed Systems, 2–5y).'
        else:
            return 'Eligible (High Match)', 'Direct SDE II backend match within candidate experience window (2–5 years).'
    if any(k in combined for k in ['sde iii', 'sde 3', 'senior software engineer', 'senior backend']):
        return 'Eligible (Stretch Match)', 'Senior level (4–6y / SDE III) achievable as stretch promotion from Microsoft SDE.'
    if 'referral' in combined or 'recruiter' in combined or 'hr' in combined or 'contact' in combined:
        return 'Referral / Pipeline Lead', 'Direct network referral, HR recruiter, or company pipeline contact.'
    return 'Potential Match / Pipeline Lead', 'Company engineering opportunity requiring candidate application or referral.'

def check_duplicate(comp, title, url):
    comp_clean = comp.strip().lower()
    title_clean = title.strip().lower()
    key_ct = f'{comp_clean}|{title_clean}'
    
    url_key = None
    if url and url.startswith('http'):
        m = re.search(r'jobs/view/(\d+)', url)
        if m:
            url_key = f'linkedin_{m.group(1)}'
        else:
            url_key = url.lower().rstrip('/')

    match_nb = nb_lookup.get(key_ct) or (nb_lookup.get(url_key) if url_key else None)
    match_rq = rq_lookup.get(key_ct) or (rq_lookup.get(url_key) if url_key else None)
    
    is_wave1_dup = key_ct in seen_in_wave1 or (url_key in seen_in_wave1 if url_key else False)
    if key_ct: seen_in_wave1.add(key_ct)
    if url_key: seen_in_wave1.add(url_key)

    dups = []
    if match_nb:
        dups.append(f'Also in New Batch (Row #{match_nb})')
    if match_rq:
        dups.append(f'Also in Review Queue (Row #{match_rq})')
    if is_wave1_dup:
        dups.append('Duplicate entry in Wave 1')
    
    if dups:
        return '; '.join(dups)
    return 'Unique in Wave 1'

# Parse Section 1: Lines 1 to 251
current_comp = ''
for idx in range(1, 252):
    r = rows[idx]
    if not any(r): continue
    
    comp = r[0].strip() if len(r) > 0 and r[0].strip() else current_comp
    if r[0].strip():
        current_comp = r[0].strip()
        
    opening = r[1].strip() if len(r) > 1 else ''
    contact = r[2].strip() if len(r) > 2 else ''
    col3 = r[3].strip() if len(r) > 3 else ''
    loc = r[4].strip() if len(r) > 4 else ''
    pay = r[5].strip() if len(r) > 5 else ''
    notes = r[6].strip() if len(r) > 6 else ''
    
    # check amazon sde 2 list
    if comp.lower() == 'amazon' and re.match(r'^\d+$', col3):
        job_id = col3
        url = f'https://www.amazon.jobs/en/jobs/{job_id}'
        role = opening if opening else 'Software Development Engineer II'
        location = contact if contact else (loc if loc else 'Bengaluru / Hyderabad')
        exp = '2-5 years (SDE II)'
        tech = 'Distributed Systems, Java/C#, AWS, Microservices'
        comp_est = 'Amazon SDE II base: ₹45L–₹65L+ LPA'
        
        elig, elig_notes = evaluate_eligibility(role, exp, notes, tech)
        link_stat = determine_link_status(url)
        dup_stat = check_duplicate(comp, role, url)
        
        entries.append({
            'section': 'Section 1 (Amazon SDE II)',
            'company': comp,
            'role': role,
            'location': location,
            'experience': exp,
            'tech': tech,
            'compensation': comp_est,
            'eligibility': elig,
            'eligibility_notes': elig_notes,
            'url': url,
            'link_status': link_stat,
            'contact_notes': notes if notes else f'Amazon Requisition ID: {job_id}',
            'dup_status': dup_stat
        })
        continue

    # check URLs in row
    line_urls = re.findall(r'https?://[^\s,\"\']+', ' '.join(r))
    if line_urls:
        for u in line_urls:
            u_clean = u.rstrip(').;\"\'')
            role = opening if opening else 'Software Engineer / Requisition Lead'
            location = loc if loc else 'Gurugram / India'
            exp = '2-5 years' if ('sde' in opening.lower() or 'ii' in opening.lower() or 'developer' in opening.lower()) else '2-5 years target'
            tech = 'Backend / Microservices / Cloud'
            comp_est = pay if pay else 'Estimated ₹25L–₹45L+ LPA'
            
            elig, elig_notes = evaluate_eligibility(role, exp, f'{opening} {notes}', tech)
            link_stat = determine_link_status(u_clean)
            dup_stat = check_duplicate(comp, role, u_clean)
            
            entries.append({
                'section': 'Section 1 (Direct URLs)',
                'company': comp,
                'role': role,
                'location': location,
                'experience': exp,
                'tech': tech,
                'compensation': comp_est,
                'eligibility': elig,
                'eligibility_notes': elig_notes,
                'url': u_clean,
                'link_status': link_stat,
                'contact_notes': f'Contact: {contact}; Notes: {opening} {notes}'.strip('; '),
                'dup_status': dup_stat
            })
    else:
        # non-URL tracker row
        role = opening if opening else 'General Engineering / Pipeline Lead'
        location = loc if loc else 'Gurugram / India'
        exp = '2-5 years target'
        tech = 'Backend / Cloud Architecture'
        comp_est = pay if pay else 'Target ≥ ₹35 LPA'
        
        elig, elig_notes = evaluate_eligibility(role, exp, f'{opening} {contact} {notes}', tech)
        link_stat = 'No Direct Link (Referral / HR Lead)'
        dup_stat = check_duplicate(comp, role, '')
        
        entries.append({
            'section': 'Section 1 (Company / Referral Lead)',
            'company': comp,
            'role': role,
            'location': location,
            'experience': exp,
            'tech': tech,
            'compensation': comp_est,
            'eligibility': elig,
            'eligibility_notes': elig_notes,
            'url': '',
            'link_status': link_stat,
            'contact_notes': f'Contact: {contact}; Status: {opening}; {notes}'.strip('; '),
            'dup_status': dup_stat
        })

# Parse Section 2: Lines 253 to 295
# ['Employer', 'Role Title', 'Location', 'Key Stack & Compensation Fit', 'Direct Application Link']
for idx in range(253, 296):
    r = rows[idx]
    if not any(r): continue
    comp = r[0].strip()
    role = r[1].strip() if len(r) > 1 else 'Software Engineer'
    loc = r[2].strip() if len(r) > 2 else 'India'
    stack_comp = r[3].strip() if len(r) > 3 else ''
    raw_link = r[4].strip() if len(r) > 4 else ''
    
    # extract URL if present
    m_url = re.search(r'https?://[^\s,\"\']+', raw_link + ' ' + ' '.join(r[5:]))
    url = m_url.group(0).rstrip(').;\"\'') if m_url else ''
    
    # check for linkedin job ID or requisition
    m_li = re.search(r'LinkedIn View (\d+)', raw_link)
    if m_li and not url:
        url = f'https://www.linkedin.com/jobs/view/{m_li.group(1)}'
        
    m_gh = re.search(r'Greenhouse Req (\d+)', raw_link)
    if m_gh and not url:
        url = f'https://boards.greenhouse.io/{comp.lower().replace(" ", "")}/jobs/{m_gh.group(1)}'
        
    m_ash = re.search(r'Ashby Req ([0-9a-zA-Z]+)', raw_link)
    if m_ash and not url:
        url = f'https://jobs.ashbyhq.com/{comp.lower().replace(" ", "")}/{m_ash.group(1)}'

    exp = '2-5 years'
    tech = stack_comp
    comp_est = stack_comp if ('$' in stack_comp or '₹' in stack_comp) else 'Target ≥ ₹35 LPA'
    
    elig, elig_notes = evaluate_eligibility(role, exp, stack_comp, tech)
    link_stat = determine_link_status(url)
    dup_stat = check_duplicate(comp, role, url)
    
    entries.append({
        'section': 'Section 2 (Discovered Opportunities)',
        'company': comp,
        'role': role,
        'location': loc,
        'experience': exp,
        'tech': tech,
        'compensation': comp_est,
        'eligibility': elig,
        'eligibility_notes': elig_notes,
        'url': url,
        'link_status': link_stat,
        'contact_notes': f'Application Route: {raw_link}; Fit: {stack_comp}',
        'dup_status': dup_stat
    })

# Parse Section 3: Lines 297 to 331
# ['Company', 'Role', 'Location', '', 'Salary (ESTIMATE)', 'Apply link', 'Referral contact', '', 'Notes', '']
current_sec3_comp = ''
for idx in range(297, 332):
    r = rows[idx]
    if not any(r): continue
    comp = r[0].strip() if r[0].strip() else current_sec3_comp
    if r[0].strip(): current_sec3_comp = r[0].strip()
    
    role = r[1].strip() if len(r) > 1 and r[1].strip() else 'Software Engineer'
    loc = r[2].strip() if len(r) > 2 and r[2].strip() else 'Gurugram / India'
    sal = r[4].strip() if len(r) > 4 else ''
    raw_link = r[5].strip() if len(r) > 5 else ''
    contact = r[6].strip() if len(r) > 6 else ''
    notes = r[8].strip() if len(r) > 8 else ''
    
    m_url = re.search(r'https?://[^\s,\"\']+', raw_link)
    url = m_url.group(0).rstrip(').;\"\'') if m_url else ''
    
    exp = '2-5 years'
    tech = 'Backend / Fullstack / Cloud'
    comp_est = sal if sal else 'Target ≥ ₹35 LPA'
    
    elig, elig_notes = evaluate_eligibility(role, exp, f'{notes} {sal}', tech)
    link_stat = determine_link_status(url)
    dup_stat = check_duplicate(comp, role, url)
    
    entries.append({
        'section': 'Section 3 (Extended Roles & Status)',
        'company': comp,
        'role': role,
        'location': loc,
        'experience': exp,
        'tech': tech,
        'compensation': comp_est,
        'eligibility': elig,
        'eligibility_notes': elig_notes,
        'url': url,
        'link_status': link_stat,
        'contact_notes': f'Contact: {contact}; Notes: {notes}',
        'dup_status': dup_stat
    })

# Parse Section 4: Lines 333 to 411
# ['Company', 'Name', 'Linkedin Title', '', '', '', '', '', 'Job role', 'Openings(Agent working on this)']
for idx in range(333, len(rows)):
    r = rows[idx]
    if not any(r): continue
    comp = r[0].strip() if r[0].strip() else 'Alumni / Referral Network'
    name = r[1].strip() if len(r) > 1 else ''
    title = r[2].strip() if len(r) > 2 else ''
    job_role = r[8].strip() if len(r) > 8 else ''
    openings = r[9].strip() if len(r) > 9 else ''
    
    role = job_role if job_role else f'Referral / HR Network Lead ({name})'
    loc = 'India / Global'
    exp = 'N/A (Referral Contact)'
    tech = f'Contact Title: {title}'
    comp_est = 'Referral Connection'
    
    elig = 'Referral / HR Network Lead'
    elig_notes = f'Internal employee / HR leader ({name}, {title}) available for warm referral outreach.'
    link_stat = 'Referral Contact (LinkedIn / Internal)'
    dup_stat = 'Network Contact'
    
    entries.append({
        'section': 'Section 4 (Referral Network Contacts)',
        'company': comp,
        'role': role,
        'location': loc,
        'experience': exp,
        'tech': tech,
        'compensation': comp_est,
        'eligibility': elig,
        'eligibility_notes': elig_notes,
        'url': '',
        'link_status': link_stat,
        'contact_notes': f'Name: {name} | Profile: {title} | Tag: {job_role} | Notes: {openings}',
        'dup_status': dup_stat
    })

print(f'Total parsed entries for Wave 1 sheet: {len(entries)}')

# 4. Create or overwrite 'wave 1' sheet in Excel workbook
sheet_name = 'wave 1'
if sheet_name in wb.sheetnames:
    del wb[sheet_name]

ws_w1 = wb.create_sheet(title=sheet_name)

# Define Headers
headers = [
    '#',
    'Source Section',
    'Company',
    'Role Title / Target Role',
    'Location',
    'Experience Required',
    'Key Stack / Technical Focus',
    'Estimated / Stated Compensation',
    'Role Eligibility Decision',
    'Eligibility Assessment & Rationale',
    'Application / Requisition URL',
    'Link Health / HTTP Status',
    'Referral Contact & Status Notes',
    'Database Duplicate Check'
]

# Styling definitions
header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
header_fill = PatternFill(start_color='1F497D', end_color='1F497D', fill_type='solid') # Professional Navy
font_data = Font(name='Calibri', size=10)
font_link = Font(name='Calibri', size=10, color='0563C1', underline='single')
align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)
thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)
fill_even = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')
fill_odd = PatternFill(start_color='F9FAFB', end_color='F9FAFB', fill_type='solid')

# Write Header
ws_w1.row_dimensions[1].height = 28
for col_idx, h in enumerate(headers, 1):
    cell = ws_w1.cell(row=1, column=col_idx, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = align_center
    cell.border = thin_border

# Write Data Rows
for r_idx, item in enumerate(entries, 2):
    ws_w1.row_dimensions[r_idx].height = 24
    row_fill = fill_odd if r_idx % 2 == 1 else fill_even
    
    row_vals = [
        r_idx - 1,
        item['section'],
        item['company'],
        item['role'],
        item['location'],
        item['experience'],
        item['tech'],
        item['compensation'],
        item['eligibility'],
        item['eligibility_notes'],
        item['url'],
        item['link_status'],
        item['contact_notes'],
        item['dup_status']
    ]
    
    for c_idx, val in enumerate(row_vals, 1):
        cell = ws_w1.cell(row=r_idx, column=c_idx, value=val)
        cell.font = font_data
        cell.fill = row_fill
        cell.border = thin_border
        
        if c_idx in (1, 9, 12, 14):
            cell.alignment = align_center
        else:
            cell.alignment = align_left
            
        if c_idx == 11 and val and str(val).startswith('http'):
            cell.hyperlink = val
            cell.font = font_link

# Set column widths
col_widths = {
    1: 6,   # #
    2: 24,  # Section
    3: 20,  # Company
    4: 30,  # Role
    5: 22,  # Location
    6: 18,  # Exp
    7: 32,  # Tech
    8: 24,  # Comp
    9: 25,  # Eligibility
    10: 38, # Rationale
    11: 35, # URL
    12: 26, # Link Status
    13: 35, # Contact Notes
    14: 30  # Duplicate Check
}

for col_idx, width in col_widths.items():
    ws_w1.column_dimensions[get_column_letter(col_idx)].width = width

ws_w1.freeze_panes = 'C2'

wb.save('data/jobs_vrinda_discovery_wave2_verified.xlsx')
print(f'Successfully created sheet "wave 1" with {len(entries)} verified entries!')
