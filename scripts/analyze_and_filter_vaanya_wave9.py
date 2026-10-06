import json
import re
import html
import openpyxl

# Load seen URLs
seen_urls = set()
with open('data/all_existing_urls.json') as f:
    for u in json.load(f):
        seen_urls.add(u.strip().lower().rstrip('/'))

wb = openpyxl.load_workbook('data/Job Search.xlsx', data_only=True)
for sname in wb.sheetnames:
    ws = wb[sname]
    for row in ws.iter_rows(values_only=True):
        for cell in row:
            if cell and isinstance(cell, str) and cell.startswith('http'):
                seen_urls.add(cell.strip().lower().rstrip('/'))

with open('data/scratch_raw_comprehensive_ats.json') as f:
    raw_jobs = json.load(f)

print(f"Total raw postings: {len(raw_jobs)}")

paused_employers = {'amazon', 'sarvam', 'sarvam ai', 'mongodb'}

# Exclude non-tech or high seniority or internships
reject_title_words = [
    r'\bintern\b', r'\binternship\b', r'\bsummer\b', r'\bwinter\b', r'\bco-op\b', r'\bcoop\b',
    r'\bsenior\b', r'\bsr\.?\b', r'\blead\b', r'\bprincipal\b', r'\bstaff\b',
    r'\bii\b', r'\biii\b', r'\biv\b', r'\b2\b', r'\b3\b', r'\b4\b', r'\b5\b',
    r'\bmanager\b', r'\bdirector\b', r'\bvp\b', r'\bhead\b',
    r'\brecruiter\b', r'\btalent\b', r'\bhr\b', r'\bsales\b', r'\bmarketing\b',
    r'\boperations\b', r'\baccountant\b', r'\blegal\b', r'\bexecutive\b',
    r'\bpartner\b', r'\bconsultant\b', r'\banalyst\b'
]

tech_title_words = [
    r'\bsoftware\b', r'\bsde\b', r'\bengineer\b', r'\bdeveloper\b',
    r'\bbackend\b', r'\bdata\b', r'\bml\b', r'\bmachine\s+learning\b',
    r'\bai\b', r'\bqa\b', r'\bsdet\b', r'\btest\b', r'\bpython\b',
    r'\bfull\s*stack\b', r'\bfrontend\b', r'\bplatform\b', r'\bcloud\b',
    r'\bsystems\b', r'\bgraduate\b', r'\btrainee\b', r'\bget\b'
]

# Strict entry/fresher markers in title
fresher_title_markers = [
    r'\bi\b', r'\b1\b', r'\bassociate\b', r'\bjunior\b', r'\bjr\.?\b',
    r'\bgraduate\b', r'\btrainee\b', r'\bget\b', r'\bearly\s+career\b',
    r'\buniversity\b', r'\bcampus\b', r'\bentry\b', r'\b2026\b'
]

india_locations = [
    'india', 'bengaluru', 'bangalore', 'gurgaon', 'gurugram', 'noida',
    'delhi', 'hyderabad', 'pune', 'mumbai', 'chennai'
]

matches = []

for j in raw_jobs:
    comp = j.get('company', '').lower()
    if any(p in comp for p in paused_employers):
        continue

    title = j.get('title', '').strip()
    loc = j.get('location', '').strip()
    url = j.get('url', '').strip()
    content = j.get('content', '')

    clean_url = url.lower().rstrip('/')
    if clean_url in seen_urls:
        continue

    loc_lower = loc.lower()
    is_india = any(loc_token in loc_lower for loc_token in india_locations)
    if not is_india:
        # Check if remote with explicit India context in description
        if 'remote' in loc_lower and ('india' in content.lower()):
            is_india = True
        else:
            continue

    t_lower = title.lower()
    if any(re.search(pat, t_lower) for pat in reject_title_words):
        continue

    if not any(re.search(pat, t_lower) for pat in tech_title_words):
        continue

    # Experience check in content
    content_lower = content.lower()
    if content_lower:
        # Exclude 3+, 4+, 5+
        if re.search(r'\b([3-9]|\d{2,})\+?\s*(?:to\s*\d+\s*)?years?\b', content_lower):
            continue
        if re.search(r'\b([3-9]|\d{2,})\s*[\u2013\u2014]\s*\d+\s*years?\b', content_lower):
            continue
        # Exclude strict 2+ years
        if re.search(r'\b(?:minimum|at least|min\.?)\s*2\+?\s*years?\b', content_lower):
            continue
        if re.search(r'\b2\+?\s*years\s+of\s+(?:relevant\s+)?experience\b', content_lower) and not ('0-2' in content_lower or '0 to 2' in content_lower):
            continue

    # Title check for entry-level signal OR content explicitly mentioning 0 YOE / fresher / entry
    is_entry_title = any(re.search(pat, t_lower) for pat in fresher_title_markers)
    has_fresher_content = False
    if content_lower:
        if any(w in content_lower for w in ['fresher', '0-1 year', '0 to 1 year', '0-2 year', '0 to 2 year', 'graduate engineer trainee', 'campus', 'early career', 'bachelor', 'new grad']):
            has_fresher_content = True

    # If title is generic (e.g. "Software Engineer"), require either entry title marker or fresher content
    if not is_entry_title and not has_fresher_content and j.get('platform') != 'Ashby':
        continue

    matches.append({
        'company': j.get('company'),
        'title': title,
        'location': loc,
        'url': url,
        'platform': j.get('platform'),
        'id': j.get('id'),
        'content': content
    })

print(f"Total matching candidates: {len(matches)}")
with open('data/scratch_wave9_raw_matches.json', 'w') as f:
    json.dump([{k: v for k, v in m.items() if k != 'content'} for m in matches], f, indent=2)

for m in matches[:25]:
    print(f"[{m['platform']}] {m['company']} | {m['title']} | {m['location']} | {m['url']}")
