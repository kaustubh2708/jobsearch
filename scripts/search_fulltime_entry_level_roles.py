import urllib.parse
import time
import json
import re
from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from bs4 import BeautifulSoup
import openpyxl

# 1. Load all existing URLs and company+role keys to prevent ANY duplicate
wb = openpyxl.load_workbook('data/Job Search.xlsx')
existing_urls = set()
existing_keys = set()

def normalize_text(text):
    if not text:
        return ""
    t = str(text).lower()
    t = re.sub(r'[^a-z0-9]', ' ', t)
    return ' '.join(t.split())

for sname in wb.sheetnames:
    ws = wb[sname]
    for r in range(2, ws.max_row + 1):
        comp = ws.cell(r, 1).value
        title = ws.cell(r, 2).value
        if comp and title:
            existing_keys.add(f"{normalize_text(comp)}_{normalize_text(title)}")
        
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(r, c)
            if cell.hyperlink and cell.hyperlink.target:
                existing_urls.add(cell.hyperlink.target.lower().rstrip('/'))
            elif cell.value and isinstance(cell.value, str) and cell.value.startswith('http'):
                existing_urls.add(cell.value.lower().rstrip('/'))

print(f"Loaded {len(existing_urls)} existing URLs and {len(existing_keys)} existing company-role keys for deduplication.")

options = Options()
options.add_argument('--headless=new')
options.add_argument('--no-sandbox')
options.add_argument('--disable-dev-shm-usage')
options.add_argument('--user-agent=Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36')

# Targeted queries strictly for Entry Level Full-Time (f_E=2: Entry Level, f_JT=F: Full-time)
searches = [
    # Noida & Gurgaon Entry-level Full-Time
    {'kw': 'Software Engineer I', 'loc': 'Noida, Uttar Pradesh, India', 'extra': '&f_E=2&f_JT=F'},
    {'kw': 'Python Backend Engineer', 'loc': 'Noida, Uttar Pradesh, India', 'extra': '&f_E=2&f_JT=F'},
    {'kw': 'Graduate Engineer Trainee', 'loc': 'Noida, Uttar Pradesh, India', 'extra': '&f_JT=F'},
    {'kw': 'Software Development Engineer I', 'loc': 'Gurgaon, Haryana, India', 'extra': '&f_E=2&f_JT=F'},
    {'kw': 'Associate Software Engineer', 'loc': 'Gurgaon, Haryana, India', 'extra': '&f_E=2&f_JT=F'},
    {'kw': 'Junior Python Developer', 'loc': 'Gurgaon, Haryana, India', 'extra': '&f_E=2&f_JT=F'},
    {'kw': 'Junior Data Engineer', 'loc': 'Gurgaon, Haryana, India', 'extra': '&f_E=2&f_JT=F'},
    {'kw': 'Graduate Engineer Trainee', 'loc': 'Gurgaon, Haryana, India', 'extra': '&f_JT=F'},
    # Delhi & NCR Tech Unicorns
    {'kw': 'Junior Machine Learning Engineer', 'loc': 'Delhi, India', 'extra': '&f_E=2&f_JT=F'},
    {'kw': 'Associate Machine Learning Engineer', 'loc': 'Gurgaon, Haryana, India', 'extra': '&f_E=2&f_JT=F'},
    {'kw': 'Junior Backend Developer', 'loc': 'Delhi, India', 'extra': '&f_E=2&f_JT=F'},
    # Remote / Bengaluru Top Tech Entry Level
    {'kw': 'Software Engineer I', 'loc': 'Remote', 'extra': '&f_WT=2&f_E=2&f_JT=F'},
    {'kw': 'Junior Backend Engineer', 'loc': 'Remote', 'extra': '&f_WT=2&f_E=2&f_JT=F'},
    {'kw': 'Software Development Engineer I', 'loc': 'Bengaluru, Karnataka, India', 'extra': '&f_E=2&f_JT=F'},
    {'kw': 'Graduate Software Engineer', 'loc': 'Bengaluru, Karnataka, India', 'extra': '&f_E=2&f_JT=F'},
    {'kw': 'Associate Data Engineer', 'loc': 'Bengaluru, Karnataka, India', 'extra': '&f_E=2&f_JT=F'}
]

driver = webdriver.Chrome(options=options)
raw_cards = {}

reject_title_words = [
    'intern', 'internship', 'summer', 'winter', 'lead', 'senior', 'sr.', 'sr ', 'manager',
    'principal', 'director', 'staff', 'ii', 'iii', '2', '3', 'recruiter', 'hr', 'sales',
    'marketing', 'business development', 'product management', 'operations'
]

try:
    for idx, s in enumerate(searches):
        kw = s['kw']
        loc = s['loc']
        extra = s.get('extra', '')
        url = f"https://www.linkedin.com/jobs/search?keywords={urllib.parse.quote(kw)}&location={urllib.parse.quote(loc)}{extra}"
        print(f"[{idx+1}/{len(searches)}] Searching: '{kw}' in '{loc}'...")
        try:
            driver.get(url)
            time.sleep(2)
            soup = BeautifulSoup(driver.page_source, 'html.parser')
            cards = soup.find_all('div', class_='base-card')
            found_count = 0
            for c in cards:
                t = c.find('h3', class_='base-search-card__title')
                comp = c.find('h4', class_='base-search-card__subtitle')
                l = c.find('span', class_='job-search-card__location')
                a = c.find('a', class_='base-card__full-link')
                
                title = t.get_text(strip=True) if t else ''
                company = comp.get_text(strip=True) if comp else ''
                location = l.get_text(strip=True) if l else ''
                link = a['href'] if a and a.has_attr('href') else ''
                
                if not link:
                    continue
                
                clean_url = link.split('?')[0]
                
                # Exclude reject words in title
                title_lower = title.lower()
                if any(w in title_lower for w in reject_title_words):
                    continue
                
                # Check duplication
                norm_key = f"{normalize_text(company)}_{normalize_text(title)}"
                if clean_url.lower().rstrip('/') in existing_urls or norm_key in existing_keys:
                    continue
                
                m = re.search(r'-([0-9]{8,12})(?:$)', clean_url)
                jid = m.group(1) if m else clean_url
                
                if jid not in raw_cards:
                    raw_cards[jid] = {
                        'company': company,
                        'title': title,
                        'location': location,
                        'url': clean_url,
                        'source': 'LinkedIn Entry-Level'
                    }
                    found_count += 1
            print(f"   -> Added {found_count} new candidate cards (Total: {len(raw_cards)})")
        except Exception as e:
            print(f"   -> Error: {e}")
        time.sleep(1)

    print(f"\nDiscovered {len(raw_cards)} potential full-time entry-level candidates.")
    with open('data/raw_entry_level_cards.json', 'w') as f:
        json.dump(list(raw_cards.values()), f, indent=2)
finally:
    driver.quit()
