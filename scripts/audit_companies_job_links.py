import urllib.request, ssl, json, re, time
from urllib.parse import urlparse
from bs4 import BeautifulSoup
import openpyxl

ctx = ssl.create_default_context()
ctx.check_hostname = False
ctx.verify_mode = ssl.CERT_NONE
headers = {
    'User-Agent': 'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36'
}

wb = openpyxl.load_workbook('data/companies.xlsx')
ws = wb['Job Links']

entries = []

for r in range(1, 130):
    row_vals = [ws.cell(row=r, column=c).value for c in range(1, 13)]
    row_links = [ws.cell(row=r, column=c).hyperlink.target if ws.cell(row=r, column=c).hyperlink else None for c in range(1, 13)]
    
    comp = ws.cell(row=r, column=1).value
    if comp and 'bot was not able to scrap' in str(comp):
        continue
        
    title = ws.cell(row=r, column=2).value
    loc = ws.cell(row=r, column=3).value
    jid = ws.cell(row=r, column=4).value
    notes = ws.cell(row=r, column=6).value

    target_url = None
    for link in row_links:
        if link and str(link).startswith('http') and not str(link).startswith('http://karya.ai') and not str(link).startswith('http://corover.ai'):
            target_url = str(link).strip()
            break
    if not target_url:
        for val in row_vals:
            if val and str(val).startswith('http') and not str(val).startswith('http://karya.ai') and not str(val).startswith('http://corover.ai'):
                target_url = str(val).strip()
                break

    entries.append({
        'row': r,
        'company': str(comp).strip() if comp else '',
        'title': str(title).strip() if title else '',
        'location': str(loc).strip() if loc else '',
        'job_id': str(jid).strip() if jid else '',
        'url': target_url,
        'notes': str(notes).strip() if notes else ''
    })

print(f"Total entries loaded: {len(entries)}")
entries_with_url = [e for e in entries if e['url']]
print(f"Entries to test: {len(entries_with_url)}")

audit_results = []

for idx, item in enumerate(entries_with_url, 1):
    url = item['url']
    comp = item['company']
    title = item['title']
    print(f"[{idx}/{len(entries_with_url)}] Auditing R{item['row']}: {comp} - {title}...")
    
    status_code = None
    final_url = url
    is_active = False
    reason = ""
    page_text_snippet = ""
    
    try:
        req = urllib.request.Request(url, headers=headers)
        with urllib.request.urlopen(req, context=ctx, timeout=12) as resp:
            status_code = resp.getcode()
            final_url = resp.geturl()
            html_bytes = resp.read()
            soup = BeautifulSoup(html_bytes, 'html.parser')
            page_text = soup.get_text(' ', strip=True).lower()
            page_text_snippet = page_text[:500]

            # Heuristics for closed/expired
            closed_markers = [
                'no longer accepting applications',
                'this job is no longer available',
                'this position has been filled',
                'job has been closed',
                'job not found',
                'page not found',
                'requisition has closed',
                'no longer open',
                'job is no longer active',
                'not accepting new applicants',
                '404 not found'
            ]

            found_closed = any(m in page_text for m in closed_markers)
            
            # Specific domain heuristics
            if 'amazon.jobs' in url:
                if 'not found' in page_text or 'job is closed' in page_text or 'no longer open' in page_text:
                    is_active = False
                    reason = "Amazon job posting has closed / not found"
                elif 'apply now' in page_text or 'job description' in page_text or 'basic qualifications' in page_text:
                    is_active = True
                    reason = "Active Amazon job page with application button"
                else:
                    is_active = True
                    reason = "HTTP 200 response"
            elif 'greenhouse.io' in url:
                if 'no longer available' in page_text or 'job has closed' in page_text:
                    is_active = False
                    reason = "Greenhouse posting closed"
                else:
                    is_active = True
                    reason = "Active Greenhouse ATS requisition"
            elif 'ashbyhq.com' in url:
                if 'no longer available' in page_text or 'closed' in page_text:
                    is_active = False
                    reason = "Ashby posting closed"
                else:
                    is_active = True
                    reason = "Active Ashby ATS requisition"
            elif 'linkedin.com' in url:
                if 'no longer accepting applications' in page_text:
                    is_active = False
                    reason = "LinkedIn listing is closed (no longer accepting applications)"
                else:
                    is_active = True
                    reason = "Active LinkedIn job listing"
            else:
                if found_closed:
                    is_active = False
                    reason = f"Closed marker found on page"
                else:
                    is_active = True
                    reason = "Active HTTP 200 page"

    except urllib.error.HTTPError as e:
        status_code = e.code
        is_active = False
        if e.code == 404:
            reason = "HTTP 404: Dead Link (Job expired or removed)"
        elif e.code == 410:
            reason = "HTTP 410: Gone (Job closed)"
        elif e.code == 403:
            # Check if domain blocks python (e.g. Workday / Cloudflare)
            reason = "HTTP 403: Bot Protection / Cloudflare WAF"
            # We will test in selenium if 403
        else:
            reason = f"HTTP Error {e.code}"
    except Exception as e:
        status_code = "ERR"
        is_active = False
        reason = f"Connection error: {e}"

    audit_results.append({
        'row': item['row'],
        'company': comp,
        'title': title,
        'location': item['location'],
        'job_id': item['job_id'],
        'url': url,
        'final_url': final_url,
        'http_status': status_code,
        'is_active': is_active,
        'reason': reason,
        'notes': item['notes']
    })
    print(f"   -> Status: {status_code} | Active: {is_active} | {reason[:60]}")
    time.sleep(0.3)

with open('data/audit_companies_links_initial.json', 'w') as f:
    json.dump(audit_results, f, indent=2)

print(f"\nInitial audit complete! Saved {len(audit_results)} results.")
