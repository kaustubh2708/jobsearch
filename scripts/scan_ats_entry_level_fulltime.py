import urllib.request
import ssl
import json
import re
import html
import openpyxl

wb = openpyxl.load_workbook('data/Job Search.xlsx')
existing_urls = set()
for sname in wb.sheetnames:
    ws = wb[sname]
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(r, c)
            if cell.hyperlink and cell.hyperlink.target:
                existing_urls.add(cell.hyperlink.target.lower().rstrip('/'))
            elif cell.value and isinstance(cell.value, str) and cell.value.startswith('http'):
                existing_urls.add(cell.value.lower().rstrip('/'))

ctx = ssl.create_default_context()
ctx.check_hostname = False
ctx.verify_mode = ssl.CERT_NONE
headers = {
    'User-Agent': 'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36',
    'Accept': 'application/json, text/html, */*'
}

greenhouse_boards = [
    'rubrik', 'stripe', 'towerresearchcapital', 'squarepointcapital',
    'postman', 'browserstack', 'inmobi', 'swiggy', 'sharechat',
    'urbancompany', 'cars24', 'physicswallah', 'sliceit', 'cred',
    'meesho', 'nutanix', 'aristanetworks', 'datadog', 'elastic',
    'couchbase', 'cloudera', 'purestorage', 'hashicorp', 'cloudflare',
    'roblox', 'snowflake', 'cohere', 'anthropic', 'gleanwork', 'mongodb', 'okta'
]

lever_boards = [
    'paytm', 'zomato', 'blinkit', 'zepto', 'incred', 'jupiter',
    'navi', 'fi-money', 'clevertap', 'urbanpiper', 'simpl', 'kuku-fm',
    'groww', 'razorpay'
]

ashby_boards = [
    'atlys', 'sarvam', 'avoca', 'forma-ai', 'atlan', 'aiprise',
    'together-ai', 'plane', 'squadstack', 'mercury', 'cursor'
]

# Strict Full-Time Entry Level Title Patterns
entry_title_patterns = [
    r'\bsoftware\s+engineer\s+i\b',
    r'\bsde\s*[-–—]?\s*(?:1|i)\b',
    r'\bassociate\s+(?:software|data|backend|ml|ai)\s+engineer\b',
    r'\bjunior\s+(?:software|data|backend|python|ml|ai)\s+(?:engineer|developer)\b',
    r'\bgraduate\s+software\s+engineer\b',
    r'\bgraduate\s+engineer\s+trainee\b',
    r'\bget\b',
    r'\bearly\s+career\s+software\s+engineer\b',
    r'\bsoftware\s+engineer\s+[-–—]?\s*(?:campus|university|2026)\b'
]

reject_patterns = [
    r'\bintern\b', r'\binternship\b', r'\bsummer\b', r'\bwinter\b',
    r'\bsenior\b', r'\bsr\.?\b', r'\blead\b', r'\bprincipal\b',
    r'\bii\b', r'\biii\b', r'\b2\b', r'\b3\b', r'\b4\b', r'\bstaff\b',
    r'\bmanager\b', r'\bdirector\b', r'\brecruiter\b', r'\bhr\b', r'\bsales\b'
]

results = []

def is_entry_fulltime_title(title):
    t = title.lower()
    if any(re.search(pat, t) for pat in reject_patterns):
        return False
    return any(re.search(pat, t) for pat in entry_title_patterns)

print("Scanning Greenhouse for Full-Time Entry Level roles...")
for slug in greenhouse_boards:
    url = f"https://boards-api.greenhouse.io/v1/boards/{slug}/jobs?content=true"
    try:
        req = urllib.request.Request(url, headers=headers)
        with urllib.request.urlopen(req, context=ctx, timeout=5) as r:
            data = json.loads(r.read())
            for j in data.get('jobs', []):
                title = j.get('title', '')
                loc = j.get('location', {}).get('name', '')
                abs_url = j.get('absolute_url', '')
                
                loc_lower = loc.lower()
                is_india = any(k in loc_lower for k in ['india', 'bengaluru', 'bangalore', 'gurgaon', 'gurugram', 'noida', 'delhi', 'hyderabad', 'pune', 'mumbai', 'remote'])
                if not is_india:
                    continue
                    
                if is_entry_fulltime_title(title):
                    content = html.unescape(re.sub(r'<[^>]+>', ' ', j.get('content', ''))).lower()
                    
                    # Hard reject 3+ YOE
                    has_high_exp = re.search(r'\b([3-9]|\d{2,})\+?\s*(?:to\s*\d+\s*)?years?\b', content)
                    has_high_exp_en = re.search(r'\b([3-9]|\d{2,})\s*[\u2013\u2014]\s*\d+\s*years?\b', content)
                    
                    # Check for 0-1 or 0-2 YOE
                    has_entry_exp = re.search(r'\b0\s*[\u2013\u2014\-to]\s*[1-2]\s*years?\b', content) or 'entry level' in content or 'new grad' in content or 'fresh' in content or '0 year' in content or 'graduate' in content
                    
                    if not has_high_exp and not has_high_exp_en:
                        clean_u = abs_url.lower().rstrip('/')
                        if clean_u not in existing_urls:
                            results.append({
                                'company': slug.capitalize(),
                                'title': title,
                                'location': loc,
                                'url': abs_url,
                                'source': 'Greenhouse',
                                'has_entry_exp': bool(has_entry_exp),
                                'content_snip': content[:400]
                            })
                            print(f"  [GH ENTRY MATCH] {slug}: {title} ({loc})")
    except:
        pass

print("\nScanning Lever for Full-Time Entry Level roles...")
for slug in lever_boards:
    url = f"https://api.lever.co/v0/postings/{slug}?mode=json"
    try:
        req = urllib.request.Request(url, headers=headers)
        with urllib.request.urlopen(req, context=ctx, timeout=5) as r:
            jobs = json.loads(r.read())
            for j in jobs:
                title = j.get('text', '')
                loc = j.get('categories', {}).get('location', '')
                h_url = j.get('hostedUrl', '')
                
                loc_lower = loc.lower() if loc else ''
                is_india = any(k in loc_lower for k in ['india', 'bengaluru', 'bangalore', 'gurgaon', 'gurugram', 'noida', 'delhi', 'hyderabad', 'pune', 'mumbai', 'remote'])
                if not is_india:
                    continue
                    
                if is_entry_fulltime_title(title):
                    desc = html.unescape(j.get('descriptionPlain', '') + ' ' + j.get('additionalPlain', '')).lower()
                    has_high_exp = re.search(r'\b([3-9]|\d{2,})\+?\s*(?:to\s*\d+\s*)?years?\b', desc)
                    has_high_exp_en = re.search(r'\b([3-9]|\d{2,})\s*[\u2013\u2014]\s*\d+\s*years?\b', desc)
                    
                    if not has_high_exp and not has_high_exp_en:
                        clean_u = h_url.lower().rstrip('/')
                        if clean_u not in existing_urls:
                            results.append({
                                'company': slug.capitalize(),
                                'title': title,
                                'location': loc,
                                'url': h_url,
                                'source': 'Lever',
                                'has_entry_exp': True,
                                'content_snip': desc[:400]
                            })
                            print(f"  [LEVER ENTRY MATCH] {slug}: {title} ({loc})")
    except:
        pass

print("\nScanning Ashby for Full-Time Entry Level roles...")
for slug in ashby_boards:
    url = f"https://api.ashbyhq.com/posting-api/job-board/{slug}"
    try:
        req = urllib.request.Request(url, headers=headers)
        with urllib.request.urlopen(req, context=ctx, timeout=5) as r:
            data = json.loads(r.read())
            for j in data.get('jobs', []):
                title = j.get('title', '')
                loc = j.get('location', '')
                jid = j.get('id', '')
                h_url = f"https://jobs.ashbyhq.com/{slug}/{jid}"
                
                loc_lower = loc.lower() if loc else ''
                is_india = any(k in loc_lower for k in ['india', 'bengaluru', 'bangalore', 'gurgaon', 'gurugram', 'noida', 'delhi', 'hyderabad', 'pune', 'mumbai', 'remote'])
                if not is_india:
                    continue
                    
                if is_entry_fulltime_title(title):
                    clean_u = h_url.lower().rstrip('/')
                    if clean_u not in existing_urls:
                        results.append({
                            'company': slug.capitalize(),
                            'title': title,
                            'location': loc,
                            'url': h_url,
                            'source': 'Ashby',
                            'has_entry_exp': True,
                            'content_snip': ''
                        })
                        print(f"  [ASHBY ENTRY MATCH] {slug}: {title} ({loc})")
    except:
        pass

with open('data/entry_level_ats_roles.json', 'w') as f:
    json.dump(results, f, indent=2)

print(f"\nTotal full-time entry level ATS roles found: {len(results)}")
