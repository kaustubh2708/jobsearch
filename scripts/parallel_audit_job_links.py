import urllib.request, ssl, json, re, time
from urllib.parse import urlparse
from bs4 import BeautifulSoup
import openpyxl
from concurrent.futures import ThreadPoolExecutor, as_completed

ctx = ssl.create_default_context()
ctx.check_hostname = False
ctx.verify_mode = ssl.CERT_NONE
headers = {
    'User-Agent': 'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36'
}

wb = openpyxl.load_workbook('data/companies.xlsx')
ws = wb['Job Links']

entries = []

for r in range(1, 130):
    row_vals = [ws.cell(row=r, column=c).value for c in range(1, 13)]
    row_links = [ws.cell(row=r, column=c).hyperlink.target if ws.cell(row=r, column=c).hyperlink else None for c in range(1, 13)]
    
    comp = ws.cell(row=r, column=1).value
    if comp and 'bot was not able to scrap' in str(comp):
        continue
        
    title = ws.cell(row=r, column=2).value
    loc = ws.cell(row=r, column=3).value
    jid = ws.cell(row=r, column=4).value
    notes = ws.cell(row=r, column=6).value

    target_url = None
    for link in row_links:
        if link and str(link).startswith('http') and not str(link).startswith('http://karya.ai') and not str(link).startswith('http://corover.ai'):
            target_url = str(link).strip()
            break
    if not target_url:
        for val in row_vals:
            if val and str(val).startswith('http') and not str(val).startswith('http://karya.ai') and not str(val).startswith('http://corover.ai'):
                target_url = str(val).strip()
                break

    entries.append({
        'row': r,
        'company': str(comp).strip() if comp else '',
        'title': str(title).strip() if title else '',
        'location': str(loc).strip() if loc else '',
        'job_id': str(jid).strip() if jid else '',
        'url': target_url,
        'notes': str(notes).strip() if notes else ''
    })

entries_with_url = [e for e in entries if e['url']]
print(f"Total entries to audit: {len(entries_with_url)}")

def audit_single_job(item):
    url = item['url']
    comp = item['company']
    title = item['title']
    
    status_code = None
    final_url = url
    is_active = False
    reason = ""
    extracted_text = ""
    
    try:
        req = urllib.request.Request(url, headers=headers)
        with urllib.request.urlopen(req, context=ctx, timeout=12) as resp:
            status_code = resp.getcode()
            final_url = resp.geturl()
            html_bytes = resp.read()
            soup = BeautifulSoup(html_bytes, 'html.parser')
            page_text = soup.get_text(' ', strip=True).lower()
            extracted_text = page_text[:1200]

            closed_markers = [
                'no longer accepting applications',
                'this job is no longer available',
                'this position has been filled',
                'job has been closed',
                'job not found',
                'page not found',
                'requisition has closed',
                'no longer open',
                'job is no longer active',
                'not accepting new applicants',
                '404 not found',
                'oops! this job is no longer available',
                'looking for a new job? this one is no longer available'
            ]

            found_closed = any(m in page_text for m in closed_markers)
            
            if 'amazon.jobs' in url:
                if 'not found' in page_text or 'job is closed' in page_text or 'no longer open' in page_text:
                    is_active = False
                    reason = "Amazon job posting has closed / not found"
                elif 'apply now' in page_text or 'basic qualifications' in page_text:
                    is_active = True
                    reason = "Active Amazon job page with Apply Now button"
                else:
                    is_active = True
                    reason = "HTTP 200 response"
            elif 'greenhouse.io' in url:
                if 'no longer available' in page_text or 'job has closed' in page_text:
                    is_active = False
                    reason = "Greenhouse posting closed / expired"
                else:
                    is_active = True
                    reason = "Active Greenhouse ATS application form"
            elif 'ashbyhq.com' in url:
                if 'no longer available' in page_text or 'closed' in page_text:
                    is_active = False
                    reason = "Ashby posting closed / expired"
                else:
                    is_active = True
                    reason = "Active Ashby ATS requisition"
            elif 'linkedin.com' in url:
                if 'no longer accepting applications' in page_text:
                    is_active = False
                    reason = "LinkedIn listing is closed (no longer accepting applications)"
                else:
                    is_active = True
                    reason = "Active LinkedIn job listing"
            elif 'myworkdayjobs.com' in url:
                if 'this job is no longer available' in page_text or 'closed' in page_text:
                    is_active = False
                    reason = "Workday posting closed / inactive"
                else:
                    is_active = True
                    reason = "Active Workday requisition"
            else:
                if found_closed:
                    is_active = False
                    reason = "Closed/expired notice detected on page"
                else:
                    is_active = True
                    reason = "HTTP 200 Active page"

    except urllib.error.HTTPError as e:
        status_code = e.code
        if e.code == 404:
            is_active = False
            reason = "HTTP 404: Dead Link (Job expired or removed)"
        elif e.code == 410:
            is_active = False
            reason = "HTTP 410: Gone (Job closed)"
        elif e.code == 403:
            # Let's mark as gated bot protection (Workday or Cloudflare)
            status_code = 403
            is_active = True # needs browser test
            reason = "HTTP 403: Cloudflare / Bot Protection (Gated ATS)"
        else:
            is_active = False
            reason = f"HTTP Error {e.code}"
    except Exception as e:
        status_code = "ERR"
        is_active = False
        reason = f"Network error: {e}"

    return {
        'row': item['row'],
        'company': comp,
        'title': title,
        'location': item['location'],
        'job_id': item['job_id'],
        'url': url,
        'final_url': final_url,
        'http_status': status_code,
        'is_active': is_active,
        'reason': reason,
        'page_snippet': extracted_text[:300],
        'notes': item['notes']
    }

results = []
start_t = time.time()

with ThreadPoolExecutor(max_workers=12) as executor:
    future_to_item = {executor.submit(audit_single_job, item): item for item in entries_with_url}
    for future in as_completed(future_to_item):
        res = future.result()
        results.append(res)
        status_str = "ACTIVE" if res['is_active'] else "INACTIVE/DEAD"
        print(f"[{status_str}] R{res['row']:03d} | {res['company']} | {res['http_status']} | {res['reason'][:50]}")

print(f"\nAudit completed in {time.time() - start_t:.2f} seconds!")
results.sort(key=lambda x: x['row'])

with open('data/audit_companies_links_final.json', 'w') as f:
    json.dump(results, f, indent=2)

print(f"Saved {len(results)} audited results to data/audit_companies_links_final.json")
