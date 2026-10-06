import json
import re
import time
import ssl
import urllib.request
from bs4 import BeautifulSoup
from selenium import webdriver
from selenium.webdriver.chrome.options import Options
import openpyxl

# 1. Load existing companies from Prev Waves, Wave 8, and Wave 9
wb = openpyxl.load_workbook('data/Job Search.xlsx', data_only=True)
existing_companies = set()
for sname in ['Prev Waves', 'Wave 8', 'Wave 9']:
    if sname in wb.sheetnames:
        ws = wb[sname]
        comp_col = 3 if sname == 'Prev Waves' else 2
        for r in range(2, ws.max_row + 1):
            val = ws.cell(r, comp_col).value
            if val and isinstance(val, str):
                existing_companies.add(val.strip().lower())

print(f"Loaded {len(existing_companies)} existing companies to avoid unless fresh grad specific.")

# 2. Load seen URLs
seen_urls = set()
with open('data/all_existing_urls.json') as f:
    for u in json.load(f):
        seen_urls.add(u.strip().lower().rstrip('/'))

for sname in wb.sheetnames:
    ws = wb[sname]
    for row in ws.iter_rows(values_only=True):
        for cell in row:
            if cell and isinstance(cell, str) and cell.startswith('http'):
                seen_urls.add(cell.strip().lower().rstrip('/'))

# 3. Load harvested cards
with open('data/scratch_wave10_harvested_cards.json') as f:
    cards = json.load(f)

# 4. Load ATS jobs
with open('data/scratch_raw_comprehensive_ats.json') as f:
    ats_raw = json.load(f)

# Paused companies
paused_companies = {'amazon', 'sarvam', 'sarvam ai', 'mongodb'}

# Reject title patterns (internships, non-tech, senior roles)
reject_title_words = [
    'intern', 'internship', 'summer', 'winter', 'co-op', 'coop',
    'lead', 'senior', 'sr.', 'sr ', 'manager', 'principal', 'director', 'staff',
    'ii', 'iii', 'iv', '2', '3', '4', '5',
    'recruiter', 'talent', 'hr', 'sales', 'marketing', 'business development',
    'operations', 'telecaller', 'customer support', 'bpo', 'civil', 'mechanical site',
    'chemical', 'metallurgy', 'hardware', 'assembly', 'procurement'
]

# Required tech keywords in title
tech_title_words = [
    'software', 'engineer', 'developer', 'backend', 'data', 'python',
    'machine learning', 'ml', 'ai', 'trainee', 'get', 'sde', 'qa', 'sdet',
    'full stack', 'frontend', 'cloud', 'systems', 'platform', 'analytics'
]

fresh_grad_markers = [
    'graduate engineer trainee', 'get', 'campus', '2026', 'university graduate',
    'fresh grad', 'fresher', 'early career', 'trainee', 'graduate software engineer',
    'associate software engineer', 'junior software developer'
]

shortlist = []

# Process ATS roles
for j in ats_raw:
    comp = j.get('company', '').strip()
    title = j.get('title', '').strip()
    loc = j.get('location', '').strip()
    url = j.get('url', '').strip()
    clean_url = url.lower().rstrip('/')

    if clean_url in seen_urls:
        continue

    c_lower = comp.lower()
    t_lower = title.lower()
    loc_lower = loc.lower()

    if any(p in c_lower for p in paused_companies):
        continue

    if not any(token in loc_lower for token in ['india', 'bengaluru', 'bangalore', 'gurgaon', 'gurugram', 'noida', 'delhi', 'hyderabad', 'pune', 'mumbai']):
        continue

    if any(w in t_lower for w in reject_title_words):
        continue

    if not any(w in t_lower for w in tech_title_words):
        continue

    is_fresh_grad_role = any(m in t_lower for m in fresh_grad_markers)
    is_existing = any(ec in c_lower or c_lower in ec for ec in existing_companies)

    if is_existing and not is_fresh_grad_role:
        continue

    shortlist.append({
        'company': comp,
        'title': title,
        'location': loc,
        'url': url,
        'source': f"Direct ATS ({j.get('platform')})",
        'is_ats': True,
        'is_existing': is_existing,
        'is_fresh_grad': is_fresh_grad_role
    })

# Process LinkedIn harvested cards
for c in cards:
    comp = c.get('company', '').strip()
    title = c.get('title', '').strip()
    loc = c.get('location', '').strip()
    url = c.get('url', '').strip()
    clean_url = url.lower().rstrip('/')

    if clean_url in seen_urls:
        continue

    c_lower = comp.lower()
    t_lower = title.lower()
    loc_lower = loc.lower()

    if any(p in c_lower for p in paused_companies):
        continue

    if any(w in t_lower for w in reject_title_words):
        continue

    if not any(w in t_lower for w in tech_title_words):
        continue

    is_fresh_grad_role = any(m in t_lower for m in fresh_grad_markers)
    is_existing = any(ec in c_lower or c_lower in ec for ec in existing_companies)

    if is_existing and not is_fresh_grad_role:
        continue

    priority = 0
    # Prioritize completely new companies
    if not is_existing:
        priority += 10
    # Prioritize fresh grad specific roles
    if is_fresh_grad_role:
        priority += 8
    # Prioritize NCR locations
    if any(ncr in loc_lower for ncr in ['noida', 'gurgaon', 'gurugram', 'delhi']):
        priority += 5
    # Prioritize Python/Backend/Data
    if any(k in t_lower for k in ['python', 'backend', 'data', 'ml', 'ai']):
        priority += 4

    c['priority'] = priority
    c['is_ats'] = False
    c['source'] = 'Public Requisition (LinkedIn Guest)'
    c['is_existing'] = is_existing
    c['is_fresh_grad'] = is_fresh_grad_role
    shortlist.append(c)

shortlist.sort(key=lambda x: x.get('priority', 0), reverse=True)
print(f"Total shortlisted candidates for deep inspection: {len(shortlist)}")
with open('data/scratch_wave10_shortlist.json', 'w') as f:
    json.dump(shortlist[:90], f, indent=2)

print("Saved top 90 to data/scratch_wave10_shortlist.json")
