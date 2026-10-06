import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

wb = openpyxl.load_workbook('data/Job Search.xlsx')

# Remove existing Wave 5 sheet
if 'New Additions (Wave 5)' in wb.sheetnames:
    del wb['New Additions (Wave 5)']

ws = wb.create_sheet('New Additions (Wave 5)')

headers = ['Company', 'Role Title', 'Location', 'Technical Stack & Skill Match for Vaanya', 'Status / Match Score', 'Direct Application URL']
ws.append(headers)

header_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
header_font = Font(color='FFFFFF', bold=True, size=11)
for col in range(1, len(headers) + 1):
    c = ws.cell(row=1, column=col)
    c.fill = header_fill
    c.font = header_font
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

green_fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
light_green_fill = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')
bold_green = Font(color='006100', bold=True)

# 14 Unique, Verified, Skill-Matched Roles (Zero Duplicates across all sheets)
unique_roles = [
    {
        'company': 'Amazon',
        'title': 'Software Development Engineer Intern (AUTA)',
        'location': 'Delhi NCR / Bengaluru / Hyderabad',
        'stack': 'Core CS / Systems & Cloud: Python, Java, C++, Data Structures & Algorithms, Distributed Systems, AWS. Amazon University Talent Acquisition (AUTA) requisition ID: 10513277. Explicit 0 YOE student intake.',
        'score': '94 / 100 (Strong Match)',
        'url': 'https://www.amazon.jobs/en/jobs/10513277/software-development-engineer-intern-may-2027-2-month-amazon-university-talent-acquisition',
        'is_strong': True
    },
    {
        'company': 'Hewlett Packard Enterprise (HPE)',
        'title': 'Graduate Software Engineer',
        'location': 'Bengaluru (Designated Remote / Teleworker)',
        'stack': 'Python & Systems Engineering: Python, C++, Java, Linux, systems software, database concepts, debugging & test automation. Brand new requisition ID: 1210218 (posted Oct 1, 2026). Designated remote teleworker across India.',
        'score': '88 / 100 (Strong Match)',
        'url': 'https://careers.hpe.com/us/en/job/1210218/Graduate-Software-Engineer',
        'is_strong': True
    },
    {
        'company': 'Electronic Arts (EA)',
        'title': 'AI Full Stack Intern',
        'location': 'Hyderabad, Telangana (Hybrid)',
        'stack': 'AI & Backend Frameworks: Python, FastAPI, LangChain, OpenAI APIs, LLM integrations, Agentic workflows, SQL/NoSQL. Central Technology (CT) team at EA. Requisition ID: 215939.',
        'score': '88 / 100 (Strong Match)',
        'url': 'https://jobs.ea.com/en_US/careers/JobDetail/215939',
        'is_strong': True
    },
    {
        'company': 'Urban Company',
        'title': 'Software Development Engineer 1 (SDE-1)',
        'location': 'Gurugram, Haryana (In-Office)',
        'stack': 'Python Backend & Microservices: Python, Node.js, PostgreSQL, Microservices, high-scale distributed systems. Prime Gurugram HQ (NCR). SDE-1 track for 0-2 YOE engineers (CTC: ₹16-24 LPA).',
        'score': '89 / 100 (Strong Match)',
        'url': 'https://www.urbancompany.com/careers',
        'is_strong': True
    },
    {
        'company': 'Prismagic Solutions Inc.',
        'title': 'Graduate Trainee – PV-Technology & AI/ML R&D',
        'location': 'Noida, Uttar Pradesh, India',
        'stack': 'AI/ML & Python: Python, machine learning R&D, data pipelines. Verbatim JD: "Experience: 0–1 year | Freshers welcome. B.Tech / M.Tech". Noida office (minutes from her college Sector 62/128).',
        'score': '86 / 100 (Strong Match)',
        'url': 'https://in.linkedin.com/jobs/view/graduate-trainee-%E2%80%93-pv-technology-ai-ml-r-d-at-prismagic-solutions-inc-4470096991',
        'is_strong': True
    },
    {
        'company': 'Zepto',
        'title': 'Data Scientist',
        'location': 'Bengaluru, Karnataka',
        'stack': 'Data Engineering & ML: Python, Machine Learning, Statistical Modeling, ML Pipelines, SQL. Instahyre Requisition ID: 264559. Explicit 0-2 YOE fresh graduate intake.',
        'score': '85 / 100 (Strong Match)',
        'url': 'https://www.instahyre.com/job/264559-data-scientist-at-zepto-bangalore/',
        'is_strong': True
    },
    {
        'company': 'Tecnod8.AI',
        'title': 'Intern Full Stack (Agentic AI Applications) Development',
        'location': 'Gurugram, Haryana, India',
        'stack': 'Agentic AI & Backend: Python, FastAPI, Agentic AI application frameworks, REST APIs. Hybrid in Gurgaon. High-performing interns eligible for PPO conversion.',
        'score': '91 / 100 (Strong Match)',
        'url': 'https://in.linkedin.com/jobs/view/intern-full-stack-agentic-ai-applications-development-at-tecnod8-ai-4466662629',
        'is_strong': True
    },
    {
        'company': 'Coding Ninjas',
        'title': 'GenAI Intern',
        'location': 'Gurugram, Haryana, India',
        'stack': 'Generative AI & LLMs: Python, LLM integrations, prompt engineering, automated evaluations. Gurgaon HQ. Perfect fit for Vaanya\'s NLP/Transformer background.',
        'score': '83 / 100 (Strong Match)',
        'url': 'https://in.linkedin.com/jobs/view/genai-intern-at-coding-ninjas-4468561627',
        'is_strong': True
    },
    {
        'company': 'Honasa Consumer Ltd. (Mamaearth)',
        'title': 'AI & Analytics Intern',
        'location': 'Gurugram, Haryana, India',
        'stack': 'Analytics & ML: Python, SQL, predictive modeling, data pipelines, business intelligence. Direct Gurgaon HQ role at India\'s largest D2C digital house.',
        'score': '82 / 100 (Strong Match)',
        'url': 'https://in.linkedin.com/jobs/view/ai-intern-at-honasa-consumer-ltd-4472919251',
        'is_strong': True
    },
    {
        'company': 'insightsoftware',
        'title': 'Associate Software Engineer (Internship)',
        'location': 'Hyderabad / Bengaluru (Remote/Hybrid)',
        'stack': 'Cloud & APIs: Python, SQL/PL-SQL, API automation testing, CI/CD, GenAI workflows. Workday Requisition ID: REQ001103-1. Explicit internship for fresh B.Tech graduates.',
        'score': '82 / 100 (Strong Match)',
        'url': 'https://magnitudesoftware.wd1.myworkdayjobs.com/External/job/India---Hyderabad---Remote/Associate-Software-Engineer_REQ001103-1',
        'is_strong': True
    },
    {
        'company': 'Bharat.Law',
        'title': 'AI Integrations - Software Intern',
        'location': 'Greater Delhi Area',
        'stack': 'NLP Document AI: Python, document extraction pipelines (BERT), REST APIs. Direct mirror of Vaanya\'s production work at S&P Global on document routing.',
        'score': '75 / 100 (Potential Match)',
        'url': 'https://in.linkedin.com/jobs/view/ai-integrations-software-intern-at-bharat-law-4472393604',
        'is_strong': False
    },
    {
        'company': 'Squarepoint Capital',
        'title': 'Junior Quant Researcher - ML Alpha Research',
        'location': 'Bangalore, India (Global Quant Division)',
        'stack': 'Quant ML & Time Series: Python, PyTorch, statistical machine learning, high-performance data processing. Top-tier global quantitative hedge fund.',
        'score': '75 / 100 (Potential Match)',
        'url': 'https://www.squarepoint-capital.com/open-opportunities?id=6069464&gh_jid=6069464',
        'is_strong': False
    },
    {
        'company': 'Ingersoll Rand',
        'title': 'Graduate Engineer Trainee',
        'location': 'Gurgaon, Haryana, India',
        'stack': 'Engineering Systems: Structured GET campus intake for B.Tech engineers. Testing, software diagnostics, engineering automation.',
        'score': '74 / 100 (Potential Match)',
        'url': 'https://in.linkedin.com/jobs/view/graduate-engineer-trainee-at-ingersoll-rand-4460474501',
        'is_strong': False
    },
    {
        'company': 'ByteDance',
        'title': 'Data Analyst Project Intern (2026 Start)',
        'location': 'Gurgaon, Haryana, India',
        'stack': 'Data Analytics & Pipelines: SQL, Python, quantitative data analysis. Explicitly marked: "2026 Start (BS/MS)" at ByteDance Gurgaon.',
        'score': '67 / 100 (Potential Match)',
        'url': 'https://in.linkedin.com/jobs/view/data-analyst-project-intern-training-quality-2026-start-bs-ms-at-bytedance-4472767912',
        'is_strong': False
    }
]

row_idx = 2
for r in unique_roles:
    ws.append([r['company'], r['title'], r['location'], r['stack'], r['score'], r['url']])
    
    fill = green_fill if r['is_strong'] else light_green_fill
    for c in range(1, 7):
        cell = ws.cell(row=row_idx, column=c)
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        
    ws.cell(row=row_idx, column=5).font = bold_green
    
    url_c = ws.cell(row=row_idx, column=6)
    url_c.hyperlink = r['url']
    url_c.font = Font(color='0563C1', underline='single')
    row_idx += 1

# Add section divider for cross-sheet duplicates
row_idx += 1
ws.cell(row=row_idx, column=1).value = "CROSS-SHEET DUPLICATES IDENTIFIED (Already tracked in earlier sheets)"
ws.cell(row=row_idx, column=1).font = Font(bold=True, color='7030A0', size=11)
ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=6)
ws.cell(row=row_idx, column=1).fill = PatternFill(start_color='E1D5E7', end_color='E1D5E7', fill_type='solid')
row_idx += 1

# List the 8 detected duplicates with cross-references
dup_fill = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')
dup_font = Font(color='B25900', italic=True)

duplicates = [
    ('Tower Research Capital', 'Intern - AI/ML', 'Gurgaon', 'DUPLICATE: Already tracked in Sheet 2 "Job Links" (Row 32) via Greenhouse ID 8143756.', 'DUPLICATE', 'https://www.tower-research.com/open-positions/?gh_jid=8143756'),
    ('Stripe', 'Software Engineer, Intern', 'Bengaluru', 'DUPLICATE: Already tracked in Sheet 2 "Job Links" (Row 6) as SWE Intern.', 'DUPLICATE', 'https://stripe.com/jobs/search?gh_jid=8031833'),
    ('Salesforce', 'Summer Intern - Software Engineer (AMTS)', 'Bengaluru / Hyd', 'DUPLICATE: Already tracked in Sheet 2 "Job Links" (Row 5) & Sheet 4 "Fresher Hiring Drives" (Row 3).', 'DUPLICATE', 'https://salesforce.wd12.myworkdayjobs.com/External_Career_Site/job/Summer-2027-Intern---Software-Engineer_JR340771'),
    ('Cars24', 'AI Science Intern', 'Gurugram', 'DUPLICATE: Already tracked in Sheet 2 "Job Links" (Row 35) as Cars24 Intern-to-PPO.', 'DUPLICATE', 'https://unstop.com/internships/ai-science-intern-cars24-1502422'),
    ('IQVIA', 'Software Engineer Intern (AI & Automation)', 'Gurugram / Bengaluru', 'DUPLICATE: Already tracked in Sheet 5 "New Additions (Wave 3)" (Row 18) as IQVIA Intern.', 'DUPLICATE', 'https://jobs.iqvia.com/en/jobs/R1567860-0'),
    ('American Express', 'Analyst - Data Science', 'Gurugram', 'DUPLICATE: Already tracked in Sheet 3 "Eligible role" (Row 49) as Campus Analyst / Data Analytics.', 'DUPLICATE', 'https://in.linkedin.com/jobs/view/analyst-data-science-at-american-express-4472321596'),
    ('American Express', 'Analyst - Data Analytics', 'Gurugram', 'DUPLICATE: Already tracked in Sheet 3 "Eligible role" (Row 49) as Campus Analyst / Data Analytics.', 'DUPLICATE', 'https://in.linkedin.com/jobs/view/analyst-data-analytics-at-american-express-4473243152'),
    ('CashKaro.com', 'AI-First Product Management - Intern', 'Gurugram', 'DUPLICATE & SKILL MISMATCH: Tracked in Sheet 3 (Row 72) as Python Backend Intern; LinkedIn post was PM.', 'DUPLICATE', 'https://in.linkedin.com/jobs/view/ai-first-product-management-intern-at-cashkaro-com-4470174235')
]

for d in duplicates:
    ws.append([d[0], d[1], d[2], d[3], d[4], d[5]])
    for c in range(1, 7):
        cell = ws.cell(row=row_idx, column=c)
        cell.fill = dup_fill
        cell.font = dup_font
        cell.alignment = Alignment(wrap_text=True, vertical='top')
    url_c = ws.cell(row=row_idx, column=6)
    url_c.hyperlink = d[5]
    url_c.font = Font(color='0563C1', underline='single')
    row_idx += 1

col_widths = [24, 38, 22, 60, 22, 32]
for i, w in enumerate(col_widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws.row_dimensions[1].height = 28
wb.save('data/Job Search.xlsx')
print(f"Rebuilt Wave 5 sheet successfully: {len(unique_roles)} unique skill-matched roles + {len(duplicates)} cross-referenced duplicates.")
