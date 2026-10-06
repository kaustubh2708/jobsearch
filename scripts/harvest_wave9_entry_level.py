import urllib.parse
import time
import json
import re
from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from bs4 import BeautifulSoup
import openpyxl

# Load seen URLs
seen_urls = set()
with open('data/all_existing_urls.json') as f:
    for u in json.load(f):
        seen_urls.add(u.strip().lower().rstrip('/'))

wb = openpyxl.load_workbook('data/Job Search.xlsx', data_only=True)
for sname in wb.sheetnames:
    ws = wb[sname]
    for row in ws.iter_rows(values_only=True):
        for cell in row:
            if cell and isinstance(cell, str) and cell.startswith('http'):
                seen_urls.add(cell.strip().lower().rstrip('/'))

print(f"Loaded {len(seen_urls)} unique existing URLs for deduplication.")

options = Options()
options.add_argument('--headless=new')
options.add_argument('--no-sandbox')
options.add_argument('--disable-dev-shm-usage')
options.add_argument('--user-agent=Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36')

searches = [
    # Noida & Gurgaon Priority (NCR)
    {'kw': 'Software Engineer I', 'loc': 'Noida, Uttar Pradesh, India', 'extra': '&f_E=2&f_JT=F'},
    {'kw': 'Python Backend Engineer', 'loc': 'Noida, Uttar Pradesh, India', 'extra': '&f_E=2&f_JT=F'},
    {'kw': 'Graduate Engineer Trainee', 'loc': 'Noida, Uttar Pradesh, India', 'extra': '&f_JT=F'},
    {'kw': 'Software Development Engineer I', 'loc': 'Gurgaon, Haryana, India', 'extra': '&f_E=2&f_JT=F'},
    {'kw': 'Associate Software Engineer', 'loc': 'Gurgaon, Haryana, India', 'extra': '&f_E=2&f_JT=F'},
    {'kw': 'Junior Python Developer', 'loc': 'Gurgaon, Haryana, India', 'extra': '&f_E=2&f_JT=F'},
    {'kw': 'Junior Data Engineer', 'loc': 'Gurgaon, Haryana, India', 'extra': '&f_E=2&f_JT=F'},
    {'kw': 'Graduate Engineer Trainee', 'loc': 'Gurgaon, Haryana, India', 'extra': '&f_JT=F'},
    {'kw': 'Associate Machine Learning Engineer', 'loc': 'Gurgaon, Haryana, India', 'extra': '&f_E=2&f_JT=F'},
    {'kw': 'Associate Software Engineer', 'loc': 'Noida, Uttar Pradesh, India', 'extra': '&f_E=2&f_JT=F'},
    # Delhi & NCR
    {'kw': 'Junior Software Engineer', 'loc': 'Delhi, India', 'extra': '&f_E=2&f_JT=F'},
    {'kw': 'Software Engineer I', 'loc': 'Delhi, India', 'extra': '&f_E=2&f_JT=F'},
    {'kw': 'Associate Data Engineer', 'loc': 'Delhi, India', 'extra': '&f_E=2&f_JT=F'},
    # Remote & Bengaluru Top Tech Entry Level
    {'kw': 'Software Engineer I', 'loc': 'Remote', 'extra': '&f_WT=2&f_E=2&f_JT=F'},
    {'kw': 'Graduate Software Engineer', 'loc': 'Bengaluru, Karnataka, India', 'extra': '&f_E=2&f_JT=F'},
    {'kw': 'Software Development Engineer I', 'loc': 'Bengaluru, Karnataka, India', 'extra': '&f_E=2&f_JT=F'},
    {'kw': 'Associate Data Engineer', 'loc': 'Bengaluru, Karnataka, India', 'extra': '&f_E=2&f_JT=F'},
    {'kw': 'Junior Machine Learning Engineer', 'loc': 'Bengaluru, Karnataka, India', 'extra': '&f_E=2&f_JT=F'},
    {'kw': 'Associate Software Engineer', 'loc': 'Bengaluru, Karnataka, India', 'extra': '&f_E=2&f_JT=F'}
]

reject_title_words = [
    'intern', 'internship', 'summer', 'winter', 'lead', 'senior', 'sr.', 'sr ', 'manager',
    'principal', 'director', 'staff', 'ii', 'iii', '2', '3', '4', 'recruiter', 'hr', 'sales',
    'marketing', 'business development', 'product management', 'operations', 'telecaller',
    'civil', 'mechanical', 'support representative'
]

paused_companies = ['amazon', 'sarvam', 'sarvam ai', 'mongodb']

driver = webdriver.Chrome(options=options)
raw_cards = {}

try:
    for idx, s in enumerate(searches):
        kw = s['kw']
        loc = s['loc']
        extra = s.get('extra', '')
        url = f"https://www.linkedin.com/jobs/search?keywords={urllib.parse.quote(kw)}&location={urllib.parse.quote(loc)}{extra}"
        print(f"[{idx+1}/{len(searches)}] Searching: '{kw}' in '{loc}'...")
        try:
            driver.get(url)
            time.sleep(1.8)
            soup = BeautifulSoup(driver.page_source, 'html.parser')
            cards = soup.find_all('div', class_='base-card')
            found_count = 0
            for c in cards:
                t = c.find('h3', class_='base-search-card__title')
                comp = c.find('h4', class_='base-search-card__subtitle')
                l = c.find('span', class_='job-search-card__location')
                a = c.find('a', class_='base-card__full-link')

                title = t.get_text(strip=True) if t else ''
                company = comp.get_text(strip=True) if comp else ''
                location = l.get_text(strip=True) if l else ''
                link = a['href'] if a and a.has_attr('href') else ''

                if not link:
                    continue

                clean_url = link.split('?')[0]
                t_lower = title.lower()
                c_lower = company.lower()

                # Check paused companies
                if any(p in c_lower for p in paused_companies):
                    continue

                # Exclude reject words in title
                if any(w in t_lower for w in reject_title_words):
                    continue

                if clean_url.lower().rstrip('/') in seen_urls:
                    continue

                m = re.search(r'-([0-9]{8,12})(?:$)', clean_url)
                jid = m.group(1) if m else clean_url

                if jid not in raw_cards:
                    raw_cards[jid] = {
                        'company': company,
                        'title': title,
                        'location': location,
                        'url': clean_url,
                        'jid': jid
                    }
                    found_count += 1
            print(f"   -> Added {found_count} new cards (Total unique: {len(raw_cards)})")
        except Exception as e:
            print(f"   -> Search error: {e}")
        time.sleep(0.5)

    print(f"\nDiscovered {len(raw_cards)} potential full-time entry-level candidates.")
    with open('data/scratch_harvested_cards.json', 'w') as f:
        json.dump(list(raw_cards.values()), f, indent=2)
finally:
    driver.quit()
