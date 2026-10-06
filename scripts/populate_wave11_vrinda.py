import openpyxl, json, shutil
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

print("Starting workbook reorganization for Wave 11...")

wb = openpyxl.load_workbook('data/Companies(1).xlsx')
print("Current sheet names:", wb.sheetnames)

ws_past = wb['past wave']
has_w9 = 'Wave 9' in wb.sheetnames
if has_w9:
    ws_w9 = wb['Wave 9']

# Thin border & styles
thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)
link_font = Font(name='Calibri', size=11, color='000563C1', underline='single')
verified_fill = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')

curr_past_row = ws_past.max_row
print(f"Current past wave max row: {curr_past_row}")

# Append Wave 9 to past wave if present
if has_w9:
    for r in range(2, ws_w9.max_row + 1):
        curr_past_row += 1
        seq_num = curr_past_row - 1
        
        comp = ws_w9.cell(row=r, column=2).value
        title = ws_w9.cell(row=r, column=3).value
        loc = ws_w9.cell(row=r, column=4).value
        exp = ws_w9.cell(row=r, column=5).value
        tech = ws_w9.cell(row=r, column=6).value
        sal = ws_w9.cell(row=r, column=7).value
        score = ws_w9.cell(row=r, column=8).value
        tier = ws_w9.cell(row=r, column=9).value
        status = ws_w9.cell(row=r, column=10).value
        
        url_cell = ws_w9.cell(row=r, column=11)
        target_url = url_cell.hyperlink.target if url_cell.hyperlink else url_cell.value
        url_text = url_cell.value or "Exact Requisition Link ↗"
        
        notes = ws_w9.cell(row=r, column=12).value

        # Write to past wave
        ws_past.cell(row=curr_past_row, column=1, value=seq_num)
        ws_past.cell(row=curr_past_row, column=2, value="Wave 9")
        ws_past.cell(row=curr_past_row, column=3, value=comp)
        ws_past.cell(row=curr_past_row, column=4, value=title)
        ws_past.cell(row=curr_past_row, column=5, value=loc)
        ws_past.cell(row=curr_past_row, column=6, value=exp)
        ws_past.cell(row=curr_past_row, column=7, value=tech)
        ws_past.cell(row=curr_past_row, column=8, value=sal)
        ws_past.cell(row=curr_past_row, column=9, value=score)
        ws_past.cell(row=curr_past_row, column=10, value=tier)
        ws_past.cell(row=curr_past_row, column=11, value=status)
        
        c_link = ws_past.cell(row=curr_past_row, column=12, value=url_text)
        if target_url:
            c_link.hyperlink = target_url
            c_link.font = link_font
            
        ws_past.cell(row=curr_past_row, column=13, value=notes)

        # Apply styling
        bg_color = 'FFFFFF' if (curr_past_row % 2 == 0) else 'F9FAFB'
        row_fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type='solid')

        for col in range(1, 14):
            c = ws_past.cell(row=curr_past_row, column=col)
            c.border = thin_border
            if col == 11 and status and 'Verified' in str(status):
                c.fill = verified_fill
            elif col != 12 or not target_url:
                c.fill = row_fill

            if col in [1, 2, 9, 10]:
                c.alignment = Alignment(horizontal='center', vertical='center')
            elif col in [11, 12]:
                c.alignment = Alignment(horizontal='center', vertical='center')
            else:
                c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

    print(f"Appended Wave 9 to past wave. New past wave max row: {ws_past.max_row} (Total data rows: {ws_past.max_row - 1})")
    wb.remove(ws_w9)
    print("Removed Wave 9 sheet from workbook.")
else:
    print(f"Wave 9 already merged. past wave current max row: {ws_past.max_row}")

# Create or recreate Wave 11 sheet
if 'Wave 11' in wb.sheetnames:
    wb.remove(wb['Wave 11'])

ws_w11 = wb.create_sheet(title='Wave 11')

headers = [
    '#',
    'Company',
    'Role Title',
    'Location (live)',
    'Experience (per JD)',
    'Key Tech Stack (estimate)',
    'Estimated Compensation (estimate)',
    'Match Score',
    'Priority / Match Tier',
    'Link Verification Status',
    'Direct Application URL',
    'Notes'
]

header_fill = PatternFill(start_color='1F497D', end_color='1F497D', fill_type='solid')
header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')

for col_idx, h in enumerate(headers, start=1):
    cell = ws_w11.cell(row=1, column=col_idx, value=h)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = thin_border

ws_w11.row_dimensions[1].height = 28

col_widths = {
    'A': 12.0,
    'B': 25.0,
    'C': 50.0,
    'D': 35.0,
    'E': 22.0,
    'F': 55.0,
    'G': 55.0,
    'H': 14.0,
    'I': 24.0,
    'J': 37.0,
    'K': 27.0,
    'L': 55.0
}
for col_l, w in col_widths.items():
    ws_w11.column_dimensions[col_l].width = w

# Populate Wave 11 rows
with open('data/wave11_bulletproof_roles.json') as f:
    wave11_roles = json.load(f)

print(f"Populating Wave 11 with {len(wave11_roles)} verified roles...")

for idx, r in enumerate(wave11_roles, start=1):
    row_num = idx + 1
    
    comp = r['company']
    title = r['title']
    loc = r['location']
    exp = r['experience_per_jd']
    tech = r['tech_stack']
    comp_est = r['comp']
    score = r['score']
    tier = r['tier']
    status_label = r.get('status_label', 'Verified: live LinkedIn job match')
    url = r['url']
    notes = r['notes']
    
    ws_w11.cell(row=row_num, column=1, value=idx)
    ws_w11.cell(row=row_num, column=2, value=comp)
    ws_w11.cell(row=row_num, column=3, value=title)
    ws_w11.cell(row=row_num, column=4, value=loc)
    ws_w11.cell(row=row_num, column=5, value=exp)
    ws_w11.cell(row=row_num, column=6, value=tech)
    ws_w11.cell(row=row_num, column=7, value=comp_est)
    ws_w11.cell(row=row_num, column=8, value=score)
    ws_w11.cell(row=row_num, column=9, value=tier)
    ws_w11.cell(row=row_num, column=10, value=status_label)
    
    c_url = ws_w11.cell(row=row_num, column=11, value="Exact Requisition Link ↗")
    c_url.hyperlink = url
    c_url.font = link_font
    
    ws_w11.cell(row=row_num, column=12, value=notes)

    bg_hex = 'FFFFFF' if (idx % 2 != 0) else 'F9FAFB'
    row_fill = PatternFill(start_color=bg_hex, end_color=bg_hex, fill_type='solid')

    for c_i in range(1, 13):
        c = ws_w11.cell(row=row_num, column=c_i)
        c.border = thin_border
        if c_i == 10:
            c.fill = verified_fill
        else:
            c.fill = row_fill

        if c_i in [1, 8, 9, 10, 11]:
            c.alignment = Alignment(horizontal='center', vertical='center')
        else:
            c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
            
    ws_w11.row_dimensions[row_num].height = 22

ws_w11.freeze_panes = 'A2'

wb.save('data/Companies(1).xlsx')
print("Successfully saved data/Companies(1).xlsx")

shutil.copyfile('data/Companies(1).xlsx', 'data/companies.xlsx')
print("Successfully mirrored to data/companies.xlsx")

# Update seen URLs and IDs index
with open('data/vrinda_all_seen_urls.json') as f:
    seen_idx = json.load(f)

seen_urls_set = set(seen_idx.get('seen_urls', []))
seen_ids_set = set(seen_idx.get('seen_ids', []))

for r in wave11_roles:
    u = r['url'].split('?')[0].lower().rstrip('/')
    seen_urls_set.add(u)
    if r.get('id'):
        seen_ids_set.add(str(r['id']))

seen_idx['seen_urls'] = sorted(list(seen_urls_set))
seen_idx['seen_ids'] = sorted(list(seen_ids_set))
seen_idx['total_urls'] = len(seen_idx['seen_urls'])
seen_idx['total_ids'] = len(seen_idx['seen_ids'])

with open('data/vrinda_all_seen_urls.json', 'w') as f:
    json.dump(seen_idx, f, indent=2)

print(f"Updated data/vrinda_all_seen_urls.json: {seen_idx['total_urls']} seen URLs, {seen_idx['total_ids']} seen IDs.")
