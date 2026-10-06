import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

wb = openpyxl.load_workbook('data/Job Search.xlsx')

if 'New Additions (Wave 5)' in wb.sheetnames:
    del wb['New Additions (Wave 5)']

ws = wb.create_sheet('New Additions (Wave 5)')

headers = ['Company', 'Role Title', 'Location', 'Employment Type', 'Technical Stack & Skill Match for Vaanya', 'Status / Match Score', 'Direct Application URL']
ws.append(headers)

header_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
header_font = Font(color='FFFFFF', bold=True, size=11)
for col in range(1, len(headers) + 1):
    c = ws.cell(row=1, column=col)
    c.fill = header_fill
    c.font = header_font
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

green_fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
light_green_fill = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')
bold_green = Font(color='006100', bold=True)

# STRICTLY Full-Time (FTE) & Graduate Engineer Trainee (GET) roles ONLY
fte_roles = [
    {
        'company': 'Urban Company',
        'title': 'Software Development Engineer 1 (SDE-1)',
        'location': 'Gurugram, Haryana (In-Office)',
        'type': 'Full-Time (FTE)',
        'stack': 'Python Backend & Microservices: Python, PostgreSQL, Microservices, high-scale distributed systems. Prime Gurugram HQ (NCR). SDE-1 track for 0-2 YOE engineers (Full-time CTC: ₹16-24 LPA). Direct route: careers@urbancompany.com.',
        'score': '89 / 100 (Strong Match)',
        'url': 'https://www.urbancompany.com/careers',
        'is_strong': True
    },
    {
        'company': 'Hewlett Packard Enterprise (HPE)',
        'title': 'Graduate Software Engineer',
        'location': 'Bengaluru (Designated Remote / Teleworker)',
        'type': 'Full-Time (Permanent)',
        'stack': 'Python & Systems Engineering: Python, C++, Java, Linux, systems software, database concepts, debugging & test automation. Brand new requisition ID: 1210218 (posted Oct 1, 2026). Designated remote teleworker across India. Full-time CTC: ₹12-16 LPA.',
        'score': '88 / 100 (Strong Match)',
        'url': 'https://careers.hpe.com/us/en/job/1210218/Graduate-Software-Engineer',
        'is_strong': True
    },
    {
        'company': 'Zepto',
        'title': 'Data Scientist',
        'location': 'Bengaluru, Karnataka',
        'type': 'Full-Time (FTE)',
        'stack': 'Data Engineering & ML: Python, Machine Learning, Statistical Modeling, ML Pipelines, SQL. Instahyre Requisition ID: 264559. Explicit 0-2 YOE fresh graduate intake. Full-time CTC: ₹15-22 LPA.',
        'score': '85 / 100 (Strong Match)',
        'url': 'https://www.instahyre.com/job/264559-data-scientist-at-zepto-bangalore/',
        'is_strong': True
    },
    {
        'company': 'Prismagic Solutions Inc.',
        'title': 'Graduate Trainee – PV-Technology & AI/ML R&D',
        'location': 'Noida, Uttar Pradesh, India',
        'type': 'Full-Time (Permanent Trainee)',
        'stack': 'AI/ML & Python: Python, machine learning R&D, data pipelines. Verbatim JD: "Experience: 0–1 year | Freshers welcome. B.Tech / M.Tech". Full-time permanent technology team in Noida (minutes from her college Sector 62/128).',
        'score': '86 / 100 (Strong Match)',
        'url': 'https://in.linkedin.com/jobs/view/graduate-trainee-%E2%80%93-pv-technology-ai-ml-r-d-at-prismagic-solutions-inc-4470096991',
        'is_strong': True
    },
    {
        'company': 'Squarepoint Capital',
        'title': 'Junior Quant Researcher - ML Alpha Research',
        'location': 'Bangalore, India (Global Quant Division)',
        'type': 'Full-Time (Permanent)',
        'stack': 'Quant ML & Time Series: Python, PyTorch, statistical machine learning, high-performance data processing. Top-tier global quantitative hedge fund. Full-time early career quantitative researcher.',
        'score': '75 / 100 (Potential Match)',
        'url': 'https://www.squarepoint-capital.com/open-opportunities?id=6069464&gh_jid=6069464',
        'is_strong': False
    },
    {
        'company': 'Ingersoll Rand',
        'title': 'Graduate Engineer Trainee',
        'location': 'Gurgaon, Haryana, India',
        'type': 'Full-Time (GET Program)',
        'stack': 'Engineering Systems: Structured GET campus intake for B.Tech engineers. Testing, software diagnostics, engineering automation in Gurgaon.',
        'score': '74 / 100 (Potential Match)',
        'url': 'https://in.linkedin.com/jobs/view/graduate-engineer-trainee-at-ingersoll-rand-4460474501',
        'is_strong': False
    }
]

row_idx = 2
for r in fte_roles:
    ws.append([r['company'], r['title'], r['location'], r['type'], r['stack'], r['score'], r['url']])
    fill = green_fill if r['is_strong'] else light_green_fill
    for c in range(1, 8):
        cell = ws.cell(row=row_idx, column=c)
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        
    ws.cell(row=row_idx, column=6).font = bold_green
    
    url_c = ws.cell(row=row_idx, column=7)
    url_c.hyperlink = r['url']
    url_c.font = Font(color='0563C1', underline='single')
    row_idx += 1

# Excluded Internships Audit Log Section
row_idx += 1
ws.cell(row=row_idx, column=1).value = "INTERNSHIP ROLES EXCLUDED (Per user rule: Full-Time / FTE & GET Only)"
ws.cell(row=row_idx, column=1).font = Font(bold=True, color='C00000', size=11)
ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=7)
ws.cell(row=row_idx, column=1).fill = PatternFill(start_color='FCE4D6', end_color='FCE4D6', fill_type='solid')
row_idx += 1

excluded_interns = [
    ('Amazon', 'Software Development Engineer Intern (AUTA)', 'Delhi NCR / Bengaluru', 'Internship (2-Month)', 'AUTA student internship track (Req ID: 10513277)', 'EXCLUDED_INTERN'),
    ('Electronic Arts (EA)', 'AI Full Stack Intern', 'Hyderabad', 'Internship (3-6 Months)', 'Central Tech AI Intern (Req ID: 215939)', 'EXCLUDED_INTERN'),
    ('Tecnod8.AI', 'Intern Full Stack (Agentic AI)', 'Gurugram', 'Internship (3-6 Months)', 'Full Stack Agentic AI Intern', 'EXCLUDED_INTERN'),
    ('Coding Ninjas', 'GenAI Intern', 'Gurugram', 'Internship', 'EdTech LLM / GenAI Intern', 'EXCLUDED_INTERN'),
    ('Honasa Consumer', 'AI & Analytics Intern', 'Gurugram', 'Internship', 'Mamaearth Analytics & AI Intern', 'EXCLUDED_INTERN'),
    ('insightsoftware', 'Associate Software Engineer (Intern)', 'Hyderabad / Bengaluru', 'Internship', 'Workday Req REQ001103-1', 'EXCLUDED_INTERN'),
    ('Bharat.Law', 'AI Integrations - Software Intern', 'Delhi Area', 'Internship', 'Legal AI Document Intern', 'EXCLUDED_INTERN'),
    ('ByteDance', 'Data Analyst Project Intern', 'Gurgaon', 'Internship', 'ByteDance 2026 Start Intern', 'EXCLUDED_INTERN'),
    ('Stripe', 'Software Engineer, Intern', 'Bengaluru', 'Internship', 'Greenhouse Req 8031833', 'EXCLUDED_INTERN'),
    ('Tower Research Capital', 'Intern - AI/ML', 'Gurgaon', 'Internship (6-Month)', 'Quant AI/ML Intern (Req 8143756)', 'EXCLUDED_INTERN'),
    ('Cars24', 'AI Science Intern', 'Gurugram', 'Internship (3-Month)', 'Unstop Req 1502422', 'EXCLUDED_INTERN'),
    ('Salesforce', 'Summer Intern - Software Engineer', 'Bengaluru / Hyd', 'Internship', 'Futureforce Req JR340771', 'EXCLUDED_INTERN')
]

red_fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
red_font = Font(color='9C0006', italic=True)

for ei in excluded_interns:
    ws.append([ei[0], ei[1], ei[2], ei[3], ei[4], ei[5], ''])
    for c in range(1, 8):
        cell = ws.cell(row=row_idx, column=c)
        cell.fill = red_fill
        cell.font = red_font
        cell.alignment = Alignment(wrap_text=True, vertical='top')
    row_idx += 1

col_widths = [24, 38, 22, 22, 60, 22, 32]
for i, w in enumerate(col_widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws.row_dimensions[1].height = 28
wb.save('data/Job Search.xlsx')
print(f"Updated Wave 5: {len(fte_roles)} Full-Time (FTE) roles + {len(excluded_interns)} excluded internships documented.")
