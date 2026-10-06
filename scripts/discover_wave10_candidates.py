import urllib.parse
import time
import json
import re
import ssl
import urllib.request
import concurrent.futures
from bs4 import BeautifulSoup
from selenium import webdriver
from selenium.webdriver.chrome.options import Options
import openpyxl

# 1. Load existing companies from Prev Waves, Wave 8, and Wave 9
wb = openpyxl.load_workbook('data/Job Search.xlsx', data_only=True)
existing_companies = set()
for sname in ['Prev Waves', 'Wave 8', 'Wave 9']:
    if sname in wb.sheetnames:
        ws = wb[sname]
        comp_col = 3 if sname == 'Prev Waves' else 2
        for r in range(2, ws.max_row + 1):
            val = ws.cell(r, comp_col).value
            if val and isinstance(val, str):
                existing_companies.add(val.strip().lower())

print(f"Loaded {len(existing_companies)} existing companies to avoid (unless fresh grad specific).")

# 2. Load all seen URLs to prevent any duplicate
seen_urls = set()
with open('data/all_existing_urls.json') as f:
    for u in json.load(f):
        seen_urls.add(u.strip().lower().rstrip('/'))

for sname in wb.sheetnames:
    ws = wb[sname]
    for row in ws.iter_rows(values_only=True):
        for cell in row:
            if cell and isinstance(cell, str) and cell.startswith('http'):
                seen_urls.add(cell.strip().lower().rstrip('/'))

print(f"Loaded {len(seen_urls)} unique existing URLs for deduplication.")

# Paused companies per operating rules
paused_companies = {'amazon', 'sarvam', 'sarvam ai', 'mongodb'}

# Setup Selenium for live LinkedIn guest search
options = Options()
options.add_argument('--headless=new')
options.add_argument('--no-sandbox')
options.add_argument('--disable-dev-shm-usage')
options.add_argument('--user-agent=Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36')

searches = [
    # NCR Priority Hubs
    {'kw': 'Graduate Engineer Trainee', 'loc': 'Noida, Uttar Pradesh, India', 'extra': '&f_JT=F'},
    {'kw': 'Software Engineer I', 'loc': 'Noida, Uttar Pradesh, India', 'extra': '&f_E=2&f_JT=F'},
    {'kw': 'Associate Software Engineer', 'loc': 'Noida, Uttar Pradesh, India', 'extra': '&f_E=2&f_JT=F'},
    {'kw': 'Python Developer', 'loc': 'Noida, Uttar Pradesh, India', 'extra': '&f_E=2&f_JT=F'},
    {'kw': 'Junior Data Engineer', 'loc': 'Noida, Uttar Pradesh, India', 'extra': '&f_E=2&f_JT=F'},
    {'kw': 'Graduate Engineer Trainee', 'loc': 'Gurgaon, Haryana, India', 'extra': '&f_JT=F'},
    {'kw': 'Software Development Engineer I', 'loc': 'Gurgaon, Haryana, India', 'extra': '&f_E=2&f_JT=F'},
    {'kw': 'Associate Software Engineer', 'loc': 'Gurgaon, Haryana, India', 'extra': '&f_E=2&f_JT=F'},
    {'kw': 'Junior Python Developer', 'loc': 'Gurgaon, Haryana, India', 'extra': '&f_E=2&f_JT=F'},
    {'kw': 'Junior Machine Learning Engineer', 'loc': 'Gurgaon, Haryana, India', 'extra': '&f_E=2&f_JT=F'},
    {'kw': 'Software Engineer I', 'loc': 'Delhi, India', 'extra': '&f_E=2&f_JT=F'},
    {'kw': 'Associate Data Engineer', 'loc': 'Delhi, India', 'extra': '&f_E=2&f_JT=F'},
    # National Tech Hubs & Remote
    {'kw': 'Software Engineer I', 'loc': 'Remote', 'extra': '&f_WT=2&f_E=2&f_JT=F'},
    {'kw': 'Graduate Software Engineer', 'loc': 'Bengaluru, Karnataka, India', 'extra': '&f_JT=F'},
    {'kw': 'Software Development Engineer I', 'loc': 'Bengaluru, Karnataka, India', 'extra': '&f_E=2&f_JT=F'},
    {'kw': 'Associate Software Engineer', 'loc': 'Bengaluru, Karnataka, India', 'extra': '&f_E=2&f_JT=F'},
    {'kw': 'Associate Data Engineer', 'loc': 'Bengaluru, Karnataka, India', 'extra': '&f_E=2&f_JT=F'},
    {'kw': 'Graduate Engineer Trainee', 'loc': 'Bengaluru, Karnataka, India', 'extra': '&f_JT=F'},
    {'kw': 'Junior Software Engineer', 'loc': 'Pune, Maharashtra, India', 'extra': '&f_E=2&f_JT=F'},
    {'kw': 'Software Engineer I', 'loc': 'Hyderabad, Telangana, India', 'extra': '&f_E=2&f_JT=F'}
]

reject_title_words = [
    'intern', 'internship', 'summer', 'winter', 'lead', 'senior', 'sr.', 'sr ', 'manager',
    'principal', 'director', 'staff', 'ii', 'iii', '2', '3', '4', 'recruiter', 'hr', 'sales',
    'marketing', 'business development', 'product management', 'operations', 'telecaller',
    'civil', 'mechanical', 'support representative'
]

fresh_grad_markers = [
    'graduate engineer trainee', 'get', 'campus', '2026', 'university graduate',
    'fresh grad', 'fresher', 'early career', 'trainee', 'graduate software engineer'
]

driver = webdriver.Chrome(options=options)
harvested = {}

try:
    for idx, s in enumerate(searches):
        kw = s['kw']
        loc = s['loc']
        extra = s.get('extra', '')
        url = f"https://www.linkedin.com/jobs/search?keywords={urllib.parse.quote(kw)}&location={urllib.parse.quote(loc)}{extra}"
        print(f"[{idx+1}/{len(searches)}] Searching: '{kw}' in '{loc}'...")
        try:
            driver.get(url)
            time.sleep(1.8)
            soup = BeautifulSoup(driver.page_source, 'html.parser')
            cards = soup.find_all('div', class_='base-card')
            found_count = 0
            for c in cards:
                t = c.find('h3', class_='base-search-card__title')
                comp = c.find('h4', class_='base-search-card__subtitle')
                l = c.find('span', class_='job-search-card__location')
                a = c.find('a', class_='base-card__full-link')

                title = t.get_text(strip=True) if t else ''
                company = comp.get_text(strip=True) if comp else ''
                location = l.get_text(strip=True) if l else ''
                link = a['href'] if a and a.has_attr('href') else ''

                if not link:
                    continue

                clean_url = link.split('?')[0]
                t_lower = title.lower()
                c_lower = company.lower()

                # Check paused companies
                if any(p in c_lower for p in paused_companies):
                    continue

                # Exclude reject words in title
                if any(w in t_lower for w in reject_title_words):
                    continue

                if clean_url.lower().rstrip('/') in seen_urls:
                    continue

                # NEW COMPANY RULE:
                # Avoid company if already listed in prior waves, UNLESS it's an explicit fresh grad role
                is_fresh_grad_role = any(m in t_lower for m in fresh_grad_markers)
                is_existing_company = any(ec in c_lower or c_lower in ec for ec in existing_companies)
                
                if is_existing_company and not is_fresh_grad_role:
                    continue

                m = re.search(r'-([0-9]{8,12})(?:$)', clean_url)
                jid = m.group(1) if m else clean_url

                if jid not in harvested:
                    harvested[jid] = {
                        'company': company,
                        'title': title,
                        'location': location,
                        'url': clean_url,
                        'jid': jid,
                        'is_existing_comp': is_existing_company,
                        'is_fresh_grad_role': is_fresh_grad_role
                    }
                    found_count += 1
            print(f"   -> Added {found_count} new candidates (Total unique: {len(harvested)})")
        except Exception as e:
            print(f"   -> Search error: {e}")
        time.sleep(0.5)

    print(f"\nTotal candidate leads harvested: {len(harvested)}")
    with open('data/scratch_wave10_harvested_cards.json', 'w') as f:
        json.dump(list(harvested.values()), f, indent=2)

finally:
    driver.quit()
