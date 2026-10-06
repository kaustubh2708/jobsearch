import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

wb = openpyxl.load_workbook('data/Job Search.xlsx')

if 'New Additions (Wave 5)' in wb.sheetnames:
    del wb['New Additions (Wave 5)']

ws = wb.create_sheet('New Additions (Wave 5)')

headers = ['Company', 'Role Title', 'Location', 'Domain / Track', 'Exact Stated Experience in JD (0–2 YOE)', 'Technical Stack & Fit for Vaanya', 'Status / Match Score', 'Direct Application URL']
ws.append(headers)

header_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
header_font = Font(color='FFFFFF', bold=True, size=11)
for col in range(1, len(headers) + 1):
    c = ws.cell(row=1, column=col)
    c.fill = header_fill
    c.font = header_font
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

green_fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
light_green_fill = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')
bold_green = Font(color='006100', bold=True)

# 1. SOFTWARE ENGINEERING & GET ROLES (0-2 YOE Full-Time)
swe_roles = [
    {
        'company': 'Urban Company',
        'title': 'Software Development Engineer 1 (SDE-1)',
        'location': 'Gurugram, Haryana (In-Office)',
        'track': 'Software Engineering (Backend)',
        'yoe': '0–2 years of experience in software engineering. Targets engineering graduates writing production code.',
        'stack': 'Python Backend & Microservices: Python, PostgreSQL, Microservices, high-scale distributed systems. Prime Gurugram HQ (NCR). Full-time CTC: ₹16-24 LPA.',
        'score': '89 / 100 (Strong Match)',
        'url': 'https://www.urbancompany.com/careers',
        'is_strong': True
    },
    {
        'company': 'Hewlett Packard Enterprise (HPE)',
        'title': 'Graduate Software Engineer',
        'location': 'Bengaluru (Designated Remote / Teleworker)',
        'track': 'Software Engineering (Systems/Cloud)',
        'yoe': 'Typically 0-2 years experience. Bachelor\'s or Master\'s degree in CS, IS, or equivalent.',
        'stack': 'Python & Systems Engineering: Python, C++, Java, Linux, systems software, database concepts, debugging & test automation. Requisition ID: 1210218 (posted Oct 1, 2026). Designated remote teleworker across India. Full-time CTC: ₹12-16 LPA.',
        'score': '88 / 100 (Strong Match)',
        'url': 'https://careers.hpe.com/us/en/job/1210218/Graduate-Software-Engineer',
        'is_strong': True
    },
    {
        'company': 'Prismagic Solutions Inc.',
        'title': 'Graduate Trainee – PV-Technology & AI/ML R&D',
        'location': 'Noida, Uttar Pradesh, India',
        'track': 'Graduate Engineer Trainee (GET)',
        'yoe': 'Experience: 0–1 year | Freshers welcome. Education: B.Tech / M.Tech / MCA.',
        'stack': 'AI/ML & Python: Python, machine learning R&D, data pipelines. Permanent technology team in Noida (minutes from her college Sector 62/128).',
        'score': '86 / 100 (Strong Match)',
        'url': 'https://in.linkedin.com/jobs/view/graduate-trainee-%E2%80%93-pv-technology-ai-ml-r-d-at-prismagic-solutions-inc-4470096991',
        'is_strong': True
    },
    {
        'company': 'Zepto',
        'title': 'Data Scientist',
        'location': 'Bengaluru, Karnataka',
        'track': 'AI/ML & Data Modeling',
        'yoe': 'Experience: 0-2 years (0 YOE / freshers explicitly eligible). Focus on data-backed experimentation.',
        'stack': 'Data Engineering & ML: Python, Machine Learning, Statistical Modeling, ML Pipelines, SQL. Instahyre Requisition ID: 264559. Full-time CTC: ₹15-22 LPA.',
        'score': '85 / 100 (Strong Match)',
        'url': 'https://www.instahyre.com/job/264559-data-scientist-at-zepto-bangalore/',
        'is_strong': True
    },
    {
        'company': 'Squarepoint Capital',
        'title': 'Junior Quant Researcher - ML Alpha Research',
        'location': 'Bangalore, India (Global Quant Division)',
        'track': 'Quant ML Research',
        'yoe': '0-2 years / Recent graduate in CS, Math, Engineering or quantitative discipline.',
        'stack': 'Quant ML & Time Series: Python, PyTorch, statistical machine learning, high-performance data processing. Top-tier global quantitative hedge fund.',
        'score': '75 / 100 (Potential Match)',
        'url': 'https://www.squarepoint-capital.com/open-opportunities?id=6069464&gh_jid=6069464',
        'is_strong': False
    },
    {
        'company': 'Ingersoll Rand',
        'title': 'Graduate Engineer Trainee',
        'location': 'Gurgaon, Haryana, India',
        'track': 'Graduate Engineer Trainee (GET)',
        'yoe': 'B.Tech / B.E. Fresher (0 YOE) Graduate Engineer Trainee program.',
        'stack': 'Engineering Systems: Structured GET campus intake for B.Tech engineers. Testing, software diagnostics, engineering automation in Gurgaon.',
        'score': '74 / 100 (Potential Match)',
        'url': 'https://in.linkedin.com/jobs/view/graduate-engineer-trainee-at-ingersoll-rand-4460474501',
        'is_strong': False
    },
    {
        'company': 'GreyOrange',
        'title': 'Graduate Engineer Trainee',
        'location': 'Gurugram, Haryana, India',
        'track': 'Graduate Engineer Trainee (Firmware/QA)',
        'yoe': 'Graduate Engineer Trainee (0 YOE) – Firmware Quality Assurance to support embedded validation.',
        'stack': 'Firmware & QA Automation: Embedded firmware testing, Python test suites, robotics systems QA in Gurugram.',
        'score': '67 / 100 (Potential Match)',
        'url': 'https://in.linkedin.com/jobs/view/graduate-engineer-trainee-at-greyorange-4464347344',
        'is_strong': False
    }
]

# 2. DATA ENGINEERING & ANALYTICS ROLES (0-2 YOE Full-Time)
de_roles = [
    {
        'company': 'Amazon',
        'title': 'Data Engineer I, CMT',
        'location': 'Bengaluru, Karnataka, India',
        'track': 'Data Engineering (Full-Time)',
        'yoe': 'Foundational entry-level engineering role (0–2 YOE / College graduates in CS, Engineering, Math).',
        'stack': 'Big Data & Cloud: Python, SQL, AWS (S3, Redshift, Glue), distributed data pipelines, ETL architecture. Amazon Data Platform & Analytics division.',
        'score': '91 / 100 (Strong Match)',
        'url': 'https://in.linkedin.com/jobs/view/data-engineer-i-cmt-at-amazon-4455916247',
        'is_strong': True
    },
    {
        'company': 'NPS Prism (Bain & Co. CoE)',
        'title': 'Associate, Product Operations (Python, PySpark, SQL)',
        'location': 'Gurugram, Haryana, India',
        'track': 'Data Engineering & Analytics',
        'yoe': '0–2 years of experience in Python, PySpark, SQL, and data transformation pipelines.',
        'stack': 'ETL & Analytics: Python, PySpark, SQL, PostgreSQL, automated data pipelines, business metrics. Prime Gurugram location (NCR). Mirrors S&P Global and KPMG advisory background.',
        'score': '87 / 100 (Strong Match)',
        'url': 'https://in.linkedin.com/jobs/view/associate-product-operations-tableau-pyspark-or-python-and-sql-at-nps-prism-4428538771',
        'is_strong': True
    },
    {
        'company': 'PwC India',
        'title': 'Associate — Data, Analytics & AI',
        'location': 'New Delhi, Delhi, India',
        'track': 'Data & AI Engineering (Full-Time)',
        'yoe': 'Entry-level Associate track for B.Tech freshers (0–1 YOE) in Data, Analytics & AI.',
        'stack': 'Data Platforms & AI: Python, SQL, database modeling, cloud ETL pipelines, machine learning integration. New Delhi office (NCR).',
        'score': '83 / 100 (Strong Match)',
        'url': 'https://in.linkedin.com/jobs/view/associate-at-pwc-india-4291067264',
        'is_strong': True
    },
    {
        'company': 'EXL',
        'title': 'Associate - Cloud Data Engineering',
        'location': 'Noida, Uttar Pradesh, India',
        'track': 'Data Engineering & Cloud',
        'yoe': '0–2 years of experience. Verbatim: "Strong hands-on programming skills in Python... Docker, Kubernetes, and cloud data pipelines."',
        'stack': 'Python & Cloud ETL: Python, SQL, Docker, Kubernetes, cloud data ingestion pipelines. Noida location (minutes from her college Sector 62/128).',
        'score': '82 / 100 (Strong Match)',
        'url': 'https://in.linkedin.com/jobs/view/associate-business-analyst-data-engineering-cloud-data-engineering-at-exl-4446740301',
        'is_strong': True
    },
    {
        'company': 'PwC Acceleration Center India',
        'title': 'Financial Market Data Engineer – Associate',
        'location': 'Bengaluru East, Karnataka, India',
        'track': 'Data Engineering (Financial Data)',
        'yoe': 'Entry-level Associate track for engineering graduates (0–2 YOE) in financial analytics.',
        'stack': 'Financial ETL & Data: SQL, Python, relational databases, data extraction pipelines. Direct match for Vaanya\'s S&P Global financial document routing experience.',
        'score': '79 / 100 (Potential Match)',
        'url': 'https://in.linkedin.com/jobs/view/assurance-financial-market-%E2%80%93-data-engineer-%E2%80%93-associate-at-pwc-acceleration-center-india-4465512947',
        'is_strong': False
    },
    {
        'company': 'Portage Point Partners',
        'title': 'Associate, Data Analytics // DevOps',
        'location': 'Gurugram, Haryana, India',
        'track': 'Data Engineering & Analytics',
        'yoe': '0–2 years of experience. Chicago advisory firm tech division in Gurgaon.',
        'stack': 'Data Infrastructure & SQL: Python, SQL, relational databases, data pipelines, DevOps automation.',
        'score': '76 / 100 (Potential Match)',
        'url': 'https://in.linkedin.com/jobs/view/associate-data-analytics-devops-at-portage-point-partners-4432722173',
        'is_strong': False
    },
    {
        'company': 'Axtria - Ingenious Insights',
        'title': 'Associate — Data & Analytics',
        'location': 'Gurgaon, Haryana, India',
        'track': 'Data Analytics & Engineering',
        'yoe': '0–2 years of experience in data analytics, SQL, and database management.',
        'stack': 'Data Analytics & ETL: Python, SQL, data warehousing, healthcare analytics pipelines in Gurgaon.',
        'score': '73 / 100 (Potential Match)',
        'url': 'https://in.linkedin.com/jobs/view/associate-at-axtria-ingenious-insights-4418733398',
        'is_strong': False
    }
]

# Write SWE section
row_idx = 2
for r in swe_roles:
    ws.append([r['company'], r['title'], r['location'], r['track'], r['yoe'], r['stack'], r['score'], r['url']])
    fill = green_fill if r['is_strong'] else light_green_fill
    for c in range(1, 9):
        cell = ws.cell(row=row_idx, column=c)
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical='top')
    ws.cell(row=row_idx, column=7).font = bold_green
    url_c = ws.cell(row=row_idx, column=8)
    url_c.hyperlink = r['url']
    url_c.font = Font(color='0563C1', underline='single')
    row_idx += 1

# Section divider for Data Engineering
ws.cell(row=row_idx, column=1).value = "DATA ENGINEERING & ANALYTICS ROLES (0–2 YOE Full-Time / Permanent)"
ws.cell(row=row_idx, column=1).font = Font(bold=True, color='002060', size=11)
ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=8)
ws.cell(row=row_idx, column=1).fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
row_idx += 1

# Write Data Engineering section
for r in de_roles:
    ws.append([r['company'], r['title'], r['location'], r['track'], r['yoe'], r['stack'], r['score'], r['url']])
    fill = green_fill if r['is_strong'] else light_green_fill
    for c in range(1, 9):
        cell = ws.cell(row=row_idx, column=c)
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical='top')
    ws.cell(row=row_idx, column=7).font = bold_green
    url_c = ws.cell(row=row_idx, column=8)
    url_c.hyperlink = r['url']
    url_c.font = Font(color='0563C1', underline='single')
    row_idx += 1

col_widths = [24, 38, 22, 28, 45, 60, 22, 32]
for i, w in enumerate(col_widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws.row_dimensions[1].height = 28
wb.save('data/Job Search.xlsx')
print(f"Successfully finalized Wave 5: {len(swe_roles)} Software Engineering roles + {len(de_roles)} Data Engineering roles (Total: {len(swe_roles) + len(de_roles)}).")
