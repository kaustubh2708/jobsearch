import openpyxl

wb = openpyxl.load_workbook('data/Job Search.xlsx')
existing_urls = set()
for sname in wb.sheetnames:
    ws = wb[sname]
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(r, c)
            if cell.hyperlink and cell.hyperlink.target:
                existing_urls.add(cell.hyperlink.target.lower().rstrip('/'))
            elif cell.value and isinstance(cell.value, str) and cell.value.startswith('http'):
                existing_urls.add(cell.value.lower().rstrip('/'))

de_roles = [
    {'comp': 'Amazon', 'title': 'Data Engineer I, CMT', 'url': 'https://in.linkedin.com/jobs/view/data-engineer-i-cmt-at-amazon-4455916247'},
    {'comp': 'NPS Prism', 'title': 'Associate, Product Operations (Python, PySpark, SQL)', 'url': 'https://in.linkedin.com/jobs/view/associate-product-operations-tableau-pyspark-or-python-and-sql-at-nps-prism-4428538771'},
    {'comp': 'PwC India', 'title': 'Associate — Data, Analytics & AI', 'url': 'https://in.linkedin.com/jobs/view/associate-at-pwc-india-4291067264'},
    {'comp': 'EXL', 'title': 'Associate - Cloud Data Engineering', 'url': 'https://in.linkedin.com/jobs/view/associate-business-analyst-data-engineering-cloud-data-engineering-at-exl-4446740301'},
    {'comp': 'PwC Acceleration Center', 'title': 'Financial Market Data Engineer – Associate', 'url': 'https://in.linkedin.com/jobs/view/assurance-financial-market-%E2%80%93-data-engineer-%E2%80%93-associate-at-pwc-acceleration-center-india-4465512947'},
    {'comp': 'Portage Point Partners', 'title': 'Associate, Data Analytics // DevOps', 'url': 'https://in.linkedin.com/jobs/view/associate-data-analytics-devops-at-portage-point-partners-4432722173'},
    {'comp': 'Axtria', 'title': 'Associate — Data & Analytics', 'url': 'https://in.linkedin.com/jobs/view/associate-at-axtria-ingenious-insights-4418733398'}
]

for d in de_roles:
    u = d['url'].lower().rstrip('/')
    is_dup = u in existing_urls
    print(f"[{'DUPLICATE' if is_dup else 'UNIQUE'}] {d['comp']} - {d['title']}")
