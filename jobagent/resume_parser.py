"""Read a resume in whatever format it happens to be in, then use the local
LLM to turn it into a structured skill + experience library."""
from __future__ import annotations

import json
import logging
import re
import zipfile
from pathlib import Path
from typing import Any, Dict, List, Optional

from sqlmodel import select

from .config import ROOT, load_config
from .db import SkillLibrary, get_session, log as dblog
from .llm import get_llm

log = logging.getLogger("jobagent.resume")

SUPPORTED = (".pdf", ".docx", ".txt", ".md", ".pages")


# --------------------------------------------------------------------------- #
#  Text extraction
# --------------------------------------------------------------------------- #
def find_resume(cfg=None) -> Optional[Path]:
    cfg = cfg or load_config()
    d = cfg.resume_dir
    if not d.exists():
        return None
    # prefer text-extractable formats, newest first
    order = {".pdf": 0, ".docx": 1, ".txt": 2, ".md": 3, ".pages": 4}
    files = [p for p in d.iterdir() if p.suffix.lower() in SUPPORTED and not p.name.startswith(".")]
    if not files:
        return None
    files.sort(key=lambda p: (order.get(p.suffix.lower(), 9), -p.stat().st_mtime))
    return files[0]


def extract_text(path: Path) -> str:
    suf = path.suffix.lower()
    if suf in (".txt", ".md"):
        return path.read_text(errors="ignore")
    if suf == ".pdf":
        return _pdf_text(path)
    if suf == ".docx":
        return _docx_text(path)
    if suf == ".pages":
        return _pages_text(path)
    raise ValueError(f"unsupported resume format: {suf}")


def _pdf_text(path: Path) -> str:
    try:
        import pdfplumber

        with pdfplumber.open(path) as pdf:
            return "\n".join((p.extract_text() or "") for p in pdf.pages)
    except Exception:
        from pypdf import PdfReader

        return "\n".join((pg.extract_text() or "") for pg in PdfReader(str(path)).pages)


def _docx_text(path: Path) -> str:
    import docx

    doc = docx.Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


def _pages_text(path: Path) -> str:
    """Apple Pages = zip of snappy-compressed protobuf (IWA). We decompress the
    frames and pull out readable text runs. Good enough for resume content."""
    try:
        import snappy
    except ImportError:  # pragma: no cover
        raise RuntimeError(
            "Reading .pages needs python-snappy. Either `pip install python-snappy` "
            "or export your resume to PDF/DOCX into the resume/ folder."
        )

    def _decompress(buf: bytes) -> bytes:
        out, i = b"", 0
        while i + 4 <= len(buf):
            ln = buf[i + 1] | (buf[i + 2] << 8) | (buf[i + 3] << 16)
            try:
                out += snappy.uncompress(buf[i + 4 : i + 4 + ln])
            except Exception:
                pass
            i += 4 + ln
        return out

    chunks: List[str] = []
    seen = set()
    with zipfile.ZipFile(path) as z:
        names = [n for n in z.namelist() if n.startswith("Index/") and n.endswith(".iwa")]
        for n in sorted(names):
            blob = _decompress(z.read(n))
            for m in re.findall(rb"[\x20-\x7e]{25,}", blob):
                s = m.decode("utf-8", "ignore")
                # drop iWork style-table noise
                if re.match(r"^[\W]?(text|table|chart|caption|imported|sticky|drawing)[-\w ]*$", s, re.I):
                    continue
                if s in seen:
                    continue
                seen.add(s)
                chunks.append(s)
    return "\n".join(chunks)


# --------------------------------------------------------------------------- #
#  Structuring
# --------------------------------------------------------------------------- #
EXTRACT_SYSTEM = """You are a precise resume analyst. You extract structured facts from a resume.
You never invent information. If something is not in the resume, omit it or use null.
You always reply with a single valid JSON object and nothing else."""

EXTRACT_PROMPT = """Extract a complete structured profile from this resume.

Return JSON with exactly this shape:
{{
  "headline": "one-line professional identity",
  "summary": "3-4 sentence positioning summary written in third person",
  "total_years_experience": number,
  "seniority": "junior|mid|senior|staff|lead|principal",
  "current_role": "title @ company",
  "target_roles": ["8-14 job titles this person should realistically be searching for on Indian job boards, ordered by fit. Use titles that actually appear in job postings."],
  "target_role_keywords": ["6-10 short boolean-ish search phrases for job boards, e.g. 'machine learning engineer', 'LLM engineer'"],
  "domains": ["industry/problem domains, e.g. 'voice AI', 'data governance'"],
  "hard_skills": ["every concrete technical skill, tool, framework, language, platform mentioned"],
  "core_skills": ["the 12-18 skills that most define this candidate's value"],
  "soft_skills": [],
  "certifications": [],
  "education": [{{"degree": "", "field": "", "institution": "", "year": ""}}],
  "experience": [{{"title": "", "company": "", "duration": "", "highlights": ["quantified achievements"]}}],
  "projects": [{{"name": "", "tech": [], "summary": ""}}],
  "publications": [],
  "achievements": ["quantified outcomes: percentages, scale, revenue, users"],
  "notable_companies": ["recognisable employer names"],
  "preferred_locations": [],
  "red_lines": ["role types this person should NOT be matched to, inferred from seniority and background"]
}}

RESUME:
---
{resume}
---"""


def build_skill_library(text: str, source_file: str = "") -> Dict[str, Any]:
    llm = get_llm()
    data = llm.chat_json(
        EXTRACT_SYSTEM,
        EXTRACT_PROMPT.format(resume=text[:22000]),
        fallback={},
        num_ctx=16384,
    )
    if not data:
        raise RuntimeError("LLM returned nothing for resume extraction — is Ollama running?")

    # normalise
    for k in ("target_roles", "target_role_keywords", "hard_skills", "core_skills",
              "soft_skills", "domains", "achievements", "notable_companies",
              "preferred_locations", "red_lines", "certifications", "publications"):
        v = data.get(k)
        data[k] = [str(x).strip() for x in v if str(x).strip()] if isinstance(v, list) else []
    for k in ("education", "experience", "projects"):
        v = data.get(k)
        data[k] = v if isinstance(v, list) else []
    data["source_file"] = source_file
    return data


def profile_embedding_text(data: Dict[str, Any]) -> str:
    """The text we embed to represent 'what this candidate is'. Deliberately
    weighted toward skills + roles, since that's what JDs talk about."""
    parts = [
        data.get("headline", ""),
        data.get("summary", ""),
        "Target roles: " + ", ".join(data.get("target_roles", [])),
        "Core skills: " + ", ".join(data.get("core_skills", [])),
        "Skills: " + ", ".join(data.get("hard_skills", [])[:80]),
        "Domains: " + ", ".join(data.get("domains", [])),
    ]
    for e in (data.get("experience") or [])[:5]:
        parts.append(f"{e.get('title','')} at {e.get('company','')}: " + " ".join(e.get("highlights", [])[:3]))
    return "\n".join(p for p in parts if p)


def ingest(path: Optional[Path] = None, force: bool = False) -> SkillLibrary:
    """Parse the resume and persist an active SkillLibrary row."""
    cfg = load_config()
    path = path or find_resume(cfg)
    if not path:
        raise FileNotFoundError(
            f"No resume found in {cfg.resume_dir}. Drop a .pdf/.docx/.txt/.pages file there."
        )

    text = extract_text(path)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) < 200:
        raise ValueError(f"Extracted only {len(text)} chars from {path.name} — try exporting to PDF instead.")

    log.info("Parsed %s (%d chars). Asking local LLM to structure it...", path.name, len(text))
    data = build_skill_library(text, source_file=str(path))

    llm = get_llm()
    emb = llm.embed(profile_embedding_text(data))

    with get_session() as s:
        for old in s.exec(select(SkillLibrary).where(SkillLibrary.is_active == True)).all():  # noqa: E712
            old.is_active = False
            s.add(old)
        row = SkillLibrary(
            source_file=str(path),
            raw_text=text,
            payload=json.dumps(data, ensure_ascii=False),
            embedding=json.dumps(emb),
            is_active=True,
        )
        s.add(row)
        s.commit()
        s.refresh(row)

    dblog("info", f"Skill library built from {path.name}",
          skills=len(data.get("hard_skills", [])), roles=len(data.get("target_roles", [])))
    return row


def get_or_build() -> SkillLibrary:
    with get_session() as s:
        row = s.exec(
            select(SkillLibrary).where(SkillLibrary.is_active == True)  # noqa: E712
            .order_by(SkillLibrary.created_at.desc())
        ).first()
    return row or ingest()
